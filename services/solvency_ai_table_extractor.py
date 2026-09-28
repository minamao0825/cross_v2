from __future__ import annotations

import base64
import io
import json
import math
import re
from dataclasses import dataclass
from typing import Callable, Mapping

import fitz
import pandas as pd
import pdfplumber
import requests
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .llm_config import model_request_parameters, normalize_model_id
from .llm_http import post_json_with_retry
from .solvency_disclosure_normalizer import (
    normalize_three_year_return_rows_core,
)
from .solvency_pdf_locator import PageMatch
from .solvency_table_boundaries import (
    TableBoundaryError,
    boundary_instruction_for_table,
    boundary_items,
    boundary_variants,
    end_item_groups,
    enforce_output_boundaries,
    item_hits_in_rows,
    line_has_item,
)
from .solvency_table_extractor import (
    TABLE_EXCLUSIONS,
    TABLE_HEADERS,
    TABLE_SIGNATURES,
    TABLE_START_MARKERS,
    ExtractedTable,
    UnitRecord,
)
from .table_strategy_handlers import (
    BoundaryRequest,
    CompletenessRequest,
    PostprocessRequest,
    PromptRequest,
)
from .table_strategy_registry import (
    active_table_strategy,
    resolve_table_strategy,
)

GRID_COLUMNS = 180
TEXT_MODE = "大模型网格重构"
VISION_MODE = "图片扫描重试"
VISION_HIGH_RES_MODE = "3倍高清图片扫描重试"
TEXT_RETRY_MODE = "网格纠错重试"
TABLE_COMPLETENESS_TERMS = {
    "SOLVENCY_MAIN": (
        "核心偿付能力充足率",
        "综合偿付能力充足率",
        "实际资本",
        "最低资本",
    ),
    "ACTUAL_CAPITAL": (
        "财务报表资产总额",
        "认可资产",
        "财务报表负债总额",
        "认可负债",
        "财务报表净资产总额",
        "实际资本",
        "核心一级资本",
        "核心二级资本",
        "附属一级资本",
        "附属二级资本",
        "实际资本合计",
    ),
    "RECOGNIZED_ASSETS": (
        "现金及流动性管理工具",
        "投资资产",
        "再保险资产",
        "认可资产合计",
    ),
}

TABLE_MIN_NUMERIC_ROWS = {
    "SOLVENCY_MAIN": 4,
}
SOURCE_ROW_COVERAGE_RATIO = 0.60
SOURCE_ITEM_RECALL_RATIO = 0.90
OPERATING_SECTION_TITLES = ("效益类指标", "规模类指标", "品质类指标")

AI_TABLE_END_MARKERS = {
    "SOLVENCY_MAIN": ("监管指标名称", "流动性覆盖率", "LCR1"),
    "OPERATING_METRICS": (
        "流动性风险监管指标", "流动性风险监测指标",
        "近三年综合投资收益率", "近三年（综合）投资收益率",
    ),
    "ACTUAL_CAPITAL": (
        "S03",
        "（三）认可资产",
        "(三)认可资产",
        "认可资产表",
        "认可负债表",
        "最低资本表",
    ),
    "THREE_YEAR_INVESTMENT_RETURN": ("S02", "实际资本表", "实际资本明细表", "最低资本表"),
    "RECOGNIZED_ASSETS": ("S04", "认可负债表", "最低资本表"),
    "MINIMUM_CAPITAL": ("S06", "风险综合评级", "偿付能力风险管理评估", "风险管理能力"),
}


class ExtractionQualityError(RuntimeError):
    pass


@dataclass
class ExtractionLog:
    table_id: str
    table_name: str
    pages: list[int]
    mode: str
    status: str
    message: str

    def to_dict(self) -> dict:
        return {
            "目标表": self.table_name,
            "物理页码": "、".join(map(str, self.pages)),
            "识别方式": self.mode,
            "状态": self.status,
            "说明": self.message,
        }


@dataclass
class AIExtractionBundle:
    tables: list[ExtractedTable]
    logs: list[ExtractionLog]


@dataclass
class PageGrid:
    page_number: int
    grid_text: str
    word_count: int


def _clean(value) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\u3000", " ")).strip()


def _compact(value: str) -> str:
    return re.sub(r"[\s：:（）()、，,。·—\-_/]", "", str(value or ""))

_UNIT_TOKEN = (
    r"(?:\u4eba\u6c11\u5e01\s*)?"
    r"(?:\u4ebf\u5143|\u4e07\u5143|\u5343\u5143|\u5143|%|\uff05|"
    r"\u767e\u5206\u6bd4|\u767e\u5206\u70b9|\u4eba|\u6237|\u4ef6|\u6b21|\u7ea7)"
)
_UNIT_DECLARATION_RE = re.compile(
    rf"(?:\u91d1\u989d)?\u5355\u4f4d\s*(?:[\uff1a:]\s*)?(?P<unit>{_UNIT_TOKEN})",
    flags=re.I,
)
_UNIT_PAREN_RE = re.compile(
    rf"[\uff08(]\s*(?P<unit>{_UNIT_TOKEN})\s*[\uff09)]",
    flags=re.I,
)
_UNIT_GROUP_PAREN_RE = re.compile(
    r"[\uff08(](?P<body>[^\uff08\uff09()]{1,40})[\uff09)]",
    flags=re.I,
)
_SUPPORTED_UNITS = {
    "\u4ebf\u5143", "\u4e07\u5143", "\u5343\u5143", "\u5143", "%",
    "\u767e\u5206\u70b9", "\u4eba", "\u6237", "\u4ef6", "\u6b21", "\u7ea7",
}
_PERCENT_LABEL_TERMS = (
    "\u7387", "\u6bd4\u4f8b", "\u5360\u6bd4", "\u589e\u901f", "\u589e\u957f",
)


def _normalize_unit(value: str) -> str:
    text = re.sub(r"[\s\uff1a:\uff08\uff09()]", "", str(value or ""))
    text = text.replace("\uff05", "%")
    if "\u767e\u5206\u6bd4" in text:
        return "%"
    if "\u767e\u5206\u70b9" in text:
        return "\u767e\u5206\u70b9"
    for unit in (
        "\u4ebf\u5143", "\u4e07\u5143", "\u5343\u5143", "\u5143", "%",
        "\u4eba", "\u6237", "\u4ef6", "\u6b21", "\u7ea7",
    ):
        if text.endswith(unit):
            return unit
    return text


def _unit_raw(value: str) -> str:
    text = re.sub(r"\s+", "", str(value or "")).strip("\uff1a:")
    return text or str(value or "").strip()


def _units_from_parenthetical(value: str) -> list[str]:
    """Read one or more explicit units from a pure parenthetical unit group."""
    units: list[str] = []
    for group in _UNIT_GROUP_PAREN_RE.finditer(str(value or "")):
        body = group.group("body")
        matches = list(re.finditer(_UNIT_TOKEN, body, flags=re.I))
        if not matches:
            continue
        remainder = re.sub(_UNIT_TOKEN, "", body, flags=re.I)
        # Only accept a parenthesis made entirely of units and separators.
        # This excludes narrative parentheses and section labels.
        if re.sub(r"[\s,，、/；;和及]+", "", remainder):
            continue
        units.extend(_unit_raw(match.group(0)) for match in matches)
    return list(dict.fromkeys(units))


def _unit_target_key(value: str) -> str:
    text = _UNIT_PAREN_RE.sub("", str(value or ""))
    normalized = _normalize_unit(text)
    if normalized in _SUPPORTED_UNITS:
        text = re.sub(
            rf"(?:{_UNIT_TOKEN})\s*$",
            "",
            text,
            flags=re.I,
        )
    return _compact(text)


def _row_label_index(row: list[str]) -> int:
    return (
        1
        if len(row) > 1 and re.fullmatch(r"\d+(?:\.\d+)*\*?", row[0].strip())
        else 0
    )


def _row_label(row: list[str]) -> str:
    if not row:
        return ""
    label_index = _row_label_index(row)
    return _clean(row[label_index] if label_index < len(row) else "")


def merge_unit_records(*groups: list[UnitRecord]) -> list[UnitRecord]:
    unique: list[UnitRecord] = []
    seen: set[tuple] = set()
    for records in groups:
        for record in records:
            key = record.dedupe_key()
            if key not in seen:
                seen.add(key)
                unique.append(record)
    return unique


def _resolve_unit_target(
    scope: str,
    target: str,
    rows: list[list[str]],
) -> str | None:
    if scope == "\u8868\u7ea7":
        return ""
    candidates = (
        [_clean(cell) for cell in rows[0] if _clean(cell)]
        if scope == "\u5217\u7ea7" and rows
        else [_row_label(row) for row in rows[1:]]
    )
    candidates = [candidate for candidate in candidates if candidate]
    target_key = _unit_target_key(target)
    if not target_key:
        return None
    exact = next(
        (
            candidate
            for candidate in candidates
            if _unit_target_key(candidate) == target_key
        ),
        None,
    )
    if exact:
        return exact
    contained = [
        candidate
        for candidate in candidates
        if (
            target_key in _unit_target_key(candidate)
            or _unit_target_key(candidate) in target_key
        )
    ]
    if not contained:
        return None
    return min(
        contained,
        key=lambda candidate: abs(len(_unit_target_key(candidate)) - len(target_key)),
    )


def unit_records_from_payload(
    payload: dict,
    rows: list[list[str]],
    pages: list[int],
    *,
    source_type: str = "\u56fe\u7247\u7ed3\u6784\u5316\u5355\u4f4d",
) -> list[UnitRecord]:
    raw_records = payload.get("unit_records", payload.get("units", []))
    if not isinstance(raw_records, list):
        return []
    page_set = set(pages)
    default_page = pages[0] if pages else 0
    scope_aliases = {
        "\u8868": "\u8868\u7ea7",
        "\u8868\u7ea7": "\u8868\u7ea7",
        "table": "\u8868\u7ea7",
        "\u5217": "\u5217\u7ea7",
        "\u5217\u7ea7": "\u5217\u7ea7",
        "column": "\u5217\u7ea7",
        "\u884c": "\u884c\u7ea7",
        "\u884c\u7ea7": "\u884c\u7ea7",
        "row": "\u884c\u7ea7",
    }
    records: list[UnitRecord] = []
    for item in raw_records:
        if not isinstance(item, dict):
            continue
        scope = scope_aliases.get(
            str(item.get("scope", item.get("\u4f5c\u7528\u8303\u56f4", ""))).strip().lower()
        )
        if not scope:
            continue
        raw = _unit_raw(
            item.get(
                "raw_unit",
                item.get(
                    "unit",
                    item.get("\u539f\u59cb\u5355\u4f4d", item.get("normalized_unit", "")),
                ),
            )
        )
        normalized = _normalize_unit(raw)
        if normalized not in _SUPPORTED_UNITS:
            continue
        target = str(
            item.get("target", item.get("\u5bf9\u8c61", item.get("label", ""))) or ""
        ).strip()
        resolved_target = _resolve_unit_target(scope, target, rows)
        if resolved_target is None:
            continue
        try:
            source_page = int(
                item.get("source_page", item.get("\u7269\u7406\u9875\u7801", default_page))
                or default_page
            )
        except (TypeError, ValueError):
            source_page = default_page
        if page_set and source_page not in page_set:
            source_page = default_page
        records.append(UnitRecord(
            scope=scope,
            target=resolved_target,
            raw_unit=raw,
            normalized_unit=normalized,
            source_page=source_page,
            source_type=source_type,
            source_text=str(
                item.get("source_text", item.get("\u6765\u6e90\u539f\u6587", "")) or ""
            ).strip(),
            confidence=str(
                item.get("confidence", item.get("\u7f6e\u4fe1\u5ea6", "\u9ad8")) or "\u9ad8"
            ).strip(),
        ))
    return merge_unit_records(records)


def unit_recovery_needed(
    rows: list[list[str]],
    records: list[UnitRecord],
) -> bool:
    if len(rows) < 2:
        return False
    table_units = {
        record.normalized_unit for record in records if record.scope == "\u8868\u7ea7"
    }
    column_units = {
        record.normalized_unit for record in records if record.scope == "\u5217\u7ea7"
    }
    for row in rows[1:]:
        label = _row_label(row)
        if not label:
            continue
        label_index = _row_label_index(row)
        value_cells = row[label_index + 1:]
        if not any(re.search(r"\d", cell) for cell in value_cells):
            continue
        row_key = _unit_target_key(label)
        available = {
            record.normalized_unit
            for record in records
            if (
                record.scope == "\u884c\u7ea7"
                and _unit_target_key(record.target) == row_key
            )
        }
        available.update(table_units)
        available.update(column_units)
        expects_percent = any(term in label for term in _PERCENT_LABEL_TERMS)
        if expects_percent:
            if "%" not in available and "\u767e\u5206\u70b9" not in available:
                return True
        elif not available:
            return True
    return False


def _unit_vision_messages(
    table_id: str,
    table_name: str,
    pages: list[int],
    rows: list[list[str]],
    images: list[tuple[int, str]],
) -> list[dict]:
    columns = rows[0] if rows else []
    row_targets = [_row_label(row) for row in rows[1:] if _row_label(row)]
    prompt = f"""你是保险公司偿付能力报告单位核对专家。
目标表ID：{table_id}
目标表名称：{table_name}
PDF物理页码：{pages}

系统已经提取的列名：
{json.dumps(columns, ensure_ascii=False)}
系统已经提取的项目名：
{json.dumps(row_targets, ensure_ascii=False)}

请只读取随后页面图片中明确出现的单位，不得依据数值大小、行业惯例或项目含义猜测。
1. 表格上方统一单位，或表格标题括号中的单位，使用scope="表级"，target=""。
   标题如“其他经营指标（元，%）”时必须分别输出“元”和“%”两条表级记录；
   金额项目使用元、率类项目使用%，不得把两个单位合成一个字符串。
2. 表头单位使用scope="列级"，target必须从上述列名中原样选择。
3. 项目名内单位使用scope="行级"，target必须从上述项目名中原样选择；即使图片原文项目名与系统项目名略有差异，也要映射到系统项目名。
4. source_text必须保留图片中实际看到的含单位原文，例如“实际资本（元）”。
5. 支持的单位：亿元、万元、千元、元、%、百分点、人、户、件、次、级。
6. 图片中没有明确单位时不要输出该记录。
7. 只输出JSON对象，不输出Markdown或解释文字。

JSON格式：
{{"unit_records":[
  {{"scope":"行级","target":"实际资本合计","raw_unit":"元",
    "source_page":8,"source_text":"实际资本（元）","confidence":"高"}}
]}}
"""
    content: list[dict] = [{"type": "text", "text": prompt}]
    for page_number, data_url in images:
        content.extend([
            {"type": "text", "text": f"PDF物理第 {page_number} 页"},
            {"type": "image_url", "image_url": {"url": data_url, "detail": "high"}},
        ])
    return [
        {
            "role": "system",
            "content": "\u4f60\u53ea\u8d1f\u8d23\u4ece\u9875\u9762\u56fe\u7247\u4e2d\u6284\u5f55\u660e\u786e\u5355\u4f4d\uff0c\u4e0d\u5f97\u63a8\u65ad\u5355\u4f4d\u3002",
        },
        {"role": "user", "content": content},
    ]


def _line_content(line: str) -> str:
    return str(line or "").partition("|")[2] if "|" in str(line or "") else str(line or "")


def _unit_from_label(
    label: str,
    *,
    scope: str,
    target: str,
    source_page: int,
    source_type: str,
) -> list[UnitRecord]:
    records: list[UnitRecord] = []
    for raw in _units_from_parenthetical(label):
        records.append(UnitRecord(
            scope=scope,
            target=target,
            raw_unit=raw,
            normalized_unit=_normalize_unit(raw),
            source_page=source_page,
            source_type=source_type,
            source_text=str(label or "").strip(),
            confidence="\u9ad8",
        ))
    return records


def _unit_records_from_grid(
    table_id: str,
    page_number: int,
    grid_text: str,
) -> list[UnitRecord]:
    lines = str(grid_text or "").splitlines()
    if not lines:
        return []
    sliced_lines = _slice_grid_for_target(table_id, grid_text).splitlines()
    if not sliced_lines:
        return []
    first = next((index for index, line in enumerate(lines) if line == sliced_lines[0]), 0)
    # Standalone unit labels normally sit between the table title and its header.
    context_start = max(0, first - 6)
    context_end = min(len(lines), first + 7)
    records: list[UnitRecord] = []
    for line in lines[context_start:context_end]:
        content = _line_content(line).strip()
        for match in _UNIT_DECLARATION_RE.finditer(content):
            raw = _unit_raw(match.group("unit"))
            records.append(UnitRecord(
                scope="\u8868\u7ea7",
                target="",
                raw_unit=raw,
                normalized_unit=_normalize_unit(raw),
                source_page=page_number,
                source_type="\u8868\u683c\u4e0a\u65b9\u6807\u6ce8",
                source_text=content,
                confidence="\u9ad8",
            ))
    # Some reports put units directly in the table title, for example
    # “S02-实际资本（元）” or “其他经营指标（元，%）”. Only inspect
    # title-side context through the first target line so later row/column
    # parentheses cannot be promoted to table-level units.
    for line in lines[context_start:first + 1]:
        content = _line_content(line).strip()
        records.extend(_unit_from_label(
            content,
            scope="\u8868\u7ea7",
            target="",
            source_page=page_number,
            source_type="\u8868\u683c\u6807\u9898\u62ec\u53f7\u5355\u4f4d",
        ))
    return records


def extract_unit_records(
    table_id: str,
    rows: list[list[str]],
    source_grids: list[PageGrid] | None = None,
) -> list[UnitRecord]:
    records: list[UnitRecord] = []
    for grid in source_grids or []:
        records.extend(_unit_records_from_grid(table_id, grid.page_number, grid.grid_text))

    first_grid = next(iter(source_grids or []), None)
    source_page_number = first_grid.page_number if first_grid else 0
    if rows:
        for cell in rows[0]:
            text = _clean(cell)
            records.extend(_unit_from_label(
                text,
                scope="\u5217\u7ea7",
                target=text,
                source_page=source_page_number,
                source_type="\u8868\u5934\u9644\u5e26\u5355\u4f4d",
            ))
        for row in rows[1:]:
            if not row:
                continue
            label_index = (
                1
                if len(row) > 1 and re.fullmatch(r"\d+(?:\.\d+)*\*?", row[0].strip())
                else 0
            )
            label = _clean(row[label_index] if label_index < len(row) else "")
            if label:
                records.extend(_unit_from_label(
                    label,
                    scope="\u884c\u7ea7",
                    target=label,
                    source_page=source_page_number,
                    source_type="\u9879\u76ee\u540d\u79f0\u9644\u5e26\u5355\u4f4d",
                ))
            value_cells = row[label_index + 1:]
            has_percent = any(
                re.search(r"(?:%|\uff05)\s*$", cell.strip())
                for cell in value_cells
                if cell.strip()
            )
            if label and has_percent:
                records.append(UnitRecord(
                    scope="\u884c\u7ea7",
                    target=label,
                    raw_unit="%",
                    normalized_unit="%",
                    source_page=source_page_number,
                    source_type="\u6570\u503c\u683c\u5f0f",
                    source_text="\uff1b".join(
                        cell.strip() for cell in value_cells if cell.strip()
                    ),
                    confidence="\u9ad8",
                ))
    return merge_unit_records(records)

def _group_matches(
    matches: list[PageMatch],
) -> list[tuple[str, str, list[int], str, dict]]:
    grouped: dict[str, dict] = {}
    order: list[str] = []
    for match in matches:
        if match.table_id not in grouped:
            grouped[match.table_id] = {
                "name": match.table_name,
                "pages": [],
                "strategy_id": match.strategy_id,
                "table_config": dict(match.table_config),
            }
            order.append(match.table_id)
        elif (
            match.strategy_id
            and grouped[match.table_id]["strategy_id"]
            and match.strategy_id != grouped[match.table_id]["strategy_id"]
        ):
            raise ValueError(
                f"目标表 {match.table_id} 同时引用了多个提取策略。"
            )
        elif match.strategy_id:
            grouped[match.table_id]["strategy_id"] = match.strategy_id
        grouped[match.table_id]["pages"].extend(match.pages)
    return [
        (
            table_id,
            grouped[table_id]["name"],
            sorted(set(grouped[table_id]["pages"])),
            grouped[table_id]["strategy_id"],
            grouped[table_id]["table_config"],
        )
        for table_id in order
    ]


def _words_to_grid(page, columns: int = GRID_COLUMNS) -> tuple[str, int]:
    words = page.extract_words(
        x_tolerance=2, y_tolerance=3, keep_blank_chars=False, use_text_flow=False
    )
    if not words:
        return "", 0
    groups: list[dict] = []
    for word in sorted(words, key=lambda item: (float(item["top"]), float(item["x0"]))):
        top = float(word["top"])
        if not groups or abs(top - groups[-1]["top"]) > 3.5:
            groups.append({"top": top, "words": [word]})
        else:
            group = groups[-1]
            group["words"].append(word)
            group["top"] = (group["top"] * (len(group["words"]) - 1) + top) / len(group["words"])

    width = max(float(page.width), 1.0)
    lines: list[str] = []
    for row_number, group in enumerate(groups):
        canvas = [" "] * columns
        for word in sorted(group["words"], key=lambda item: float(item["x0"])):
            text = _clean(word.get("text", ""))
            column = min(columns - 1, max(0, int(float(word["x0"]) / width * (columns - 1))))
            while column < columns and canvas[column] != " ":
                column += 1
            for character in text:
                if column >= columns:
                    break
                canvas[column] = character
                column += 1
        line = "".join(canvas).rstrip()
        if line:
            lines.append(f"{row_number:03d}|{line}")
    return "\n".join(lines), len(words)


def build_pdf_grids(pdf_bytes: bytes, pages: list[int]) -> list[PageGrid]:
    result: list[PageGrid] = []
    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        for page_number in pages:
            if page_number < 1 or page_number > len(pdf.pages):
                raise ValueError(f"页码 {page_number} 超出PDF范围。")
            grid, word_count = _words_to_grid(pdf.pages[page_number - 1])
            result.append(PageGrid(page_number, grid, word_count))
    return result


def _instructions(
    table_id: str,
    table_name: str,
    pages: list[int],
    table_config: Mapping[str, object] | None = None,
) -> str:
    config = dict(table_config or {})
    signatures = tuple(
        config.get("content_terms") or TABLE_SIGNATURES.get(table_id, ())
    )
    headers = tuple(
        config.get("header_terms") or TABLE_HEADERS.get(table_id, ())
    )
    configured_exclusions = tuple(dict.fromkeys([
        *config.get("exclude_item_terms", ()),
        *config.get("profile_exclude_items", ()),
        *config.get("stop_terms", ()),
    ]))
    exclusions = configured_exclusions or TABLE_EXCLUSIONS.get(table_id, ())
    completeness = tuple(
        config.get("completeness_terms")
        or TABLE_COMPLETENESS_TERMS.get(table_id, ())
    )
    strategy = active_table_strategy(table_id)
    return strategy.build_prompt(PromptRequest(
        mode="full_table",
        table_id=table_id,
        table_name=table_name,
        signatures=signatures,
        headers=headers,
        exclusions=exclusions,
        completeness_terms=completeness,
        boundary_text=boundary_instruction_for_table(config or {"table_id": table_id}),
        pages=tuple(pages),
        table_config=config,
    ))


def _slice_grid_for_target_core(
    table_id: str,
    grid_text: str,
    *,
    operating_metrics: bool = False,
    respect_variant_activation: bool = False,
    require_value_for_boundary_items: bool = False,
    table_config: Mapping[str, object] | None = None,
) -> str:
    config = dict(table_config or {})
    lines = grid_text.splitlines()
    if not lines:
        return grid_text

    compact_lines = [_compact(line.partition("|")[2]) for line in lines]

    def boundary_line_has_item(line: str, item: str) -> bool:
        matched = line_has_item(
            line,
            item,
            require_value=require_value_for_boundary_items,
        )
        if matched or not require_value_for_boundary_items:
            return matched
        # ``compact_item`` intentionally removes punctuation, so a business
        # row whose disclosed values are only "-" placeholders cannot satisfy
        # ``require_value=True`` there. Preserve that legitimate row while
        # still rejecting a bare numbered section title such as
        # "十一、最低资本".
        if not line_has_item(line, item):
            return False
        content = str(line or "").partition("|")[2]
        return bool(re.search(
            r"(?:--|—|-|[<＜]不适用[>＞]|不适用|N/?A)\s*$",
            content,
            flags=re.I,
        ))

    variants = boundary_variants(table_id, table_config)
    boundary = variants[0] if variants else {}
    variant_profiles: list[tuple[int, int, dict, int | None, bool]] = []
    for variant_index, candidate in enumerate(variants):
        activation_terms = tuple(candidate.get("activation_terms", ()))
        if (
            respect_variant_activation
            and activation_terms
            and not any(
                any(_compact(term) in line for term in activation_terms)
                for line in compact_lines
            )
        ):
            continue
        candidate_starts = tuple(candidate.get("start_items", ()))
        candidate_end_groups = end_item_groups(table_id, candidate)
        candidate_ends = tuple(
            item for group in candidate_end_groups for item in group
        )
        start_hits = [
            index
            for index, line in enumerate(lines)
            if any(boundary_line_has_item(line, item) for item in candidate_starts)
        ]
        end_hits = [
            index
            for index, line in enumerate(lines)
            if any(boundary_line_has_item(line, item) for item in candidate_ends)
        ]
        preferred_end_items = (
            candidate_end_groups[0]
            if candidate_end_groups
            else ()
        )
        preferred_end_hits = [
            index
            for index, line in enumerate(lines)
            if any(
                boundary_line_has_item(line, item)
                for item in preferred_end_items
            )
        ]
        ordered = bool(
            start_hits
            and end_hits
            and any(end >= start for start in start_hits for end in end_hits)
        )
        end_only_continuation = bool(
            operating_metrics
            and preferred_end_hits
            and (
                not start_hits
                or min(preferred_end_hits) < min(start_hits)
            )
        )
        if end_only_continuation:
            # A continuation tail can end before a footnote later repeats the
            # table's start item. Treat it as end-only evidence so the footnote
            # cannot replace the real page start.
            evidence_rank = 3
        elif ordered:
            evidence_rank = 4
        elif start_hits and end_hits:
            evidence_rank = 1
        elif end_hits:
            evidence_rank = 3
        elif start_hits:
            evidence_rank = 2
        else:
            evidence_rank = 0
        variant_profiles.append((
            evidence_rank,
            -variant_index,
            candidate,
            (
                min(start_hits)
                if start_hits and not end_only_continuation
                else None
            ),
            end_only_continuation,
        ))
    if variant_profiles:
        _, _, boundary, profiled_start_hit, end_only_continuation = max(
            variant_profiles,
            key=lambda item: (item[0], item[1]),
        )
    else:
        profiled_start_hit = None
        end_only_continuation = False

    start_items = tuple(boundary.get("start_items", ()))
    configured_end_groups = end_item_groups(table_id, boundary)
    title_terms = tuple(boundary.get("title_terms", ()))

    title_index = next(
        (
            index
            for index, line in enumerate(compact_lines)
            if any(_compact(term) in line for term in title_terms)
        ),
        None,
    )
    search_from = title_index + 1 if title_index is not None else 0

    start_hit = (
        profiled_start_hit
        if profiled_start_hit is not None and profiled_start_hit >= search_from
        else None
    )
    if start_hit is None and not end_only_continuation:
        for item in start_items:
            start_hit = next(
                (
                    index
                    for index, line in enumerate(
                        lines[search_from:],
                        start=search_from,
                    )
                    if boundary_line_has_item(line, item)
                ),
                None,
            )
            if start_hit is not None:
                break

    if (
        operating_metrics
        and title_index is not None
        and start_hit is not None
    ):
        # Keep the report's real table header. Some operating-metrics pages
        # contain a category row immediately before the first business item;
        # using only ``start_hit - 1`` would drop both the title and columns.
        start = title_index
    elif start_hit is not None:
        lower_bound = title_index if title_index is not None else 0
        start = max(lower_bound, start_hit - 1)
    else:
        start = title_index if title_index is not None else 0

    if (
        operating_metrics
        and start_hit is None
        and title_index is None
    ):
        continuation_titles = tuple(map(
            _compact,
            ("效益类指标", "规模类指标", "品质类指标"),
        ))
        continuation_start = next(
            (
                index
                for index, line in enumerate(compact_lines)
                if any(title in line for title in continuation_titles)
            ),
            None,
        )
        product_detail_markers = tuple(map(
            _compact,
            (
                "前五大产品的信息",
                "第一大产品的信息",
                "第二大产品的信息",
                "第三大产品的信息",
                "第四大产品的信息",
                "第五大产品的信息",
            ),
        ))
        product_details_before_continuation = (
            continuation_start is not None
            and any(
                any(marker in line for marker in product_detail_markers)
                for line in compact_lines[:continuation_start]
            )
        )
        if continuation_start is not None and product_details_before_continuation:
            start = continuation_start

    end_search_from = start_hit + 1 if start_hit is not None else start
    end_hit = None
    candidate_end_groups = (
        configured_end_groups
        if operating_metrics
        else configured_end_groups[:1]
    )
    for group_index, end_items in enumerate(candidate_end_groups):
        candidate_end_hit = next(
            (
                index
                for index, line in enumerate(
                    lines[end_search_from:], start=end_search_from
                )
                if any(boundary_line_has_item(line, item) for item in end_items)
            ),
            None,
        )
        if candidate_end_hit is not None:
            if (
                operating_metrics
                and group_index > 0
                and any(
                    any(
                        line_has_item(lines[index], title)
                        for title in (
                            config.get("section_titles")
                            or OPERATING_SECTION_TITLES
                        )
                    )
                    for index in range(candidate_end_hit + 1, len(lines))
                )
            ):
                # “综合投资收益率”只是未披露后续类别时的备用终点。
                # 同页已出现效益/规模/品质类标题，说明目标表仍在继续。
                continue
            end_hit = candidate_end_hit
            break
    if end_hit is not None:
        end = end_hit + 1
    else:
        if start_hit is None and title_index is None:
            for index, line in enumerate(compact_lines):
                if any(
                    _compact(term) in line
                    for term in (
                        config.get("start_item_terms")
                        or config.get("title_terms")
                        or TABLE_START_MARKERS.get(table_id, ())
                    )
                ):
                    start = index
                    break
        end = len(lines)
        for index in range(start + 1, len(lines)):
            if any(
                _compact(term) in compact_lines[index]
                for term in (
                    config.get("stop_terms")
                    or AI_TABLE_END_MARKERS.get(table_id, ())
                )
            ):
                end = index
                break
    return "\n".join(lines[start:end])


def _slice_grid_for_target(
    table_id: str,
    grid_text: str,
    table_config: Mapping[str, object] | None = None,
) -> str:
    strategy = active_table_strategy(table_id)
    return strategy.enforce_boundaries(BoundaryRequest(
        mode="grid_slice",
        table_id=table_id,
        grid_text=grid_text,
        engine=_slice_grid_for_target_core,
        table_config=dict(table_config or {}),
    ))


_TABLE_FOOTNOTE_START_RE = re.compile(
    r"^\s*(?:\d+\s*)?(?:注释|注|(?:填表)?说明)\s*[：:]"
)


def _source_table_lines_core(
    table_id: str,
    grids: list[PageGrid] | None,
    *,
    preserve_adjustment_after_footnote: bool = False,
    table_config: Mapping[str, object] | None = None,
) -> list[str]:
    """Return target-table lines without page-local footnotes and their continuations."""
    source_lines: list[str] = []
    for grid in grids or []:
        sliced_lines = _slice_grid_for_target(
            table_id,
            grid.grid_text,
            table_config,
        ).splitlines()
        index = 0
        while index < len(sliced_lines):
            line = sliced_lines[index]
            content = line.partition("|")[2]
            if _TABLE_FOOTNOTE_START_RE.match(content):
                if preserve_adjustment_after_footnote:
                    continuation_index = next(
                        (
                            candidate
                            for candidate in range(index + 1, len(sliced_lines))
                            if any(
                                term in _compact(
                                    sliced_lines[candidate].partition("|")[2]
                                )
                                for term in (
                                    _compact("核心一级资本调整表"),
                                    _compact("核心资本调整表"),
                                )
                            )
                        ),
                        None,
                    )
                    if continuation_index is not None:
                        index = continuation_index
                        continue
                break
            source_lines.append(line)
            index += 1
    return source_lines


def _source_table_lines(
    table_id: str,
    grids: list[PageGrid] | None,
    table_config: Mapping[str, object] | None = None,
) -> list[str]:
    strategy = active_table_strategy(table_id)
    return strategy.enforce_boundaries(BoundaryRequest(
        mode="source_lines",
        table_id=table_id,
        source_grids=grids,
        engine=_source_table_lines_core,
        table_config=dict(table_config or {}),
    ))


_GRID_VALUE_RE = re.compile(
    r"(?<![\d.])(?:[<＜]\s*不适用\s*[>＞]|不适用|N\s*/?\s*A|"
    r"--|—|-|[（(]?-?\d[\d,，]*(?:\.\d+)?%?[）)]?)(?![\d.])",
    flags=re.I,
)
_SOURCE_CELL_VALUE_RE = re.compile(
    r"(?:[<＜]\s*不适用\s*[>＞]|不适用|N\s*/?\s*A|"
    r"(?<!\d)[(（\-]?\d[\d,，.％%]*[)）]?)",
    flags=re.I,
)
_SOURCE_PLACEHOLDER_VALUE_RE = re.compile(
    r"(?<!\S)(?:--|—|-)(?!\S)"
)


def _source_value_matches(
    table_id: str,
    content: str,
) -> list[re.Match]:
    matches = list(_SOURCE_CELL_VALUE_RE.finditer(content))
    if table_id == "OPERATING_METRICS":
        matches.extend(_SOURCE_PLACEHOLDER_VALUE_RE.finditer(content))
    return sorted(matches, key=lambda match: match.start())


def _is_operating_section_only_line(
    content: str,
    value_matches: list[re.Match],
) -> bool:
    if not value_matches:
        return False
    label = _compact(content[:value_matches[0].start()])
    label = re.sub(
        r"^(?:[一二三四五六七八九十]+|\d+(?:\.\d+)*)",
        "",
        label,
    )
    non_business_titles = (
        *OPERATING_SECTION_TITLES,
        "主要经营指标",
        "主要经营情况",
    )
    return label in {_compact(item) for item in non_business_titles}


def _cell_has_disclosed_value(value: object) -> bool:
    text = str(value or "")
    return bool(
        re.search(r"\d", text)
        or re.fullmatch(
            r"\s*(?:[<＜]\s*不适用\s*[>＞]|不适用|N\s*/?\s*A)\s*",
            text,
            flags=re.I,
        )
    )


def _grid_business_row(line: str, width: int) -> list[str] | None:
    content = line.partition("|")[2].strip()
    value_matches = list(_GRID_VALUE_RE.finditer(content))
    if len(value_matches) < 3:
        return None
    current_match, previous_match = value_matches[-2:]
    prefix = content[:current_match.start()].strip()
    row_number_match = re.match(r"(?P<number>\d+(?:\.\d+)*\*?)\s+", prefix)
    row_number = row_number_match.group("number") if row_number_match else ""
    label = (
        prefix[row_number_match.end():].strip()
        if row_number_match
        else prefix
    )
    if not label:
        return None
    parsed = (
        [
            row_number,
            label,
            current_match.group(0),
            previous_match.group(0),
        ]
        if width >= 4
        else [label, current_match.group(0), previous_match.group(0)]
    )
    parsed += [""] * max(0, width - len(parsed))
    return parsed[:width]


def _normalize_minimum_capital_total_alias_rows(
    rows: list[list[str]],
) -> tuple[list[list[str]], bool]:
    """Normalize a row-numbered final “4 合计” to canonical minimum capital."""
    normalized: list[list[str]] = []
    changed = False
    for row in rows:
        cells = list(row)
        if (
            len(cells) >= 2
            and _compact(cells[0]) == "4"
            and _compact(cells[1]).startswith("合计")
        ):
            cells[1] = re.sub("合计", "最低资本", cells[1], count=1)
            changed = True
        elif cells and re.match(r"^4\s*合计", str(cells[0]).strip()):
            cells[0] = re.sub(
                r"^(4\s*)合计",
                r"\1最低资本",
                cells[0],
                count=1,
            )
            changed = True
        elif (
            len(cells) >= 3
            and _compact(cells[0]).startswith("合计")
        ):
            cells[0] = re.sub("合计", "最低资本", cells[0], count=1)
            changed = True
        normalized.append(cells)
    return normalized, changed


def recover_minimum_capital_terminal_rows_core(
    table_id: str,
    rows: list[list[str]],
    grids: list[PageGrid] | None,
) -> tuple[list[list[str]], str]:
    """Recover the explicitly printed tail after the main additional-capital row.

    This is deliberately limited to the quantity-risk-first minimum-capital
    layout: both source and model output must show the main “附加资本” row,
    and the source must show a later terminal “最低资本”. Therefore it cannot
    move the summary-first layout's opening minimum-capital row.
    """
    rows, alias_normalized = _normalize_minimum_capital_total_alias_rows(rows)
    alias_note = (
        "已将源表末行‘4 合计’标准化为‘4 最低资本’"
        if alias_normalized
        else ""
    )
    if not rows or not grids:
        return rows, alias_note
    terminal_aliases = (
        "最低资本",
        "最低资本合计",
        "最低资本总额",
        "合计",
    )

    source_lines = _source_table_lines(table_id, grids)
    terminal_candidates: list[tuple[int, str, str]] = []
    for index, line in enumerate(source_lines):
        alias = next(
            (
                item
                for item in terminal_aliases
                if line_has_item(line, item, require_value=True)
            ),
            "",
        )
        if alias:
            terminal_candidates.append((index, alias, line))
    if not terminal_candidates:
        return rows, alias_note

    terminal_index, alias, _ = terminal_candidates[-1]
    width = max(len(row) for row in rows)
    source_additional_indices = [
        index
        for index, line in enumerate(source_lines[:terminal_index])
        if (
            line_has_item(line, "附加资本")
            and _grid_business_row(line, width) is not None
        )
    ]
    output_additional_index = next(
        (
            index
            for index, row in enumerate(rows)
            if item_hits_in_rows([row], ("附加资本",))
        ),
        None,
    )
    if not source_additional_indices or output_additional_index is None:
        return rows, alias_note

    source_additional_index = source_additional_indices[-1]
    source_tail = [
        parsed
        for line in source_lines[source_additional_index + 1:terminal_index + 1]
        if (parsed := _grid_business_row(line, width)) is not None
    ]
    source_tail, _ = _normalize_minimum_capital_total_alias_rows(source_tail)
    canonical_terminal_aliases = (
        "最低资本",
        "最低资本合计",
        "最低资本总额",
    )
    if not source_tail or not item_hits_in_rows(
        source_tail[-1:], canonical_terminal_aliases
    ):
        return rows, alias_note

    existing_tail = rows[output_additional_index + 1:]
    if existing_tail == source_tail:
        return rows, alias_note
    recovered_rows = [*rows[:output_additional_index + 1], *source_tail]
    recovery_note = (
        f"模型遗漏最低资本尾段，已依据源PDF网格补回"
        f"{len(source_tail)}行（终止于“{alias}”）"
    )
    return recovered_rows, "；".join(
        note for note in (alias_note, recovery_note) if note
    )


def recover_minimum_capital_terminal_rows(
    table_id: str,
    rows: list[list[str]],
    grids: list[PageGrid] | None,
) -> tuple[list[list[str]], str]:
    """Backward-compatible table-aware entry point."""
    if table_id != "MINIMUM_CAPITAL":
        return rows, ""
    return recover_minimum_capital_terminal_rows_core(table_id, rows, grids)


def _source_completeness_profile_core(
    table_id: str,
    grids: list[PageGrid] | None,
    *,
    allow_single_value_boundary_rows: bool = False,
    table_config: Mapping[str, object] | None = None,
) -> tuple[int, tuple[str, ...]]:
    """Estimate source rows and the boundary items visibly present in the source."""
    config = dict(table_config or {})
    source_lines = _source_table_lines(table_id, grids, config)

    required_terms = [
        term
        for term in (
            config.get("completeness_terms")
            or TABLE_COMPLETENESS_TERMS.get(table_id, ())
        )
        if any(
            line_has_item(line, term, require_value=True)
            for line in source_lines
        )
    ]
    semantic_items = boundary_items(table_id, config)
    for item in semantic_items:
        if any(
            line_has_item(line, item, require_value=True)
            for line in source_lines
        ):
            canonical_item = item
            if item == "合计":
                if table_id == "MINIMUM_CAPITAL":
                    canonical_item = "最低资本"
                elif table_id == "ACTUAL_CAPITAL":
                    canonical_item = "实际资本合计"
                elif table_id == "RECOGNIZED_ASSETS":
                    canonical_item = "认可资产合计"
            required_terms.append(canonical_item)

    numeric_rows = 0
    for line in source_lines:
        content = line.partition("|")[2]
        compact_content = _compact(content)
        if any(
            compact_content.startswith(_compact(term))
            for term in (
                "公司名称",
                "报告期",
                "报表日期",
                "编制日期",
                "单位",
                "币种",
            )
        ):
            continue
        value_matches = _source_value_matches(table_id, content)
        if (
            bool(config.get("completeness_track_sections"))
            or table_id == "OPERATING_METRICS"
        ) and _is_operating_section_only_line(content, value_matches):
            continue
        if len(value_matches) >= 2:
            numeric_rows += 1
        elif allow_single_value_boundary_rows and (
            "%" in content or "％" in content
        ) and any(line_has_item(line, item) for item in semantic_items):
            numeric_rows += 1
    return numeric_rows, tuple(dict.fromkeys(required_terms))


def _source_completeness_profile(
    table_id: str,
    grids: list[PageGrid] | None,
    table_config: Mapping[str, object] | None = None,
) -> tuple[int, tuple[str, ...]]:
    strategy = active_table_strategy(table_id)
    return strategy.validate_completeness(CompletenessRequest(
        mode="source_profile",
        table_id=table_id,
        rows=[],
        validator=_source_completeness_profile_core,
        validator_kwargs={"grids": grids},
        table_config=dict(table_config or {}),
    ))


def _normalized_item_label(value: str) -> str:
    text = _compact(value)
    text = re.sub(r"^\d+(?:\.\d+)*\*?", "", text)
    text = re.sub(r"[（(](?:元|千元|万元|亿元|%|％|人|件|次)[）)]$", "", text)
    return text


def _source_item_labels(
    table_id: str,
    grids: list[PageGrid] | None,
    table_config: Mapping[str, object] | None = None,
) -> tuple[str, ...]:
    """Recover row-label fingerprints from the source PDF grid.

    Labels are deliberately conservative: a source line must contain a
    financial value and a meaningful text prefix. This avoids treating titles,
    page numbers, and narrative notes as required table rows.
    """
    labels: list[str] = []
    number_pattern = _SOURCE_CELL_VALUE_RE
    config = dict(table_config or {})
    header_terms = {
        _compact(term)
        for term in (
            config.get("header_terms")
            or TABLE_HEADERS.get(table_id, ())
        )
    }
    metadata_terms = tuple(map(
        _compact,
        ("公司名称", "报告期", "报表日期", "编制日期", "单位", "币种"),
    ))
    source_lines = _source_table_lines(table_id, grids, config)
    numbered_row_re = re.compile(
        r"^\s*(?:[（(]?[一二三四五六七八九十]+[）)]|"
        r"\d+(?:\.\d+)*\.?)"
    )
    for line_index, line in enumerate(source_lines):
        content = line.partition("|")[2]
        compact_content = _compact(content)
        if any(
            compact_content.startswith(term)
            for term in metadata_terms
        ):
            continue
        number_matches = _source_value_matches(table_id, content)
        if len(number_matches) < 2:
            continue
        if (
            (
                bool(config.get("completeness_track_sections"))
                or table_id == "OPERATING_METRICS"
            )
            and _is_operating_section_only_line(content, number_matches)
        ):
            continue
        # The item label ends at the first value column. When a row starts with
        # a lone row-number (\u884c\u6b21/\u5e8f\u53f7) cell, that index is the first numeric
        # token, so the first value is the following token; otherwise the first
        # token is already a value. Anchoring to the first value keeps
        # multi-value-column tables (e.g. RECOGNIZED_ASSETS with \u8d26\u9762\u4ef7\u503c/\u975e\u8ba4\u53ef
        # \u4ef7\u503c/\u8ba4\u53ef\u4ef7\u503c \u00d7 \u671f\u672b/\u671f\u521d) from gluing value cells into the label.
        first_prefix = content[: number_matches[0].start()].strip()
        is_leading_row_number = (
            not first_prefix
            and re.match(
                r"^\s*\d+(?:\.\d+)*\*?\.?\s*[一-鿿A-Za-z]",
                content,
            )
        )
        value_index = 1 if is_leading_row_number else 0
        number_match = number_matches[value_index]
        label_prefix = content[:number_match.start()].strip()
        recovered_wrapped_label = False
        if (
            not numbered_row_re.match(label_prefix)
            and not re.search(r"[\u4e00-\u9fffA-Za-z]", label_prefix)
        ):
            fragments: list[str] = []
            for previous_index in range(
                line_index - 1,
                max(-1, line_index - 3),
                -1,
            ):
                previous = source_lines[previous_index].partition("|")[2].strip()
                if not previous:
                    continue
                fragments.insert(0, previous)
                if numbered_row_re.match(previous):
                    label_prefix = "".join([*fragments, label_prefix])
                    recovered_wrapped_label = True
                    break
                if len(number_pattern.findall(previous)) >= 2:
                    break
        if recovered_wrapped_label:
            for next_index in range(
                line_index + 1,
                min(len(source_lines), line_index + 3),
            ):
                following = source_lines[next_index].partition("|")[2].strip()
                if (
                    number_pattern.search(following)
                    or numbered_row_re.match(following)
                ):
                    break
                if following:
                    label_prefix += following
        label = _normalized_item_label(label_prefix)
        if len(label) < 2 or not re.search(r"[\u4e00-\u9fffA-Za-z]", label):
            continue
        if label in header_terms or any(
            header and header in label and len(label) <= len(header) + 2
            for header in header_terms
        ):
            continue
        if any(label.startswith(term) for term in metadata_terms):
            continue
        labels.append(label)
    return tuple(dict.fromkeys(labels))


def _source_section_titles_core(
    table_id: str,
    grids: list[PageGrid] | None,
    *,
    track_operating_sections: bool = False,
    table_config: Mapping[str, object] | None = None,
) -> tuple[str, ...]:
    if not track_operating_sections:
        return ()
    ordered: list[str] = []
    config = dict(table_config or {})
    section_titles = tuple(
        config.get("section_titles") or OPERATING_SECTION_TITLES
    )
    for line in _source_table_lines(table_id, grids, config):
        compact_line = _compact(line.partition("|")[2])
        for title in section_titles:
            if _compact(title) in compact_line and title not in ordered:
                ordered.append(title)
    return tuple(ordered)


def _source_section_titles(
    table_id: str,
    grids: list[PageGrid] | None,
    table_config: Mapping[str, object] | None = None,
) -> tuple[str, ...]:
    strategy = active_table_strategy(table_id)
    return strategy.validate_completeness(CompletenessRequest(
        mode="source_sections",
        table_id=table_id,
        rows=[],
        validator=_source_section_titles_core,
        validator_kwargs={"grids": grids},
        table_config=dict(table_config or {}),
    ))


def _validate_source_section_order(
    rows: list[list[str]],
    source_section_titles: tuple[str, ...],
) -> None:
    if not source_section_titles:
        return
    output_order: list[str] = []
    for row in rows:
        row_text = _compact("".join(row))
        for title in source_section_titles:
            if _compact(title) in row_text and title not in output_order:
                output_order.append(title)
    missing = [
        title for title in source_section_titles if title not in output_order
    ]
    if missing:
        raise ExtractionQualityError(
            f"遗漏源PDF类别标题：{'、'.join(missing)}。"
        )
    if tuple(output_order) != source_section_titles:
        raise ExtractionQualityError(
            "类别顺序与源PDF不一致：原文为"
            f"{' → '.join(source_section_titles)}，输出为"
            f"{' → '.join(output_order)}。"
        )


def _output_item_labels(rows: list[list[str]]) -> tuple[str, ...]:
    labels: list[str] = []
    for row in rows[1:] if len(rows) > 1 else rows:
        cells = [_clean(cell) for cell in row]
        label_index = _row_label_index(cells)
        if label_index >= len(cells):
            continue
        label = _normalized_item_label(cells[label_index])
        if len(label) >= 2 and re.search(r"[\u4e00-\u9fffA-Za-z]", label):
            labels.append(label)
    return tuple(dict.fromkeys(labels))


def _item_recall_profile(
    rows: list[list[str]],
    source_labels: tuple[str, ...],
) -> tuple[float, tuple[str, ...]]:
    if not source_labels:
        return 1.0, ()
    output_labels = _output_item_labels(rows)
    missing = [
        source
        for source in source_labels
        if not any(
            source == output
            or source in output
            or output in source
            for output in output_labels
        )
    ]
    recall = (len(source_labels) - len(missing)) / len(source_labels)
    return recall, tuple(missing)


def _trim_adjacent_rows(
    table_id: str,
    rows: list[list[str]],
    table_config: Mapping[str, object] | None = None,
) -> tuple[list[list[str]], str]:
    config = dict(table_config or {})
    markers = tuple(
        config.get("stop_terms")
        or AI_TABLE_END_MARKERS.get(table_id, ())
    )
    if len(rows) < 2 or not markers:
        return rows, ""
    for index, row in enumerate(rows[1:], start=1):
        row_text = _compact("".join(row))
        hit = next((_compact(term) for term in markers if _compact(term) in row_text), "")
        if hit:
            trimmed = rows[:index]
            if len(trimmed) < 2:
                raise ExtractionQualityError(f"目标表边界出现在首个数据行：{hit}。")
            return trimmed, f"已在相邻表边界“{hit}”前自动截断"
    return rows, ""

def _grid_messages(
    table_id: str,
    table_name: str,
    pages: list[int],
    grids: list[PageGrid],
    table_config: Mapping[str, object] | None = None,
) -> list[dict]:
    blocks = "\n\n".join(
        f"===== PDF物理第 {item.page_number} 页 =====\n"
        f"{_slice_grid_for_target(table_id, item.grid_text, table_config)}"
        for item in grids
    )
    config = dict(table_config or {})
    prompt_role = str(
        config.get("prompt_role")
        or "保险公司偿付能力报告表格重构专家"
    )
    return [
        {"role": "system", "content": f"你是{prompt_role}。必须依据页面网格恢复表格，不能编造披露数据。"},
        {"role": "user", "content": f"{_instructions(table_id, table_name, pages, table_config)}\n\n{blocks}"},
    ]


def _grid_repair_messages(
    table_id: str,
    table_name: str,
    pages: list[int],
    grids: list[PageGrid],
    previous_error: str,
    table_config: Mapping[str, object] | None = None,
) -> list[dict]:
    messages = _grid_messages(table_id, table_name, pages, grids, table_config)
    messages[-1]["content"] += f"""

上一轮重构未通过本地质量检查，失败原因：{previous_error}
请重新检查列对齐并返回完整JSON。特别注意：
- “行次/序号”列只是行编号，不是金额列；不得把它判作数值错位。
- 实际资本明细表应保留“行次、项目、期末数、期初数”等完整列。
- 实际资本的首个业务行可能是“财务报表资产总额”“认可资产”或“核心一级资本”；
  若原文表头后直接从“核心一级资本”开始，应从该行提取到“实际资本合计”，
  不得要求或补造原文未披露的“认可资产”等汇总行。
- 若原文从“实际资本”开始，随后另有“核心一级资本调整表”，两张子表都必须保留，
  并在“银保监会规定的其他调整项目”或同义项目结束；后续“认可资产表”和“认可负债表”
  明细不属于本目标，不得提取。
- 若实际资本先披露“实际资本汇总”、后披露“核心资本”，必须同时保留两张子表，
  从汇总首行“财务报表资产总额”或“认可资产”连续提取到核心资本末行“实际资本合计”；
  汇总表的“本季度数/上季度可比数”与明细表的“期末数/期初数”按本期、上期一一对应，
  统一输出为“期末数/期初数”，不得把第二张子表表头当作数据行。
- 组合经营指标按原文实际披露范围提取：后续类别缺失时允许以“综合投资收益率”结束；若出现品质类指标，则必须继续保留至最后的“营销员脱落率”（或同义名称），不得补造缺失项目。
- 偿付能力充足率指标不得只返回核心/综合充足率两行，必须保留实际资本、最低资本等整张表的所有数据行。
"""
    return messages


def _is_unsupported_vision_error(error: Exception | str) -> bool:
    message = str(error).lower().replace("`", "")
    signals = (
        "unknown variant image_url",
        "expected text",
        "image input is not supported",
        "does not support image",
        "unsupported image",
        "unsupported content type",
    )
    return any(signal in message for signal in signals)

def _render_images(
    pdf_bytes: bytes,
    pages: list[int],
    *,
    zoom: float = 1.8,
    jpeg_quality: int = 84,
) -> list[tuple[int, str]]:
    document = fitz.open(stream=pdf_bytes, filetype="pdf")
    result: list[tuple[int, str]] = []
    try:
        for page_number in pages:
            if page_number < 1 or page_number > len(document):
                raise ValueError(f"页码 {page_number} 超出PDF范围。")
            pixmap = document.load_page(page_number - 1).get_pixmap(
                matrix=fitz.Matrix(zoom, zoom), alpha=False
            )
            try:
                image = pixmap.tobytes("jpeg", jpg_quality=jpeg_quality)
                mime = "image/jpeg"
            except (TypeError, ValueError):
                image = pixmap.tobytes("png")
                mime = "image/png"
            result.append((page_number, f"data:{mime};base64,{base64.b64encode(image).decode('ascii')}"))
    finally:
        document.close()
    return result


def _vision_messages(
    table_id: str,
    table_name: str,
    pages: list[int],
    images: list[tuple[int, str]],
    table_config: Mapping[str, object] | None = None,
) -> list[dict]:
    content: list[dict] = [{
        "type": "text",
        "text": _instructions(table_id, table_name, pages, table_config) + "\n以下图片按物理页码顺序排列，请直接读取图片网格。",
    }]
    for page_number, data_url in images:
        content.extend([
            {"type": "text", "text": f"PDF物理第 {page_number} 页"},
            {"type": "image_url", "image_url": {"url": data_url, "detail": "high"}},
        ])
    return [
        {"role": "system", "content": "你是偿付能力报告图片表格识别专家。严格按图片网格恢复跨页表格，不得编造数据。"},
        {"role": "user", "content": content},
    ]


def _completion_url(base_url: str) -> str:
    value = base_url.strip().rstrip("/")
    if not value:
        raise ValueError("模型接口地址不能为空。")
    return value if value.endswith("/chat/completions") else f"{value}/chat/completions"


def _response_error(response) -> str:
    try:
        error = response.json().get("error", {})
        if isinstance(error, dict):
            return _clean(error.get("message", ""))
    except Exception:
        pass
    return _clean(getattr(response, "text", ""))[:500]


def _call_model(
    *, api_key: str, base_url: str, model: str, messages: list[dict],
    timeout: int, post_func: Callable | None,
    source_grids: list[PageGrid] | None = None,
) -> str:
    post = post_func or requests.post
    payload = {
        "model": normalize_model_id(base_url, model),
        "messages": messages,
        "response_format": {"type": "json_object"},
    }
    payload.update(model_request_parameters(base_url, model))
    headers = {"Authorization": f"Bearer {api_key.strip()}", "Content-Type": "application/json"}
    response = post_json_with_retry(
        post,
        _completion_url(base_url),
        headers=headers,
        json=payload,
        timeout=timeout,
    )
    if not getattr(response, "ok", False) and int(getattr(response, "status_code", 0) or 0) in {400, 422}:
        payload.pop("response_format", None)
        response = post_json_with_retry(
            post,
            _completion_url(base_url),
            headers=headers,
            json=payload,
            timeout=timeout,
        )
    if not getattr(response, "ok", False):
        raise RuntimeError(f"模型接口调用失败：{_response_error(response) or 'HTTP错误'}")
    try:
        content = response.json()["choices"][0]["message"]["content"]
    except (KeyError, IndexError, TypeError) as exc:
        raise RuntimeError("模型接口返回结构中没有choices/message/content。") from exc
    if isinstance(content, list):
        content = "".join(
            str(part.get("text", "")) if isinstance(part, dict) else str(part) for part in content
        )
    return str(content)

def _parse_json(content: str) -> dict:
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", content.strip(), flags=re.I)
    try:
        payload = json.loads(text)
    except json.JSONDecodeError:
        start, end = text.find("{"), text.rfind("}")
        if start < 0 or end <= start:
            raise ExtractionQualityError("模型没有返回可解析的JSON对象。")
        try:
            payload = json.loads(text[start:end + 1])
        except json.JSONDecodeError as exc:
            raise ExtractionQualityError(f"模型JSON解析失败：{exc}") from exc
    if not isinstance(payload, dict):
        raise ExtractionQualityError("模型返回结果不是JSON对象。")
    for key in ("result", "table", "data"):
        if isinstance(payload.get(key), dict):
            payload = payload[key]
            break
    if isinstance(payload.get("tables"), list) and payload["tables"]:
        payload = payload["tables"][0]
    return payload


def recover_unit_records_from_images(
    *,
    table_id: str,
    table_name: str,
    pages: list[int],
    rows: list[list[str]],
    images: list[tuple[int, str]],
    api_key: str,
    base_url: str,
    model: str,
    timeout: int,
    post_func: Callable | None,
) -> list[UnitRecord]:
    payload = _parse_json(_call_model(
        api_key=api_key,
        base_url=base_url,
        model=model,
        messages=_unit_vision_messages(
            table_id,
            table_name,
            pages,
            rows,
            images,
        ),
        timeout=timeout,
        post_func=post_func,
    ))
    return unit_records_from_payload(
        payload,
        rows,
        pages,
        source_type="\u56fe\u7247\u5355\u4f4d\u4e13\u7528\u8865\u63d0\u53d6",
    )


def _header_similarity(row: list[str], columns: list[str]) -> float:
    width = max(len(row), len(columns))
    if not width:
        return 0.0
    left = row + [""] * (width - len(row))
    right = columns + [""] * (width - len(columns))
    matches = sum(bool(_compact(a)) and _compact(a) == _compact(b) for a, b in zip(left, right))
    return matches / width


def _coerce_table(payload: dict) -> tuple[list[list[str]], float, str]:
    columns = payload.get("columns") or payload.get("headers") or payload.get("header")
    raw_rows = payload.get("rows") or payload.get("records")
    if not isinstance(raw_rows, list):
        raise ExtractionQualityError("模型结果缺少rows数组。")
    if not isinstance(columns, list) or not columns:
        first_dict = next((row for row in raw_rows if isinstance(row, dict)), None)
        if first_dict is None:
            raise ExtractionQualityError("模型结果缺少columns数组。")
        columns = list(first_dict)
    columns = [_clean(value) or f"列{index + 1}" for index, value in enumerate(columns)]
    width = len(columns)
    normalized: list[list[str]] = []
    ragged = 0
    for raw_row in raw_rows:
        if isinstance(raw_row, dict) and isinstance(raw_row.get("cells"), list):
            cells = raw_row["cells"]
        elif isinstance(raw_row, dict):
            cells = [raw_row.get(column, "") for column in columns]
        elif isinstance(raw_row, list):
            cells = raw_row
        else:
            continue
        cleaned = [_clean(value) for value in cells]
        if not any(cleaned):
            continue
        ragged += len(cleaned) != width
        cleaned = (cleaned + [""] * width)[:width]
        if _header_similarity(cleaned, columns) >= 0.6:
            continue
        if normalized and cleaned == normalized[-1]:
            continue
        normalized.append(cleaned)
    if not normalized:
        raise ExtractionQualityError("模型结果没有有效数据行。")
    return [columns, *normalized], ragged / max(len(normalized), 1), _clean(payload.get("notes", ""))


def _validate(
    table_id: str,
    rows: list[list[str]],
    ragged_ratio: float,
    *,
    expected_source_rows: int = 0,
    source_required_terms: tuple[str, ...] = (),
    source_item_labels: tuple[str, ...] = (),
    source_section_titles: tuple[str, ...] = (),
    required_signature_hits: int | None = None,
    source_item_recall_ratio: float = SOURCE_ITEM_RECALL_RATIO,
    configured_required: tuple[str, ...] = (),
    configured_exclusions: tuple[str, ...] = (),
    configured_minimum_rows: int = 0,
    table_config: Mapping[str, object] | None = None,
) -> tuple[float, str]:
    config = dict(table_config or {})
    if len(rows) < 2 or len(rows[0]) < 2:
        raise ExtractionQualityError("重构结果少于2行或2列。")
    if ragged_ratio > 0.2:
        raise ExtractionQualityError(f"超过20%的数据行列数与表头不一致（{ragged_ratio:.0%}）。")
    flat = _compact("".join(cell for row in rows for cell in row))
    required = (
        configured_required
        or tuple(config.get("content_terms", ()))
        or TABLE_SIGNATURES.get(table_id, ())
    )
    required_hits = [term for term in required if _compact(term) in flat]
    minimum = (
        required_signature_hits
        if required_signature_hits is not None
        else (
            len(required)
            if table_id == "SOLVENCY_MAIN"
            else min(1, len(required))
        )
    )
    if required and len(required_hits) < minimum:
        raise ExtractionQualityError(f"缺少目标表核心内容：{'、'.join(required)}。")
    exclusions = (
        tuple(configured_exclusions)
        or tuple(config.get("stop_terms", ()))
        or TABLE_EXCLUSIONS.get(table_id, ())
    )
    exclusion_hits = [term for term in exclusions if _compact(term) in flat]
    if exclusion_hits:
        raise ExtractionQualityError(f"疑似混入相邻表内容：{'、'.join(exclusion_hits)}。")

    completeness_terms = (
        source_required_terms
        or tuple(config.get("completeness_terms", ()))
        or TABLE_COMPLETENESS_TERMS.get(table_id, ())
    )
    semantic_items = set(boundary_items(table_id, config))
    exact_hits = set(item_hits_in_rows(rows, semantic_items))
    missing_completeness = [
        term for term in completeness_terms
        if (
            term not in exact_hits
            if term in semantic_items
            else _compact(term) not in flat
        )
    ]
    if missing_completeness:
        raise ExtractionQualityError(
            f"整表提取不完整，缺少原页面项目：{'、'.join(missing_completeness)}。"
        )

    data_rows = rows[1:]
    header_text = _compact("".join(rows[0][:2]))
    row_number_table = any(term in header_text for term in ("行次", "序号", "编号"))
    numeric_rows = 0
    shifted_rows = 0
    for row in data_rows:
        value_start = 2 if row_number_table and len(row) > 2 else 1
        has_number = any(
            _cell_has_disclosed_value(cell)
            for cell in row[value_start:]
        )
        numeric_rows += bool(has_number)
        if not row_number_table:
            shifted_rows += bool(re.search(r"\d", row[0]) and not has_number)
    configured_minimum = (
        int(config.get("minimum_rows", 0) or 0)
        if "minimum_rows" in config
        else max(
            int(TABLE_MIN_NUMERIC_ROWS.get(table_id, 1)),
            int(configured_minimum_rows or 0),
        )
    )
    minimum_numeric_rows = (
        min(configured_minimum, expected_source_rows)
        if expected_source_rows > 0
        else configured_minimum
    )
    if numeric_rows < minimum_numeric_rows:
        raise ExtractionQualityError(
            f"整表提取不完整：仅识别到{numeric_rows}条数值行，至少应有{minimum_numeric_rows}条。"
        )
    if expected_source_rows >= 3:
        coverage_minimum = max(
            minimum_numeric_rows,
            min(expected_source_rows, math.ceil(expected_source_rows * SOURCE_ROW_COVERAGE_RATIO)),
        )
        if numeric_rows < coverage_minimum:
            raise ExtractionQualityError(
                f"整表覆盖不足：原PDF网格约有{expected_source_rows}条数值行，"
                f"模型仅返回{numeric_rows}条，至少应返回{coverage_minimum}条。"
            )
    item_recall, missing_items = _item_recall_profile(rows, source_item_labels)
    _validate_source_section_order(rows, source_section_titles)
    if len(source_item_labels) >= 4 and item_recall < source_item_recall_ratio:
        preview = "、".join(missing_items[:8])
        raise ExtractionQualityError(
            f"项目行召回率不足：{item_recall:.0%}，要求至少"
            f"{source_item_recall_ratio:.0%}；疑似遗漏：{preview}。"
        )
    if len(data_rows) >= 3 and numeric_rows / len(data_rows) < 0.25:
        raise ExtractionQualityError("有效数值行比例过低，疑似列错位。")
    if shifted_rows >= 2 and shifted_rows / len(data_rows) > 0.35:
        raise ExtractionQualityError("数值集中在首列，疑似网格列错位。")

    header_hits = [
        term
        for term in (
            config.get("header_terms")
            or TABLE_HEADERS.get(table_id, ())
        )
        if _compact(term) in flat
    ]
    score = min(99.0, 65 + len(required_hits) * 8 + len(header_hits) * 3 + min(numeric_rows, 15) * 0.6)
    evidence_parts = [*required_hits, *header_hits]
    if source_item_labels:
        evidence_parts.append(f"项目行召回率{item_recall:.0%}")
    if row_number_table:
        evidence_parts.append("首列已识别为行次，不参与数值错位判断")
    evidence = "、".join(dict.fromkeys(evidence_parts))
    return round(score, 2), evidence or "跨页结构与数值列校验通过"


def _reconstruct(
    *, table_id: str, table_name: str, pages: list[int], mode: str,
    messages: list[dict], api_key: str, base_url: str, model: str,
    timeout: int, post_func: Callable | None,
    source_grids: list[PageGrid] | None = None,
    table_config: Mapping[str, object] | None = None,
) -> ExtractedTable:
    config = dict(table_config or {})
    strategy = active_table_strategy(table_id)
    payload = _parse_json(_call_model(
        api_key=api_key, base_url=base_url, model=model, messages=messages,
        timeout=timeout, post_func=post_func,
    ))
    rows, ragged_ratio, notes = _coerce_table(payload)
    rows, postprocess_notes = strategy.postprocess(PostprocessRequest(
        table_id=table_id,
        rows=rows,
        source_grids=source_grids,
        normalize_three_year=normalize_three_year_return_rows_core,
        recover_minimum_capital=recover_minimum_capital_terminal_rows_core,
        trim_adjacent=lambda target_id, target_rows: _trim_adjacent_rows(
            target_id,
            target_rows,
            config,
        ),
        table_config=config,
    ))
    if any("归一化" in note for note in postprocess_notes):
        ragged_ratio = 0.0
    try:
        rows, item_boundary_note = strategy.enforce_boundaries(BoundaryRequest(
            mode="output_rows",
            table_id=table_id,
            rows=rows,
            require_complete=True,
            engine=enforce_output_boundaries,
            table_config=config,
        ))
    except TableBoundaryError as exc:
        raise ExtractionQualityError(str(exc)) from exc
    boundary_note = "；".join(
        item
        for item in (
            *postprocess_notes,
            item_boundary_note,
        )
        if item
    )
    expected_source_rows, source_required_terms = _source_completeness_profile(
        table_id, source_grids, config
    )
    configured_completeness = tuple(config.get("completeness_terms", ()))
    if configured_completeness:
        source_required_terms = configured_completeness
    source_item_labels = _source_item_labels(table_id, source_grids, config)
    source_section_titles = _source_section_titles(table_id, source_grids, config)
    score, evidence = strategy.validate_completeness(CompletenessRequest(
        mode="full_table",
        table_id=table_id,
        rows=rows,
        ragged_ratio=ragged_ratio,
        signatures=tuple(
            config.get("content_terms")
            or TABLE_SIGNATURES.get(table_id, ())
        ),
        validator=_validate,
        validator_kwargs={
            "expected_source_rows": expected_source_rows,
            "source_required_terms": source_required_terms,
            "source_item_labels": source_item_labels,
            "source_section_titles": source_section_titles,
            "configured_required": tuple(config.get("content_terms", ())),
            "configured_exclusions": tuple(dict.fromkeys([
                *config.get("exclude_item_terms", ()),
                *config.get("profile_exclude_items", ()),
                *config.get("stop_terms", ()),
            ])),
            "configured_minimum_rows": int(config.get("minimum_rows", 0) or 0),
        },
        table_config=config,
    ))
    if boundary_note:
        evidence = f"{evidence}；{boundary_note}"
    if notes:
        evidence = f"{evidence}；{notes}"
    payload_units = unit_records_from_payload(
        payload,
        rows,
        pages,
        source_type="\u8868\u683c\u7ed3\u6784\u5316\u540c\u6b65\u5355\u4f4d",
    )
    return ExtractedTable(
        table_id=table_id, table_name=table_name, page=pages[0], table_index=1,
        rows=rows, strategy=mode, quality_score=score, evidence=evidence,
        source_pages=pages,
        unit_records=merge_unit_records(
            extract_unit_records(table_id, rows, source_grids),
            payload_units,
        ),
    )


def extract_tables_with_llm(
    pdf_bytes: bytes,
    matches: list[PageMatch],
    *,
    api_key: str,
    base_url: str,
    model: str,
    auto_vision_retry: bool = True,
    force_vision: bool = False,
    timeout: int = 180,
    post_func: Callable | None = None,
    progress_callback: Callable[[ExtractionLog], None] | None = None,
    _grid_builder: Callable[[bytes, list[int]], list[PageGrid]] | None = None,
    _render_func: Callable[..., list[tuple[int, str]]] | None = None,
) -> AIExtractionBundle:
    if not pdf_bytes:
        raise ValueError("请先上传PDF。")
    if not matches:
        raise ValueError("请先完成STEP1页码定位。")
    if not api_key.strip() or not base_url.strip() or not model.strip():
        raise ValueError("请完整填写模型接口地址、模型名称和API Key。")

    tables: list[ExtractedTable] = []
    logs: list[ExtractionLog] = []
    build_grids = _grid_builder or build_pdf_grids
    render_images = _render_func or _render_images

    def add_log(log: ExtractionLog) -> None:
        logs.append(log)
        if progress_callback:
            progress_callback(log)

    for table_id, table_name, pages, strategy_id, table_config in _group_matches(matches):
        table_strategy = resolve_table_strategy(
            table_id,
            strategy_id or None,
            table_config,
        )
        grids: list[PageGrid] = []
        text_error = ""
        if not force_vision:
            try:
                grids = build_grids(pdf_bytes, pages)
                if sum(item.word_count for item in grids) < max(12, len(pages) * 6):
                    raise ExtractionQualityError("页面可检索文字不足，判断为扫描件或文字层异常。")
                table = table_strategy.run(
                    "reconstruct_table",
                    _reconstruct,
                    table_name=table_name,
                    pages=pages,
                    mode=TEXT_MODE,
                    messages=_grid_messages(
                        table_id, table_name, pages, grids, table_config,
                    ),
                    source_grids=grids,
                    table_config=table_config,
                    api_key=api_key, base_url=base_url, model=model,
                    timeout=timeout, post_func=post_func,
                )
                tables.append(table)
                add_log(ExtractionLog(
                    table_id, table_name, pages, TEXT_MODE, "成功",
                    f"已完成多页网格重构和自动拼接，质量评分 {table.quality_score:.1f}。",
                ))
                continue
            except Exception as exc:
                text_error = str(exc)
                add_log(ExtractionLog(
                    table_id, table_name, pages, TEXT_MODE,
                    "触发图片重试" if auto_vision_retry else "失败", text_error,
                ))

        if not (force_vision or auto_vision_retry):
            continue

        try:
            images = render_images(pdf_bytes, pages)
            table = table_strategy.run(
                "reconstruct_table",
                _reconstruct,
                table_name=table_name,
                pages=pages,
                mode=VISION_MODE,
                messages=_vision_messages(
                    table_id, table_name, pages, images, table_config,
                ),
                source_grids=grids,
                table_config=table_config,
                api_key=api_key, base_url=base_url, model=model,
                timeout=timeout, post_func=post_func,
            )
            unit_note = ""
            if unit_recovery_needed(table.rows, table.unit_records):
                try:
                    recovered_units = recover_unit_records_from_images(
                        table_id=table_id,
                        table_name=table_name,
                        pages=pages,
                        rows=table.rows,
                        images=images,
                        api_key=api_key,
                        base_url=base_url,
                        model=model,
                        timeout=timeout,
                        post_func=post_func,
                    )
                    table.unit_records = merge_unit_records(
                        table.unit_records,
                        recovered_units,
                    )
                    if recovered_units:
                        unit_note = f"；单位专用补提取获得{len(recovered_units)}条明确单位"
                    else:
                        unit_note = "；单位专用补提取未发现明确单位，已保留待核对状态"
                except Exception as unit_exc:
                    unit_note = f"；单位专用补提取失败，已保留待核对状态：{unit_exc}"
            tables.append(table)
            add_log(ExtractionLog(
                table_id, table_name, pages, VISION_MODE, "成功",
                f"图片扫描重试成功，质量评分 {table.quality_score:.1f}{unit_note}。",
            ))
            continue
        except Exception as vision_exc:
            vision_error = str(vision_exc)
            vision_unsupported = _is_unsupported_vision_error(vision_exc)
            if vision_unsupported:
                add_log(ExtractionLog(
                    table_id, table_name, pages, VISION_MODE, "接口不支持图片",
                    "当前模型接口仅接受文本消息，已自动改用加强版网格纠错重试。",
                ))
            else:
                prefix = f"网格模式失败：{text_error}；" if text_error else ""
                add_log(ExtractionLog(
                    table_id, table_name, pages, VISION_MODE, "失败，转3倍高清图片",
                    f"{prefix}图片扫描重试未通过：{vision_error}",
                ))

            if not vision_unsupported:
                try:
                    high_res_images = render_images(
                        pdf_bytes,
                        pages,
                        zoom=3.0,
                        jpeg_quality=92,
                    )
                    table = table_strategy.run(
                        "reconstruct_table",
                        _reconstruct,
                        table_name=table_name,
                        pages=pages,
                        mode=VISION_HIGH_RES_MODE,
                        messages=_vision_messages(
                            table_id, table_name, pages, high_res_images,
                            table_config,
                        ),
                        source_grids=grids,
                        table_config=table_config,
                        api_key=api_key,
                        base_url=base_url,
                        model=model,
                        timeout=timeout,
                        post_func=post_func,
                    )
                    tables.append(table)
                    add_log(ExtractionLog(
                        table_id,
                        table_name,
                        pages,
                        VISION_HIGH_RES_MODE,
                        "成功",
                        f"3倍高清图片重试成功，质量评分{table.quality_score:.1f}。",
                    ))
                    continue
                except Exception as high_res_exc:
                    vision_error = (
                        f"{vision_error}；3倍高清图片重试未通过：{high_res_exc}"
                    )
                    add_log(ExtractionLog(
                        table_id,
                        table_name,
                        pages,
                        VISION_HIGH_RES_MODE,
                        "失败，转网格纠错",
                        str(high_res_exc),
                    ))

            if not grids:
                try:
                    grids = build_grids(pdf_bytes, pages)
                except Exception:
                    grids = []
            searchable_words = sum(item.word_count for item in grids)
            if searchable_words < max(12, len(pages) * 6):
                add_log(ExtractionLog(
                    table_id, table_name, pages, TEXT_RETRY_MODE, "失败",
                    "PDF没有足够文字层，图片重试也未成功。请改用支持图片的视觉模型，或人工核对页码后重试。",
                ))
                continue

            try:
                repair_reason = "；".join(item for item in (text_error, vision_error) if item)
                table = table_strategy.run(
                    "reconstruct_table",
                    _reconstruct,
                    table_name=table_name,
                    pages=pages,
                    mode=TEXT_RETRY_MODE,
                    messages=_grid_repair_messages(
                        table_id, table_name, pages, grids, repair_reason,
                        table_config,
                    ),
                    source_grids=grids,
                    table_config=table_config,
                    api_key=api_key, base_url=base_url, model=model,
                    timeout=timeout, post_func=post_func,
                )
                tables.append(table)
                add_log(ExtractionLog(
                    table_id, table_name, pages, TEXT_RETRY_MODE, "成功",
                    f"加强版网格纠错成功，质量评分 {table.quality_score:.1f}。",
                ))
            except Exception as repair_exc:
                add_log(ExtractionLog(
                    table_id, table_name, pages, TEXT_RETRY_MODE, "失败",
                    f"图片重试和网格纠错均未通过完整性检查：{repair_exc}",
                ))
    return AIExtractionBundle(tables=tables, logs=logs)

def extraction_logs_frame(logs: list[ExtractionLog]) -> pd.DataFrame:
    columns = ["目标表", "物理页码", "识别方式", "状态", "说明"]
    return pd.DataFrame([item.to_dict() for item in logs], columns=columns)


def reconstructed_workbook_bytes(
    bundle: AIExtractionBundle,
    metadata: Mapping[str, object] | None = None,
) -> bytes:
    output = io.BytesIO()
    used: set[str] = {"提取日志"}
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        extraction_logs_frame(bundle.logs).to_excel(writer, sheet_name="提取日志", index=False)
        metadata_rows = [
            {"字段": key, "值": value}
            for key in ("公司", "报告年度", "报告季度", "报告期", "披露日期", "来源文件")
            if (value := (metadata or {}).get(key)) not in (None, "")
        ]
        if metadata_rows:
            metadata_sheet = "报告元信息"
            pd.DataFrame(metadata_rows).to_excel(
                writer,
                sheet_name=metadata_sheet,
                index=False,
            )
            used.add(metadata_sheet)
        unit_frames = [
            table.units_frame() for table in bundle.tables if table.unit_records
        ]
        if unit_frames:
            unit_sheet = "\u5355\u4f4d\u4fe1\u606f"
            pd.concat(unit_frames, ignore_index=True).to_excel(writer, sheet_name=unit_sheet, index=False)
            used.add(unit_sheet)
        for table in bundle.tables:
            pages = table.source_pages or [table.page]
            page_label = str(pages[0]) if len(pages) == 1 else f"{pages[0]}-{pages[-1]}"
            base = re.sub(r"[\\/*?:\[\]]", "", f"{table.table_name}_P{page_label}")[:31] or "Table"
            name, suffix = base, 1
            while name in used:
                suffix += 1
                name = f"{base[:27]}_{suffix}"
            used.add(name)
            table.to_frame(include_unit_footer=True).to_excel(writer, sheet_name=name, index=False, header=False)

        fill = PatternFill("solid", fgColor="00338D")
        font = Font(color="FFFFFF", bold=True)
        unit_fill = PatternFill("solid", fgColor="FFF2CC")
        unit_footer_label = "\u3010\u5355\u4f4d\u5907\u6ce8\u3011"
        for sheet in writer.book.worksheets:
            sheet.freeze_panes = "A2"
            for cell in sheet[1]:
                cell.fill, cell.font = fill, font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            for row in sheet.iter_rows():
                if row and row[0].value == unit_footer_label:
                    for cell in row:
                        cell.fill = unit_fill
                        cell.alignment = Alignment(vertical="top", wrap_text=True)
            for index, cells in enumerate(sheet.iter_cols(1, sheet.max_column), start=1):
                length = max((len(str(cell.value or "")) for cell in cells), default=8)
                sheet.column_dimensions[get_column_letter(index)].width = min(max(length + 2, 10), 42)
    return output.getvalue()
