from __future__ import annotations

import io
import math
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font

from .solvency_metric_registry import (
    DERIVED_METRICS,
    INDUSTRY_METRICS_BY_SOURCE_CODE,
)
from .solvency_normalizer import (
    NARROW_TABLE_COLUMNS,
    STANDARD_COLUMNS,
    narrow_table_view,
    parse_numeric,
    standardize_uploaded_frame,
)
from .solvency_company_identity import (
    apply_company_identities,
    known_company_standard_name,
    reconcile_known_company_aliases,
    resolve_company_identity,
)


QUARTER_SHEET_PATTERN = re.compile(r"^(?P<year>20\d{2})(?P<quarter>Q[1-4])$", re.IGNORECASE)
REQUIRED_STANDARD_HEADERS = {"公司", "指标编码", "指标名称", "数值"}
INTERNAL_STANDARD_SHEET_NAME = "_系统字段"
FORMULA_SOURCE_SHEET_NAME = "计算依据"
IDENTIFIER_COLUMNS = {"分类", "公司"}
PERCENT_FROM_RATIO_CODES = {"CORE_SOLVENCY_RATIO", "COMBINED_SOLVENCY_RATIO"}
POLICY_SURPLUS_CORE_CODES = ("POLICY_SURPLUS_CORE_T1", "POLICY_SURPLUS_CORE_T2")
POLICY_SURPLUS_COMPONENT_CODES = (
    *POLICY_SURPLUS_CORE_CODES, "POLICY_SURPLUS_ANC_T1", "POLICY_SURPLUS_ANC_T2",
)
POLICY_SURPLUS_RATIO_COMPONENTS = {
    "POLICY_SURPLUS_CORE_TO_ACTUAL_CAPITAL": POLICY_SURPLUS_CORE_CODES,
    "POLICY_SURPLUS_CORE_TO_CORE_CAPITAL": POLICY_SURPLUS_CORE_CODES,
    "POLICY_SURPLUS_TO_INSURANCE_LIABILITIES": POLICY_SURPLUS_COMPONENT_CODES,
}
AUTHORITATIVE_POLICY_SURPLUS_DERIVED_CODES = {
    "POLICY_SURPLUS_CORE_TO_CORE_CAPITAL",
    "POLICY_SURPLUS_TO_INSURANCE_LIABILITIES",
}
EXTERNAL_INVESTMENT_COLUMNS = {
    "投资收益率(累计数)": ("投资收益率", "本年累计数"),
    "综合投资收益率(累计数)": ("综合投资收益率", "本年累计数"),
    "投资收益率(当季数)": ("投资收益率", "本季度数"),
    "综合投资收益率(当季数)": ("综合投资收益率", "本季度数"),
    "近三年平均投资收益率": ("投资收益率", "近三年平均"),
    "近三年平均综合投资收益率": ("综合投资收益率", "近三年平均"),
}
EXTERNAL_OPERATING_QUARTER_COLUMNS = {
    "综合退保率": ("综合退保率", True),
    "签单保费": ("签单保费", False),
    "新业务利润率": ("新业务利润率", True),
    "新业务价值": ("新业务价值", False),
}
EXTERNAL_CORE_T1_DETAIL_CODES = (
    "NON_RECOGNIZED_ASSET_BOOK_VALUE",
    "LONG_TERM_EQUITY_VALUATION_DIFFERENCE",
    "CORE_T1_INVESTMENT_PROPERTY_FAIR_VALUE_ADJUSTMENT",
    "DEFERRED_TAX_ASSET_ADJUSTMENT",
    "AGRICULTURAL_CATASTROPHE_RISK_RESERVE",
    "POLICY_SURPLUS_CORE_T1",
    "QUALIFYING_CORE_T1_LIABILITY_CAPITAL",
    "OTHER_CORE_T1_ADJUSTMENT",
)
EXTERNAL_ANC_T1_DETAIL_CODES = (
    "ANC_T1_SUBORDINATED_TERM_DEBT",
    "ANC_T1_CAPITAL_SUPPLEMENTARY_BONDS",
    "ANC_T1_CONVERTIBLE_SUBORDINATED_DEBT",
    "ANC_T1_DEFERRED_TAX_ASSET",
    "ANC_T1_INVESTMENT_PROPERTY_FAIR_VALUE",
    "POLICY_SURPLUS_ANC_T1",
    "OTHER_ANC_T1_CAPITAL",
)
CAPITAL_DETAIL_RECONCILIATION_ABSOLUTE_TOLERANCE = 0.05
CAPITAL_DETAIL_RECONCILIATION_RELATIVE_TOLERANCE = 2e-4  # 0.02%, filters source rounding tails.
EXTERNAL_UNDISCLOSED_MARKERS = frozenset({"-", "--", "—", "–", "/", "未披露", "不适用", "n/a", "na"})
INDUSTRY_TOTAL_COMPANY_CODE = "INDUSTRY_LIFE_TOTAL"
INDUSTRY_OUTPUT_CODES = (
    "COMBINED_SOLVENCY_RATIO",
    "CORE_SOLVENCY_RATIO",
    "COMBINED_SOLVENCY_SURPLUS",
    "CORE_SOLVENCY_SURPLUS",
    "RECOGNIZED_ASSETS",
    "RECOGNIZED_LIABILITIES",
    "ACTUAL_CAPITAL",
    "CORE_T1_CAPITAL",
    "CORE_T2_CAPITAL",
    "ANC_T1_CAPITAL",
    "ANC_T2_CAPITAL",
    "MINIMUM_CAPITAL",
    "TOTAL_ASSETS",
    "QUANT_RISK_CAPITAL",
    "INSURANCE_RISK_CAPITAL",
    "NON_LIFE_INSURANCE_RISK_CAPITAL",
    "MARKET_RISK_CAPITAL",
    "CREDIT_RISK_CAPITAL",
    "QUANT_RISK_DIVERSIFICATION_EFFECT",
    "CONTRACT_LOSS_ABSORPTION_EFFECT",
    "CONTROL_RISK_CAPITAL",
    "INSURANCE_CONTRACT_LIABILITY",
    "SEPARATE_ACCOUNT_LIABILITY",
    "POLICY_SURPLUS_CORE_T1",
    "POLICY_SURPLUS_CORE_T2",
    "POLICY_SURPLUS_ANC_T1",
    "POLICY_SURPLUS_ANC_T2",
    "FEATURE_FACTOR_IMPACT",
    "ACTUAL_CAPITAL_TO_RECOGNIZED_ASSETS",
    "MINIMUM_CAPITAL_TO_RECOGNIZED_LIABILITIES",
    "POLICY_SURPLUS_TO_INSURANCE_LIABILITIES",
    "RECOGNIZED_ASSETS_TO_ACTUAL_CAPITAL",
    "RECOGNIZED_ASSETS_TO_MINIMUM_CAPITAL",
)
INDUSTRY_RATIO_FORMULAS = {
    "COMBINED_SOLVENCY_RATIO": "行业实际资本合计/行业最低资本合计×100",
    "CORE_SOLVENCY_RATIO": "(行业核心一级资本合计+行业核心二级资本合计)/行业最低资本合计×100",
    "ACTUAL_CAPITAL_TO_RECOGNIZED_ASSETS": "行业实际资本合计/行业认可资产合计",
    "MINIMUM_CAPITAL_TO_RECOGNIZED_LIABILITIES": "行业最低资本合计/行业认可负债合计",
    "POLICY_SURPLUS_TO_INSURANCE_LIABILITIES": "行业已披露保单未来盈余合计/行业保险合同负债合计",
    "RECOGNIZED_ASSETS_TO_ACTUAL_CAPITAL": "行业认可资产合计/行业实际资本合计",
    "RECOGNIZED_ASSETS_TO_MINIMUM_CAPITAL": "行业认可资产合计/行业最低资本合计",
}


@dataclass
class ExternalConversionResult:
    data: pd.DataFrame
    sheet_summary: pd.DataFrame
    mapping_summary: pd.DataFrame
    logic_checks: pd.DataFrame
    warnings: list[str]


def canonical_metric_name(value: Any) -> str:
    text = unicodedata.normalize("NFKC", str(value or "")).replace("％", "%")
    return re.sub(r"\s+", "", text).strip()


def canonical_period_scope(value: Any) -> str:
    text = str(value or "").strip()
    if "预测" in text:
        return "下季度末预测数"
    if "上季度末" in text or "期初" in text:
        return "上季度末数"
    if (
        "累计" in text
        and any(term in text for term in ("本年", "本年度", "年度"))
        and not any(term in text for term in ("上年", "去年", "同期"))
    ):
        return "本年累计数"
    if (
        "本季度数" in text
        or "本季度（末）数" in text
        or "本季度(末)数" in text
        or "当季数" in text
    ) and "累计" not in text:
        return "本季度数"
    if (
        "本季度末" in text
        or "期末" in text
        or "本季度" in text
        or "本期" in text
    ) and "累计" not in text:
        return "本季度末数"
    return text


def _number(value: Any) -> float | None:
    numeric = parse_numeric(value)
    if pd.isna(numeric) or not math.isfinite(float(numeric)):
        return None
    return float(numeric)


def _blank_external_cell(value: Any) -> bool:
    return pd.isna(value) or (isinstance(value, str) and not value.strip())


def _external_undisclosed_cell(value: Any) -> bool:
    if _blank_external_cell(value):
        return True
    return str(value).strip().lower() in EXTERNAL_UNDISCLOSED_MARKERS


def _infer_capital_detail_zeros(
    source: pd.DataFrame,
    numeric: pd.DataFrame,
    source_column_by_code: dict[str, str],
) -> tuple[set[tuple[int, str]], list[tuple[int, str, float]]]:
    """Apply the external workbook convention within each populated T1 detail section.

    True blanks become inferred zeros when another component has a number and
    the disclosed components reconcile to the reported capital total. Explicit
    markers, entirely blank sections and materially unreconciled sections stay
    missing. An ancillary section with only zero values also stays missing:
    zero-only rows do not establish that its composition was disclosed.
    """
    inferred: set[tuple[int, str]] = set()
    discrepancies: list[tuple[int, str, float]] = []

    def amount(row_id: int, code: str) -> float | None:
        column = source_column_by_code.get(code)
        if column is None:
            return None
        value = numeric.iloc[row_id][column]
        return None if pd.isna(value) else float(value)

    for row_id in range(len(numeric)):
        for detail_codes, total_code, start_codes in (
            (EXTERNAL_CORE_T1_DETAIL_CODES, "CORE_T1_CAPITAL", ("FINANCIAL_STATEMENT_NET_ASSETS", "NET_ASSETS")),
            (EXTERNAL_ANC_T1_DETAIL_CODES, "ANC_T1_CAPITAL", ()),
        ):
            columns = [source_column_by_code[code] for code in detail_codes if code in source_column_by_code]
            disclosed = [float(value) for column in columns if pd.notna(value := numeric.iloc[row_id][column])]
            blanks = [
                column for column in columns
                if pd.isna(numeric.iloc[row_id][column])
                and _blank_external_cell(source.iloc[row_id][column])
            ]
            if not disclosed or not blanks:
                continue
            if detail_codes == EXTERNAL_ANC_T1_DETAIL_CODES and not any(disclosed):
                continue
            total = amount(row_id, total_code)
            start = next((value for code in start_codes if (value := amount(row_id, code)) is not None), None) if start_codes else 0.0
            if total is None or start is None:
                continue
            disclosed_total = math.fsum((start, *disclosed))
            residual = total - disclosed_total
            tolerance = max(
                CAPITAL_DETAIL_RECONCILIATION_ABSOLUTE_TOLERANCE,
                CAPITAL_DETAIL_RECONCILIATION_RELATIVE_TOLERANCE
                * max(abs(total), abs(disclosed_total), 1.0),
            )
            if abs(residual) > tolerance:
                discrepancies.append((row_id, total_code, residual))
                continue
            for column in blanks:
                numeric.iat[row_id, numeric.columns.get_loc(column)] = 0.0
                inferred.add((row_id, column))
    return inferred, discrepancies


def _safe_sum(values: dict[str, Any], codes: tuple[str, ...]) -> float | None:
    items = [_number(values.get(code)) for code in codes]
    if any(item is None for item in items):
        return None
    return float(sum(items))


def _sum_available(values: dict[str, Any], codes: tuple[str, ...]) -> float | None:
    """Sum disclosed amounts without turning undisclosed components into zero."""
    items = [_number(values.get(code)) for code in codes]
    available = [item for item in items if item is not None]
    return float(math.fsum(available)) if available else None


def _safe_divide(numerator: float | None, denominator: float | None) -> float | None:
    if numerator is None or denominator is None:
        return None
    if denominator == 0:
        return 0.0 if numerator == 0 else None
    return float(numerator / denominator)


def calculate_derived_values(values: dict[str, Any]) -> dict[str, Any]:
    """Calculate the CROSS derived fields from canonical base metric codes."""
    result: dict[str, Any] = {}
    get = lambda code: _number(result.get(code, values.get(code)))

    core_policy_surplus = _sum_available(values, POLICY_SURPLUS_CORE_CODES)
    all_policy_surplus = _sum_available(values, POLICY_SURPLUS_COMPONENT_CODES)
    quant_before_factor = _safe_sum(
        values,
        (
            "INSURANCE_RISK_CAPITAL",
            "NON_LIFE_INSURANCE_RISK_CAPITAL",
            "MARKET_RISK_CAPITAL",
            "CREDIT_RISK_CAPITAL",
            "QUANT_RISK_DIVERSIFICATION_EFFECT",
            "CONTRACT_LOSS_ABSORPTION_EFFECT",
        ),
    )

    result["POLICY_SURPLUS_CORE_TO_ACTUAL_CAPITAL"] = _safe_divide(core_policy_surplus, get("ACTUAL_CAPITAL"))
    ratio = result["POLICY_SURPLUS_CORE_TO_ACTUAL_CAPITAL"]
    if ratio is None:
        result["POLICY_SURPLUS_CORE_BAND"] = None
    elif ratio <= 0:
        result["POLICY_SURPLUS_CORE_BAND"] = "小于等于0%"
    elif ratio <= 0.20:
        result["POLICY_SURPLUS_CORE_BAND"] = "(0,20%]"
    elif ratio <= 0.35:
        result["POLICY_SURPLUS_CORE_BAND"] = "(20%,35%]"
    else:
        result["POLICY_SURPLUS_CORE_BAND"] = "大于35%"

    quant_capital = get("QUANT_RISK_CAPITAL")
    result["FEATURE_FACTOR_IMPACT"] = (
        None if quant_capital is None or quant_before_factor is None else quant_capital - quant_before_factor
    )
    result["ACTUAL_CAPITAL_TO_RECOGNIZED_ASSETS"] = _safe_divide(get("ACTUAL_CAPITAL"), get("RECOGNIZED_ASSETS"))
    result["CORE_T1_TO_ACTUAL_CAPITAL"] = _safe_divide(get("CORE_T1_CAPITAL"), get("ACTUAL_CAPITAL"))
    result["CORE_T2_TO_ACTUAL_CAPITAL"] = _safe_divide(get("CORE_T2_CAPITAL"), get("ACTUAL_CAPITAL"))
    result["ANC_T1_TO_ACTUAL_CAPITAL"] = _safe_divide(get("ANC_T1_CAPITAL"), get("ACTUAL_CAPITAL"))
    result["ANC_T2_TO_ACTUAL_CAPITAL"] = _safe_divide(get("ANC_T2_CAPITAL"), get("ACTUAL_CAPITAL"))
    result["MINIMUM_CAPITAL_TO_RECOGNIZED_LIABILITIES"] = _safe_divide(get("MINIMUM_CAPITAL"), get("RECOGNIZED_LIABILITIES"))
    result["POLICY_SURPLUS_TO_INSURANCE_LIABILITIES"] = _safe_divide(
        all_policy_surplus,
        get("INSURANCE_CONTRACT_LIABILITY"),
    )
    result["RECOGNIZED_ASSETS_TO_ACTUAL_CAPITAL"] = _safe_divide(get("RECOGNIZED_ASSETS"), get("ACTUAL_CAPITAL"))
    result["RECOGNIZED_ASSETS_TO_MINIMUM_CAPITAL"] = _safe_divide(get("RECOGNIZED_ASSETS"), get("MINIMUM_CAPITAL"))
    result["RECOGNIZED_ASSETS_TO_REGISTERED_CAPITAL"] = _safe_divide(get("RECOGNIZED_ASSETS"), get("REGISTERED_CAPITAL"))
    result["ACTUAL_CAPITAL_TO_REGISTERED_CAPITAL"] = _safe_divide(get("ACTUAL_CAPITAL"), get("REGISTERED_CAPITAL"))
    result["FEATURE_FACTOR_CHECK"] = _safe_divide(get("FEATURE_FACTOR_IMPACT"), quant_before_factor)
    result["LIFE_INSURANCE_RISK_TO_LIABILITIES"] = _safe_divide(get("INSURANCE_RISK_CAPITAL"), get("RECOGNIZED_LIABILITIES"))
    result["NON_LIFE_INSURANCE_RISK_TO_LIABILITIES"] = _safe_divide(get("NON_LIFE_INSURANCE_RISK_CAPITAL"), get("RECOGNIZED_LIABILITIES"))
    result["MARKET_RISK_TO_ASSETS"] = _safe_divide(get("MARKET_RISK_CAPITAL"), get("RECOGNIZED_ASSETS"))
    result["CREDIT_RISK_TO_ASSETS"] = _safe_divide(get("CREDIT_RISK_CAPITAL"), get("RECOGNIZED_ASSETS"))
    core_capital = _safe_sum(values, ("CORE_T1_CAPITAL", "CORE_T2_CAPITAL"))
    result["CORE_CAPITAL_TO_REGISTERED_CAPITAL"] = _safe_divide(core_capital, get("REGISTERED_CAPITAL"))
    result["REGISTERED_CAPITAL_TO_CORE_CAPITAL"] = _safe_divide(get("REGISTERED_CAPITAL"), core_capital)
    if result["REGISTERED_CAPITAL_TO_CORE_CAPITAL"] is None:
        result["REGISTERED_CAPITAL_TO_CORE_CAPITAL"] = _safe_divide(
            1.0,
            _number(values.get("CORE_CAPITAL_TO_REGISTERED_CAPITAL")),
        )
    result["POLICY_SURPLUS_CORE_TO_CORE_CAPITAL"] = _safe_divide(core_policy_surplus, core_capital)
    result["CORE_T1_POLICY_SURPLUS_SHARE"] = _safe_divide(get("POLICY_SURPLUS_CORE_T1"), get("CORE_T1_CAPITAL"))
    result["ANC_T1_POLICY_SURPLUS_SHARE"] = _safe_divide(get("POLICY_SURPLUS_ANC_T1"), get("ANC_T1_CAPITAL"))
    result["INTEREST_RATE_RISK_TO_ASSETS"] = _safe_divide(get("INTEREST_RATE_RISK_CAPITAL"), get("RECOGNIZED_ASSETS"))
    result["EQUITY_RISK_TO_ASSETS"] = _safe_divide(get("EQUITY_RISK_CAPITAL"), get("RECOGNIZED_ASSETS"))
    result["SPREAD_RISK_TO_ASSETS"] = _safe_divide(get("SPREAD_RISK_CAPITAL"), get("RECOGNIZED_ASSETS"))
    result["COUNTERPARTY_RISK_TO_ASSETS"] = _safe_divide(get("COUNTERPARTY_RISK_CAPITAL"), get("RECOGNIZED_ASSETS"))
    result["LIFE_INSURANCE_RISK_TO_QUANT_CAPITAL"] = _safe_divide(get("INSURANCE_RISK_CAPITAL"), quant_capital)
    result["NON_LIFE_INSURANCE_RISK_TO_QUANT_CAPITAL"] = _safe_divide(get("NON_LIFE_INSURANCE_RISK_CAPITAL"), quant_capital)
    result["MARKET_RISK_TO_QUANT_CAPITAL"] = _safe_divide(get("MARKET_RISK_CAPITAL"), quant_capital)
    result["CREDIT_RISK_TO_QUANT_CAPITAL"] = _safe_divide(get("CREDIT_RISK_CAPITAL"), quant_capital)
    result["DIVERSIFICATION_EFFECT_TO_QUANT_CAPITAL"] = _safe_divide(get("QUANT_RISK_DIVERSIFICATION_EFFECT"), quant_capital)
    result["LOSS_ABSORPTION_TO_QUANT_CAPITAL"] = _safe_divide(get("CONTRACT_LOSS_ABSORPTION_EFFECT"), quant_capital)
    result["TOTAL_ASSETS_TO_REGISTERED_CAPITAL"] = _safe_divide(get("TOTAL_ASSETS"), get("REGISTERED_CAPITAL"))
    return result


def calculate_industry_values(company_totals: dict[str, Any]) -> dict[str, Any]:
    """Calculate the populated industry row in the PPT draft ``汇总`` sheet.

    Amount fields are supplied as company totals. Ratios are recalculated from
    those totals instead of averaging company ratios. Two formulas in the draft
    contradict their labels, so this implementation follows the metric meaning:
    separate-account liabilities are summed and recognized-assets/minimum-capital
    uses recognized assets as its numerator.
    """
    result = dict(company_totals)
    get = lambda code: _number(result.get(code))

    result["COMBINED_SOLVENCY_RATIO"] = _safe_divide(
        get("ACTUAL_CAPITAL"), get("MINIMUM_CAPITAL")
    )
    core_capital = _safe_sum(result, ("CORE_T1_CAPITAL", "CORE_T2_CAPITAL"))
    result["CORE_SOLVENCY_RATIO"] = _safe_divide(core_capital, get("MINIMUM_CAPITAL"))
    for code in PERCENT_FROM_RATIO_CODES:
        if result.get(code) is not None:
            result[code] = float(result[code]) * 100

    all_policy_surplus = _sum_available(result, POLICY_SURPLUS_COMPONENT_CODES)
    result["ACTUAL_CAPITAL_TO_RECOGNIZED_ASSETS"] = _safe_divide(
        get("ACTUAL_CAPITAL"), get("RECOGNIZED_ASSETS")
    )
    result["MINIMUM_CAPITAL_TO_RECOGNIZED_LIABILITIES"] = _safe_divide(
        get("MINIMUM_CAPITAL"), get("RECOGNIZED_LIABILITIES")
    )
    result["POLICY_SURPLUS_TO_INSURANCE_LIABILITIES"] = _safe_divide(
        all_policy_surplus, get("INSURANCE_CONTRACT_LIABILITY")
    )
    result["RECOGNIZED_ASSETS_TO_ACTUAL_CAPITAL"] = _safe_divide(
        get("RECOGNIZED_ASSETS"), get("ACTUAL_CAPITAL")
    )
    result["RECOGNIZED_ASSETS_TO_MINIMUM_CAPITAL"] = _safe_divide(
        get("RECOGNIZED_ASSETS"), get("MINIMUM_CAPITAL")
    )
    return result


def _industry_company_totals(
    frame: pd.DataFrame,
    code_by_column: dict[str, str],
) -> dict[str, float]:
    source_column_by_code: dict[str, str] = {}
    for column, code in code_by_column.items():
        source_column_by_code.setdefault(code, column)

    ratio_codes = set(INDUSTRY_RATIO_FORMULAS)
    totals: dict[str, float] = {}
    for code in INDUSTRY_OUTPUT_CODES:
        if code in ratio_codes:
            continue
        column = source_column_by_code.get(code)
        if column is None:
            continue
        numeric = frame[column].map(_number).dropna()
        if not numeric.empty:
            totals[code] = float(numeric.sum())
    return totals


def _industry_metric_records(
    frame: pd.DataFrame,
    code_by_column: dict[str, str],
    lookup: dict[str, dict],
    *,
    filename: str,
    sheet_name: str,
    report_profile_id: str,
    year: int,
    quarter: str,
    report_period: str,
) -> list[dict]:
    totals = _industry_company_totals(frame, code_by_column)
    values = calculate_industry_values(totals)
    metrics_by_code: dict[str, dict] = {}
    for metric in lookup.values():
        metrics_by_code.setdefault(str(metric.get("指标编码", "")), metric)

    records: list[dict] = []
    for source_code in INDUSTRY_OUTPUT_CODES:
        value = values.get(source_code)
        if value is None or not math.isfinite(float(value)):
            continue
        source_metric = metrics_by_code.get(source_code)
        if source_metric is None:
            continue

        renamed = INDUSTRY_METRICS_BY_SOURCE_CODE.get(source_code)
        if renamed is None:
            metric_code = source_code
            metric_name = source_metric["指标名称"]
            level1 = source_metric["一级模块"]
            level2 = source_metric["二级模块"]
            unit = source_metric["单位"]
            data_type = source_metric["数据类型"]
            rename_note = ""
        else:
            metric_code = renamed.code
            metric_name = renamed.name
            level1 = renamed.level1
            level2 = renamed.level2
            unit = renamed.unit
            data_type = renamed.data_type
            rename_note = f"；行业指标由“{source_metric['指标名称']}”重命名"

        formula = INDUSTRY_RATIO_FORMULAS.get(
            source_code,
            f"所有有效公司的{source_metric['指标名称']}求和（忽略空值）",
        )
        records.append({
            "公司": "行业合计",
            "原始公司名称": "行业合计",
            "标准公司名称": "行业合计",
            "公司统一编码": INDUSTRY_TOTAL_COMPANY_CODE,
            "公司类型": "行业合计",
            "同业分类": "全行业",
            "报告类型": report_profile_id,
            "报告年度": year,
            "报告季度": quarter,
            "报告期": report_period,
            "披露日期": "",
            "一级模块": level1,
            "二级模块": level2,
            "行次": "",
            "指标编码": metric_code,
            "指标名称": metric_name,
            "期间口径": "本季度末数",
            "数值": float(value),
            "单位": unit,
            "数据类型": data_type,
            "是否预测": "否",
            "来源页码": "",
            "原始披露值": "",
            "备注": f"按PPT底稿“汇总”页行业口径，由本工作表公司数据计算{rename_note}",
            "来源类型": "系统计算",
            "指标属性": "行业计算",
            "来源文件": filename,
            "来源工作表": sheet_name,
            "导入批次": f"{Path(filename).stem}:{report_period}",
            "计算逻辑": formula,
        })
    return records


def _standardize_with_company_identities(frame: pd.DataFrame) -> pd.DataFrame:
    """Standardize a frame without re-resolving identities already supplied."""
    source = standardize_uploaded_frame(frame)
    identity_columns = ("公司", "原始公司名称", "标准公司名称", "公司统一编码", "公司类型")
    if source.empty or all(
        source[column].fillna("").astype(str).str.strip().ne("").all()
        for column in identity_columns
    ):
        return reconcile_known_company_aliases(source)
    return apply_company_identities(source)


def append_derived_metrics(frame: pd.DataFrame) -> pd.DataFrame:
    source = _standardize_with_company_identities(frame)
    if source.empty:
        return source
    if "报告类型" not in source.columns:
        source["报告类型"] = ""
    source = source[~source["指标编码"].isin([item.code for item in DERIVED_METRICS])].copy()
    source["_期间组"] = source["期间口径"].map(canonical_period_scope)
    grouping = ["报告类型", "公司统一编码", "报告年度", "报告季度", "报告期", "_期间组"]
    derived_records: list[dict] = []

    for _, group in source.groupby(grouping, dropna=False, sort=False):
        values: dict[str, Any] = {}
        for code, code_rows in group.groupby("指标编码", sort=False):
            if str(code) in POLICY_SURPLUS_COMPONENT_CODES:
                code_rows = code_rows.loc[
                    ~code_rows["来源类型"].fillna("").astype(str).eq("宽表空白推定")
                ]
            candidates = pd.to_numeric(code_rows["数值"], errors="coerce").dropna()
            if not candidates.empty:
                values[str(code)] = float(candidates.iloc[0])
        calculated = calculate_derived_values(values)
        base = group.iloc[0]
        source_pages = "、".join(dict.fromkeys(str(item) for item in group["来源页码"] if str(item).strip()))
        source_sheets = "、".join(dict.fromkeys(str(item) for item in group["来源工作表"] if str(item).strip()))
        for definition in DERIVED_METRICS:
            value = calculated.get(definition.code)
            if value is None or (isinstance(value, float) and not math.isfinite(value)):
                continue
            note = "由标准化基础指标自动计算"
            if definition.code in POLICY_SURPLUS_RATIO_COMPONENTS:
                components = POLICY_SURPLUS_RATIO_COMPONENTS[definition.code]
                if any(code not in values for code in components):
                    note = "仅汇总有披露数值的保单未来盈余层级；未披露层级未计入分子"
            derived_records.append({
                "公司": base.get("公司", ""),
                "原始公司名称": base.get("原始公司名称", ""),
                "标准公司名称": base.get("标准公司名称", ""),
                "公司统一编码": base.get("公司统一编码", ""),
                "公司类型": base.get("公司类型", ""),
                "同业分类": base.get("同业分类", ""),
                "报告类型": base.get("报告类型", ""),
                "报告年度": base.get("报告年度", ""),
                "报告季度": base.get("报告季度", ""),
                "报告期": base.get("报告期", ""),
                "披露日期": base.get("披露日期", ""),
                "一级模块": definition.level1,
                "二级模块": definition.level2,
                "行次": "",
                "指标编码": definition.code,
                "指标名称": definition.name,
                "期间口径": base.get("_期间组", "本季度末数"),
                "数值": value,
                "单位": definition.unit,
                "数据类型": definition.data_type,
                "是否预测": "是" if "预测" in str(base.get("_期间组", "")) else "否",
                "来源页码": source_pages,
                "原始披露值": "",
                "备注": note,
                "来源类型": "系统计算",
                "指标属性": definition.attribute,
                "来源文件": base.get("来源文件", ""),
                "来源工作表": source_sheets,
                "导入批次": base.get("导入批次", ""),
                "计算逻辑": definition.formula,
            })

    source = source.drop(columns=["_期间组"])
    if not derived_records:
        return standardize_uploaded_frame(source)
    return pd.concat(
        [standardize_uploaded_frame(source), pd.DataFrame(derived_records, columns=STANDARD_COLUMNS)],
        ignore_index=True,
    )


def add_missing_derived_metrics(frame: pd.DataFrame) -> pd.DataFrame:
    """Fill calculable derived metrics and enforce agreed policy-surplus formulas.

    ``append_derived_metrics`` is intentionally authoritative for STEP3: it removes
    existing derived rows and recalculates them. Integrated workbooks preserve reviewed
    external ratios except the policy-surplus ratios whose business rules require the
    available disclosed capital layers to be recalculated consistently across STEP3-7.
    """
    source = _standardize_with_company_identities(frame)
    if source.empty:
        return source

    derived_codes = {item.code for item in DERIVED_METRICS}
    calculated = append_derived_metrics(source)
    candidates = calculated[calculated["指标编码"].isin(derived_codes)].copy()
    legacy_rows = source[
        source["指标编码"].astype(str).eq("CORE_CAPITAL_TO_REGISTERED_CAPITAL")
    ].copy()
    if not legacy_rows.empty:
        legacy_values = pd.to_numeric(legacy_rows["数值"], errors="coerce")
        legacy_rows = legacy_rows[legacy_values.notna() & legacy_values.ne(0)].copy()
        if not legacy_rows.empty:
            definition = next(
                item for item in DERIVED_METRICS
                if item.code == "REGISTERED_CAPITAL_TO_CORE_CAPITAL"
            )
            legacy_rows["数值"] = 1.0 / pd.to_numeric(legacy_rows["数值"], errors="coerce")
            legacy_rows["指标编码"] = definition.code
            legacy_rows["指标名称"] = definition.name
            legacy_rows["一级模块"] = definition.level1
            legacy_rows["二级模块"] = definition.level2
            legacy_rows["单位"] = definition.unit
            legacy_rows["数据类型"] = definition.data_type
            legacy_rows["指标属性"] = definition.attribute
            legacy_rows["来源类型"] = "系统计算"
            legacy_rows["备注"] = "由旧版核心资本/注册资本指标取倒数自动转换"
            legacy_rows["计算逻辑"] = "1/(核心资本/注册资本)"
            candidates = pd.concat([candidates, legacy_rows], ignore_index=True)

    existing = source[source["指标编码"].isin(derived_codes)].copy()
    key_columns = [
        "报告类型",
        "公司统一编码",
        "报告年度",
        "报告季度",
        "报告期",
        "_期间组",
        "指标编码",
    ]

    def row_keys(rows: pd.DataFrame) -> pd.Series:
        keyed = rows.copy()
        keyed["_期间组"] = keyed["期间口径"].map(canonical_period_scope)
        values = keyed[key_columns].fillna("").astype(str)
        return values.apply(tuple, axis=1)

    if candidates.empty:
        return source
    candidates = candidates.loc[~row_keys(candidates).duplicated(keep="first")].copy()
    candidate_by_key = dict(zip(row_keys(candidates), candidates.index))
    authoritative = existing["指标编码"].isin(AUTHORITATIVE_POLICY_SURPLUS_DERIVED_CODES)
    legacy_unavailable = (
        existing["指标编码"].isin(POLICY_SURPLUS_RATIO_COMPONENTS)
        & pd.to_numeric(existing["数值"], errors="coerce").isna()
        & existing["披露状态"].fillna("").astype(str).eq("无法计算")
    )
    replaceable = existing.loc[authoritative | legacy_unavailable]
    for index, key in zip(replaceable.index, row_keys(replaceable)):
        candidate_index = candidate_by_key.get(key)
        if candidate_index is None:
            continue
        original_value = source.at[index, "数值"]
        original_disclosure = source.at[index, "原始披露值"]
        if (
            (pd.isna(original_disclosure) or not str(original_disclosure).strip())
            and pd.notna(original_value)
        ):
            source.at[index, "原始披露值"] = original_value
        for column in ("数值", "备注", "来源类型", "指标属性", "计算逻辑", "披露状态"):
            source.at[index, column] = candidates.at[candidate_index, column]
        if str(source.at[index, "指标编码"]) in AUTHORITATIVE_POLICY_SURPLUS_DERIVED_CODES:
            source.at[index, "备注"] = (
                "外部派生值按系统统一口径重算；"
                + str(candidates.at[candidate_index, "备注"] or "")
            ).rstrip("；")
        source.at[index, "披露状态"] = "已计算"
    existing_keys = set(row_keys(existing).tolist()) if not existing.empty else set()
    missing_mask = ~row_keys(candidates).isin(existing_keys)
    additions = candidates.loc[missing_mask]
    if additions.empty:
        return source
    return pd.concat([source, additions], ignore_index=True)


def _append_legacy_reversed_ratio(frame: pd.DataFrame) -> pd.DataFrame:
    """Add only the legacy reciprocal ratio needed by direct wide conversion."""
    source = standardize_uploaded_frame(frame)
    legacy_rows = source[
        source["指标编码"].astype(str).eq("CORE_CAPITAL_TO_REGISTERED_CAPITAL")
    ].copy()
    if legacy_rows.empty:
        return source

    legacy_values = pd.to_numeric(legacy_rows["数值"], errors="coerce")
    legacy_rows = legacy_rows[legacy_values.notna() & legacy_values.ne(0)].copy()
    if legacy_rows.empty:
        return source

    definition = next(
        item for item in DERIVED_METRICS
        if item.code == "REGISTERED_CAPITAL_TO_CORE_CAPITAL"
    )
    legacy_rows["数值"] = 1.0 / pd.to_numeric(legacy_rows["数值"], errors="coerce")
    legacy_rows["指标编码"] = definition.code
    legacy_rows["指标名称"] = definition.name
    legacy_rows["一级模块"] = definition.level1
    legacy_rows["二级模块"] = definition.level2
    legacy_rows["单位"] = definition.unit
    legacy_rows["数据类型"] = definition.data_type
    legacy_rows["指标属性"] = definition.attribute
    legacy_rows["来源类型"] = "系统计算"
    legacy_rows["备注"] = "由旧版核心资本/注册资本指标取倒数自动转换"
    legacy_rows["计算逻辑"] = "1/(核心资本/注册资本)"

    existing = source[
        source["指标编码"].astype(str).eq(definition.code)
    ]
    key_columns = [
        "报告类型", "公司统一编码", "报告年度", "报告季度",
        "报告期", "期间口径", "指标编码",
    ]
    existing_keys = set(
        existing[key_columns].fillna("").astype(str).itertuples(index=False, name=None)
    )
    candidate_keys = legacy_rows[key_columns].fillna("").astype(str).apply(tuple, axis=1)
    additions = legacy_rows.loc[~candidate_keys.isin(existing_keys)]
    if additions.empty:
        return source
    return pd.concat([source, additions], ignore_index=True)


def _metric_lookup(
    taxonomy: pd.DataFrame,
    *,
    include_step3_only: bool = False,
) -> dict[str, dict]:
    lookup: dict[str, dict] = {}
    scoped_taxonomy = taxonomy
    if not include_step3_only and "STEP5宽表映射" in taxonomy.columns:
        wide_flag = taxonomy["STEP5宽表映射"].fillna("").astype(str).str.strip()
        scoped_taxonomy = taxonomy.loc[~wide_flag.isin({"否", "false", "False", "0"})]
    for _, row in scoped_taxonomy.iterrows():
        names = [row.get("指标名称", "")]
        names.extend(str(row.get("别名", "")).split("|"))
        payload = {
            "指标编码": str(row.get("指标编码", "")).strip(),
            "指标名称": str(row.get("指标名称", "")).strip(),
            "一级模块": str(row.get("一级模块", "")).strip(),
            "二级模块": str(row.get("二级模块", "")).strip(),
            "单位": str(row.get("标准单位", "")).strip(),
            "数据类型": str(row.get("数据类型", "")).strip(),
            "指标属性": "披露",
            "计算逻辑": "",
        }
        for name in names:
            key = canonical_metric_name(name)
            if key:
                lookup.setdefault(key, payload)
    for definition in DERIVED_METRICS:
        lookup[canonical_metric_name(definition.name)] = {
            "指标编码": definition.code,
            "指标名称": definition.name,
            "一级模块": definition.level1,
            "二级模块": definition.level2,
            "单位": definition.unit,
            "数据类型": definition.data_type,
            "指标属性": definition.attribute,
            "计算逻辑": definition.formula,
        }
    return lookup


def supported_metric_catalog(
    taxonomy: pd.DataFrame,
    *,
    include_derived: bool = True,
    include_step3_only: bool = False,
) -> pd.DataFrame:
    """Return the canonical metric-code range accepted by STEP5 exact mapping.

    STEP3 uses this public catalog as its target-table source so both workflow
    steps share names, modules, units, data types, attributes and formulas.
    """
    lookup = _metric_lookup(
        taxonomy,
        include_step3_only=include_step3_only,
    )
    catalog = pd.DataFrame(lookup.values())
    if include_step3_only:
        # STEP3 selects canonical codes. Name-indexed lookup can lose distinct
        # checklist items that share a label (e.g. net assets and deductions).
        disclosed = taxonomy.rename(columns={'标准单位': '单位'}).copy()
        disclosed['指标属性'] = '披露'
        disclosed['计算逻辑'] = ''
        derived_codes = {definition.code for definition in DERIVED_METRICS}
        disclosed = disclosed[~disclosed['指标编码'].isin(derived_codes)]
        catalog = pd.concat([disclosed, catalog[catalog['指标编码'].isin(derived_codes)]], ignore_index=True)
    if catalog.empty:
        return pd.DataFrame(columns=[
            "指标编码", "指标名称", "一级模块", "二级模块", "标准单位",
            "数据类型", "指标属性", "计算逻辑",
        ])
    if not include_derived:
        derived_codes = {definition.code for definition in DERIVED_METRICS}
        catalog = catalog[~catalog["指标编码"].isin(derived_codes)]
    catalog = catalog.drop_duplicates(subset="指标编码", keep="last").copy()
    catalog = catalog.rename(columns={"单位": "标准单位"})
    return catalog.reindex(columns=[
        "指标编码", "指标名称", "一级模块", "二级模块", "标准单位",
        "数据类型", "指标属性", "计算逻辑",
    ]).reset_index(drop=True)


def _external_value(raw_value: Any, metric: dict) -> Any:
    if metric["数据类型"] in {"文本", "评级", "布尔"}:
        return "" if pd.isna(raw_value) else str(raw_value).strip()
    numeric = _number(raw_value)
    if numeric is None:
        return np.nan
    if metric["指标编码"] in PERCENT_FROM_RATIO_CODES or metric.get("源比例小数"):
        text = str(raw_value).strip()
        return numeric if text.endswith(("%", "％")) else numeric * 100
    return numeric


def _company_type(value: Any) -> str:
    text = str(value or "").strip()
    return {"养老": "养老险", "健康": "健康险"}.get(text, text or "未分类")


def convert_external_workbook(
    workbook_bytes: bytes,
    filename: str,
    taxonomy: pd.DataFrame,
    company_type_map: dict[str, str] | None = None,
    report_profile_id: str = "LIFE_SOLVENCY",
) -> ExternalConversionResult:
    excel = pd.ExcelFile(io.BytesIO(workbook_bytes))
    quarter_sheets = [name for name in excel.sheet_names if QUARTER_SHEET_PATTERN.fullmatch(str(name).strip())]
    if not quarter_sheets:
        raise ValueError("未找到名称形如 2025Q4 的季度数据工作表。")

    lookup = _metric_lookup(taxonomy)
    for source_label, (metric_name, period_scope) in EXTERNAL_INVESTMENT_COLUMNS.items():
        metric = lookup.get(canonical_metric_name(metric_name))
        existing = lookup.get(source_label)
        if metric is not None and (existing is None or existing["指标编码"] == metric["指标编码"]):
            lookup[source_label] = {
                **metric,
                "期间口径": period_scope,
                "源比例小数": True,
            }
    for source_label, (metric_name, source_ratio) in EXTERNAL_OPERATING_QUARTER_COLUMNS.items():
        metric = lookup.get(canonical_metric_name(metric_name))
        existing = lookup.get(source_label)
        if metric is not None and (existing is None or existing["指标编码"] == metric["指标编码"]):
            lookup[source_label] = {
                **metric,
                "期间口径": "本季度数",
                "源比例小数": source_ratio,
            }
    company_type_map = {str(key).strip(): _company_type(value) for key, value in (company_type_map or {}).items()}
    output_frames: list[pd.DataFrame] = []
    sheet_rows: list[dict] = []
    mapping_rows: dict[str, dict] = {}
    logic_rows: list[dict] = []
    warnings: list[str] = []
    unknown_companies: set[str] = set()

    for sheet_name in quarter_sheets:
        match = QUARTER_SHEET_PATTERN.fullmatch(sheet_name.strip())
        assert match is not None
        year = int(match.group("year"))
        quarter = match.group("quarter").upper()
        report_period = f"{year}{quarter}"
        frame = pd.read_excel(excel, sheet_name=sheet_name, header=0)
        frame.columns = [str(column).strip() for column in frame.columns]
        missing_identifiers = IDENTIFIER_COLUMNS - set(frame.columns)
        if missing_identifiers:
            raise ValueError(f"工作表 {sheet_name} 缺少字段：{', '.join(sorted(missing_identifiers))}")
        frame = frame[frame["公司"].notna() & (frame["公司"].astype(str).str.strip() != "")].copy()
        metric_columns = [column for column in frame.columns if column not in IDENTIFIER_COLUMNS]
        unmapped = [column for column in metric_columns if canonical_metric_name(column) not in lookup]
        if unmapped:
            raise ValueError(f"工作表 {sheet_name} 存在未精确映射指标：{', '.join(unmapped)}")

        metric_by_column = {
            column: lookup[canonical_metric_name(column)]
            for column in metric_columns
        }
        code_by_column = {
            column: metric_by_column[column]["指标编码"]
            for column in metric_columns
        }
        for column in metric_columns:
            metric = metric_by_column[column]
            mapping_rows.setdefault(column, {
                "来源字段": column,
                "匹配方式": "指标名称精确匹配",
                **{key: value for key, value in metric.items() if key != "源比例小数"},
            })

        failed_logic: dict[str, list[float]] = {item.code: [] for item in DERIVED_METRICS}
        checked_logic: dict[str, int] = {item.code: 0 for item in DERIVED_METRICS}
        missing_values: dict[str, int] = {}
        invalid_values: dict[str, int] = {}

        identities: list[dict[str, Any]] = []
        for row_id, (company_value, peer_value) in enumerate(
            zip(frame["公司"].tolist(), frame["分类"].tolist())
        ):
            company = str(company_value).strip()
            identity = resolve_company_identity(company, company_type_map)
            if identity.company_type == "未分类":
                unknown_companies.add(company)
            identities.append({
                "_row_id": row_id,
                "公司": identity.standard_name,
                "原始公司名称": identity.original_name,
                "标准公司名称": identity.standard_name,
                "公司统一编码": identity.company_code,
                "公司类型": identity.company_type,
                "同业分类": str(peer_value).strip(),
            })

        numeric_frame = frame[metric_columns].apply(
            lambda series: series.map(_number)
        )
        source_column_by_code: dict[str, str] = {}
        for column, code in code_by_column.items():
            source_column_by_code.setdefault(code, column)
        inferred_zero_keys, capital_discrepancies = _infer_capital_detail_zeros(
            frame, numeric_frame, source_column_by_code
        )
        raw_arrays = {column: frame[column].tolist() for column in metric_columns}
        for row_id, numeric_values in enumerate(
            numeric_frame.itertuples(index=False, name=None)
        ):
            raw_by_code = {
                code_by_column[column]: value
                for column, value in zip(metric_columns, numeric_values)
                if value is not None and not pd.isna(value)
            }
            calculated = calculate_derived_values(raw_by_code)
            for definition in DERIVED_METRICS:
                source_column = source_column_by_code.get(definition.code)
                if not source_column:
                    continue
                actual = raw_arrays[source_column][row_id]
                expected = calculated.get(definition.code)
                if expected is None or pd.isna(actual):
                    continue
                checked_logic[definition.code] += 1
                if definition.data_type == "文本":
                    if str(actual).strip() != str(expected).strip():
                        failed_logic[definition.code].append(float("nan"))
                else:
                    actual_number = _number(actual)
                    if actual_number is None:
                        failed_logic[definition.code].append(float("nan"))
                    else:
                        difference = abs(actual_number - float(expected))
                        tolerance = 1e-8 * max(1.0, abs(float(expected)))
                        if difference > tolerance:
                            failed_logic[definition.code].append(difference)

        normalized_columns: dict[str, pd.Series] = {}
        for column in metric_columns:
            metric = metric_by_column[column]
            if metric["数据类型"] in {"文本", "评级", "布尔"}:
                values = frame[column].map(
                    lambda value: "" if pd.isna(value) else str(value).strip()
                )
                valid = values.ne("")
                undisclosed = ~valid
            else:
                values = numeric_frame[column].astype(float)
                if metric["指标编码"] in PERCENT_FROM_RATIO_CODES or metric.get("源比例小数"):
                    percent_mask = frame[column].astype(str).str.strip().str.endswith(("%", "％"))
                    values = values.where(percent_mask, values * 100)
                valid = values.notna()
                undisclosed = frame[column].map(_external_undisclosed_cell)
            missing_count = int(((~valid) & undisclosed).sum())
            invalid_count = int(((~valid) & ~undisclosed).sum())
            if missing_count:
                missing_values[column] = missing_count
            if invalid_count:
                invalid_values[column] = invalid_count
            normalized_columns[column] = values

        normalized_values = pd.DataFrame(normalized_columns, index=frame.index)

        raw_wide = frame[metric_columns].reset_index(drop=True).copy()
        raw_wide.insert(0, "_row_id", np.arange(len(raw_wide)))
        value_wide = normalized_values.reset_index(drop=True).copy()
        value_wide.insert(0, "_row_id", np.arange(len(value_wide)))
        raw_long = raw_wide.melt(
            id_vars="_row_id",
            value_vars=metric_columns,
            var_name="_source_column",
            value_name="原始披露值",
        )
        value_long = value_wide.melt(
            id_vars="_row_id",
            value_vars=metric_columns,
            var_name="_source_column",
            value_name="数值",
        )
        value_long["原始披露值"] = raw_long["原始披露值"]

        metric_metadata = pd.DataFrame([
            {
                "_source_column": column,
                "_metric_order": order,
                **metric_by_column[column],
            }
            for order, column in enumerate(metric_columns)
        ])
        company_records = (
            value_long
            .merge(metric_metadata, on="_source_column", how="left", validate="many_to_one")
            .merge(pd.DataFrame(identities), on="_row_id", how="left", validate="many_to_one")
        )
        text_mask = company_records["数据类型"].isin({"文本", "评级", "布尔"})
        company_records = company_records.loc[
            (text_mask & company_records["数值"].ne(""))
            | (~text_mask & company_records["数值"].notna())
        ].sort_values(["_row_id", "_metric_order"], kind="stable")
        company_records["报告类型"] = report_profile_id
        company_records["报告年度"] = year
        company_records["报告季度"] = quarter
        company_records["报告期"] = report_period
        company_records["披露日期"] = ""
        company_records["行次"] = ""
        if "期间口径" in company_records:
            company_records["期间口径"] = company_records["期间口径"].fillna("本季度末数")
        else:
            company_records["期间口径"] = "本季度末数"
        company_records["是否预测"] = "否"
        company_records["来源页码"] = ""
        company_records["备注"] = "外部宽表按指标名称精确映射"
        company_records["来源类型"] = "外部数据集"
        inferred_mask = np.fromiter(
            ((int(row_id), column) in inferred_zero_keys for row_id, column in zip(
                company_records["_row_id"], company_records["_source_column"]
            )),
            dtype=bool,
            count=len(company_records),
        )
        company_records.loc[inferred_mask, "备注"] = "外部宽表同组资本明细已有数值，空白按0推定"
        company_records.loc[inferred_mask, "来源类型"] = "宽表空白推定"
        company_records.loc[inferred_mask, "指标属性"] = "推定"
        company_records.loc[inferred_mask, "披露状态"] = "推定零值"
        company_records["来源文件"] = filename
        company_records["来源工作表"] = sheet_name
        company_records["导入批次"] = f"{Path(filename).stem}:{report_period}"
        company_records = company_records.reindex(columns=STANDARD_COLUMNS)
        company_record_count = len(company_records)
        output_frames.append(company_records)

        industry_records = _industry_metric_records(
            frame,
            code_by_column,
            lookup,
            filename=filename,
            sheet_name=sheet_name,
            report_profile_id=report_profile_id,
            year=year,
            quarter=quarter,
            report_period=report_period,
        )
        if industry_records:
            output_frames.append(pd.DataFrame(industry_records, columns=STANDARD_COLUMNS))

        for definition in DERIVED_METRICS:
            if checked_logic[definition.code] == 0:
                continue
            differences = [item for item in failed_logic[definition.code] if math.isfinite(item)]
            failed = bool(failed_logic[definition.code])
            status = "通过"
            if failed:
                status = (
                    "已按系统口径重算"
                    if definition.code in AUTHORITATIVE_POLICY_SURPLUS_DERIVED_CODES
                    else "需复核"
                )
            logic_rows.append({
                "来源工作表": sheet_name,
                "指标编码": definition.code,
                "指标名称": definition.name,
                "校验行数": checked_logic[definition.code],
                "不一致行数": len(failed_logic[definition.code]),
                "最大绝对差异": max(differences) if differences else 0.0,
                "状态": status,
                "计算逻辑": definition.formula,
            })
        sheet_rows.append({
            "来源文件": filename,
            "来源工作表": sheet_name,
            "报告期": report_period,
            "公司数": int(frame["公司"].nunique()),
            "指标数": len(metric_columns),
            "源数据单元格数": len(frame) * len(metric_columns),
            "转换记录数": company_record_count + len(industry_records),
            "公司转换记录数": company_record_count,
            "行业指标数": len(industry_records),
            "跳过空值或错误值": sum(missing_values.values()) + sum(invalid_values.values()),
            "空值或未披露标记数": sum(missing_values.values()),
            "无效值数": sum(invalid_values.values()),
            "资本明细推定零值数": len(inferred_zero_keys),
            "资本明细与合计不一致组数": len(capital_discrepancies),
        })
        if capital_discrepancies:
            capital_names = {
                "CORE_T1_CAPITAL": "核心一级资本",
                "ANC_T1_CAPITAL": "附属一级资本",
            }
            examples = "、".join(
                f"{identities[row_id]['原始公司名称']} {capital_names.get(code, code)}差额 {difference:,.2f} 万元"
                for row_id, code, difference in capital_discrepancies[:3]
            )
            warnings.append(
                f"工作表 {sheet_name} 有 {len(capital_discrepancies)} 组资本明细已披露部分"
                f"与披露合计不一致，相关空白保持未披露，请复核原宽表；示例：{examples}"
            )
        if invalid_values:
            details = "、".join(f"{name} {count} 条" for name, count in invalid_values.items())
            warnings.append(
                f"工作表 {sheet_name} 有 {sum(invalid_values.values())} 个非空值无法解析，已跳过：{details}"
            )

    if unknown_companies:
        warnings.append(f"{len(unknown_companies)} 家公司未在公司主数据中找到，已标记为“未分类”：{'、'.join(sorted(unknown_companies))}")
    logic_frame = pd.DataFrame(logic_rows)
    if not logic_frame.empty and logic_frame["状态"].eq("需复核").any():
        warnings.append("部分外部派生指标与系统计算逻辑不一致，请在确认集成前查看逻辑校验结果。")

    converted_data = (
        pd.concat(output_frames, ignore_index=True)
        if output_frames
        else pd.DataFrame(columns=STANDARD_COLUMNS)
    )
    converted_data = _append_legacy_reversed_ratio(converted_data)
    return ExternalConversionResult(
        data=converted_data,
        sheet_summary=pd.DataFrame(sheet_rows),
        mapping_summary=pd.DataFrame(mapping_rows.values()),
        logic_checks=logic_frame,
        warnings=warnings,
    )


def _standard_row_keys(frame: pd.DataFrame) -> pd.DataFrame:
    public = narrow_table_view(frame).copy()
    normalized = pd.DataFrame(index=public.index)
    for column in NARROW_TABLE_COLUMNS:
        text = public[column].fillna("").astype(str).str.strip()
        if column == "数值":
            numeric = pd.to_numeric(public[column], errors="coerce")
            text = text.where(
                numeric.isna(),
                numeric.map(lambda value: format(float(value), ".15g")),
            )
        normalized[column] = text
    keys = normalized.agg("\x1f".join, axis=1)
    return pd.DataFrame({
        "_标准行键": keys,
        "_重复序号": keys.groupby(keys, sort=False).cumcount(),
    })


def _formula_entity_key(row: pd.Series) -> tuple[str, ...]:
    company_key = str(row.get("公司统一编码", "")).strip()
    if not company_key:
        company_key = str(row.get("标准公司名称", "") or row.get("公司", "")).strip()
    return (
        str(row.get("报告类型", "")).strip(),
        company_key,
        str(row.get("报告年度", "")).replace(".0", "").strip(),
        str(row.get("报告季度", "")).strip(),
        str(row.get("报告期", "")).strip(),
        canonical_period_scope(row.get("期间口径", "")),
    )


def _derived_excel_formula(code: str, references: dict[str, str]) -> str | None:
    """Return the Excel equivalent of ``calculate_derived_values`` for one row."""

    def ref(metric_code: str) -> str:
        return references[metric_code]

    def add(*metric_codes: str) -> str:
        return "+".join(ref(metric_code) for metric_code in metric_codes)

    def add_available(*metric_codes: str) -> str:
        return "+".join(ref(metric_code) for metric_code in metric_codes if metric_code in references)

    def divide(numerator: str, denominator: str) -> str:
        return f'=IFERROR(({numerator})/({denominator}),"")'

    simple_ratios = {
        "ACTUAL_CAPITAL_TO_RECOGNIZED_ASSETS": ("ACTUAL_CAPITAL", "RECOGNIZED_ASSETS"),
        "CORE_T1_TO_ACTUAL_CAPITAL": ("CORE_T1_CAPITAL", "ACTUAL_CAPITAL"),
        "CORE_T2_TO_ACTUAL_CAPITAL": ("CORE_T2_CAPITAL", "ACTUAL_CAPITAL"),
        "ANC_T1_TO_ACTUAL_CAPITAL": ("ANC_T1_CAPITAL", "ACTUAL_CAPITAL"),
        "ANC_T2_TO_ACTUAL_CAPITAL": ("ANC_T2_CAPITAL", "ACTUAL_CAPITAL"),
        "MINIMUM_CAPITAL_TO_RECOGNIZED_LIABILITIES": ("MINIMUM_CAPITAL", "RECOGNIZED_LIABILITIES"),
        "RECOGNIZED_ASSETS_TO_ACTUAL_CAPITAL": ("RECOGNIZED_ASSETS", "ACTUAL_CAPITAL"),
        "RECOGNIZED_ASSETS_TO_MINIMUM_CAPITAL": ("RECOGNIZED_ASSETS", "MINIMUM_CAPITAL"),
        "RECOGNIZED_ASSETS_TO_REGISTERED_CAPITAL": ("RECOGNIZED_ASSETS", "REGISTERED_CAPITAL"),
        "ACTUAL_CAPITAL_TO_REGISTERED_CAPITAL": ("ACTUAL_CAPITAL", "REGISTERED_CAPITAL"),
        "LIFE_INSURANCE_RISK_TO_LIABILITIES": ("INSURANCE_RISK_CAPITAL", "RECOGNIZED_LIABILITIES"),
        "NON_LIFE_INSURANCE_RISK_TO_LIABILITIES": ("NON_LIFE_INSURANCE_RISK_CAPITAL", "RECOGNIZED_LIABILITIES"),
        "LIFE_INSURANCE_RISK_TO_QUANT_CAPITAL": ("INSURANCE_RISK_CAPITAL", "QUANT_RISK_CAPITAL"),
        "NON_LIFE_INSURANCE_RISK_TO_QUANT_CAPITAL": ("NON_LIFE_INSURANCE_RISK_CAPITAL", "QUANT_RISK_CAPITAL"),
        "MARKET_RISK_TO_QUANT_CAPITAL": ("MARKET_RISK_CAPITAL", "QUANT_RISK_CAPITAL"),
        "CREDIT_RISK_TO_QUANT_CAPITAL": ("CREDIT_RISK_CAPITAL", "QUANT_RISK_CAPITAL"),
        "DIVERSIFICATION_EFFECT_TO_QUANT_CAPITAL": ("QUANT_RISK_DIVERSIFICATION_EFFECT", "QUANT_RISK_CAPITAL"),
        "LOSS_ABSORPTION_TO_QUANT_CAPITAL": ("CONTRACT_LOSS_ABSORPTION_EFFECT", "QUANT_RISK_CAPITAL"),
        "MARKET_RISK_TO_ASSETS": ("MARKET_RISK_CAPITAL", "RECOGNIZED_ASSETS"),
        "CREDIT_RISK_TO_ASSETS": ("CREDIT_RISK_CAPITAL", "RECOGNIZED_ASSETS"),
        "CORE_T1_POLICY_SURPLUS_SHARE": ("POLICY_SURPLUS_CORE_T1", "CORE_T1_CAPITAL"),
        "ANC_T1_POLICY_SURPLUS_SHARE": ("POLICY_SURPLUS_ANC_T1", "ANC_T1_CAPITAL"),
        "INTEREST_RATE_RISK_TO_ASSETS": ("INTEREST_RATE_RISK_CAPITAL", "RECOGNIZED_ASSETS"),
        "EQUITY_RISK_TO_ASSETS": ("EQUITY_RISK_CAPITAL", "RECOGNIZED_ASSETS"),
        "SPREAD_RISK_TO_ASSETS": ("SPREAD_RISK_CAPITAL", "RECOGNIZED_ASSETS"),
        "COUNTERPARTY_RISK_TO_ASSETS": ("COUNTERPARTY_RISK_CAPITAL", "RECOGNIZED_ASSETS"),
        "TOTAL_ASSETS_TO_REGISTERED_CAPITAL": ("TOTAL_ASSETS", "REGISTERED_CAPITAL"),
    }
    if code in simple_ratios:
        numerator_code, denominator_code = simple_ratios[code]
        return divide(ref(numerator_code), ref(denominator_code))

    if code == "POLICY_SURPLUS_CORE_TO_ACTUAL_CAPITAL":
        return divide(add_available(*POLICY_SURPLUS_CORE_CODES), ref("ACTUAL_CAPITAL"))
    if code == "POLICY_SURPLUS_CORE_BAND":
        value = ref("POLICY_SURPLUS_CORE_TO_ACTUAL_CAPITAL")
        return (
            f'=IFERROR(IF({value}<=0,"小于等于0%",'
            f'IF({value}<=0.2,"(0,20%]",IF({value}<=0.35,"(20%,35%]","大于35%"))),"")'
        )
    if code == "FEATURE_FACTOR_IMPACT":
        components = add(
            "INSURANCE_RISK_CAPITAL",
            "NON_LIFE_INSURANCE_RISK_CAPITAL",
            "MARKET_RISK_CAPITAL",
            "CREDIT_RISK_CAPITAL",
            "QUANT_RISK_DIVERSIFICATION_EFFECT",
            "CONTRACT_LOSS_ABSORPTION_EFFECT",
        )
        return f'={ref("QUANT_RISK_CAPITAL")}-({components})'
    if code == "POLICY_SURPLUS_CORE_TO_CORE_CAPITAL":
        return divide(
            add_available(*POLICY_SURPLUS_CORE_CODES),
            add("CORE_T1_CAPITAL", "CORE_T2_CAPITAL"),
        )
    if code == "POLICY_SURPLUS_TO_INSURANCE_LIABILITIES":
        return divide(
            add_available(*POLICY_SURPLUS_COMPONENT_CODES),
            ref("INSURANCE_CONTRACT_LIABILITY"),
        )
    if code == "FEATURE_FACTOR_CHECK":
        return divide(
            ref("FEATURE_FACTOR_IMPACT"),
            add(
                "INSURANCE_RISK_CAPITAL",
                "NON_LIFE_INSURANCE_RISK_CAPITAL",
                "MARKET_RISK_CAPITAL",
                "CREDIT_RISK_CAPITAL",
                "QUANT_RISK_DIVERSIFICATION_EFFECT",
                "CONTRACT_LOSS_ABSORPTION_EFFECT",
            ),
        )
    if code == "CORE_CAPITAL_TO_REGISTERED_CAPITAL":
        return divide(add("CORE_T1_CAPITAL", "CORE_T2_CAPITAL"), ref("REGISTERED_CAPITAL"))
    if code == "REGISTERED_CAPITAL_TO_CORE_CAPITAL":
        return divide(ref("REGISTERED_CAPITAL"), add("CORE_T1_CAPITAL", "CORE_T2_CAPITAL"))
    return None


def _formula_plan(
    canonical: pd.DataFrame,
    formula_source: pd.DataFrame,
) -> dict[int, tuple[str, str]]:
    definitions = {definition.code: definition for definition in DERIVED_METRICS}
    source_rows: dict[tuple[tuple[str, ...], str], int] = {}
    for index, row in formula_source.reset_index(drop=True).iterrows():
        if pd.isna(row.get('数值')) or str(row.get('数值', '')).strip() == '':
            continue
        metric_code = str(row.get("指标编码", "")).strip()
        if metric_code in POLICY_SURPLUS_COMPONENT_CODES and str(row.get("来源类型", "")).strip() == "宽表空白推定":
            continue
        key = (_formula_entity_key(row), metric_code)
        source_rows.setdefault(key, int(index) + 2)

    plan: dict[int, tuple[str, str]] = {}
    for index, row in canonical.reset_index(drop=True).iterrows():
        if pd.isna(row.get('数值')) or str(row.get('数值', '')).strip() == '':
            continue
        code = str(row.get("指标编码", "")).strip()
        definition = definitions.get(code)
        if definition is None:
            continue
        entity_key = _formula_entity_key(row)
        references: dict[str, str] = {}
        optional_components = POLICY_SURPLUS_RATIO_COMPONENTS.get(code, ())
        for dependency in definition.dependencies:
            source_row = source_rows.get((entity_key, dependency))
            if source_row is None:
                if dependency in optional_components:
                    continue
                break
            references[dependency] = f"'{FORMULA_SOURCE_SHEET_NAME}'!$I${source_row}"
        else:
            if optional_components and not any(item in references for item in optional_components):
                continue
            formula = _derived_excel_formula(code, references)
            if formula:
                plan[int(index) + 2] = (formula, definition.formula)
    return plan


def _restore_internal_fields(frame: pd.DataFrame, internal: pd.DataFrame) -> pd.DataFrame:
    if frame.empty or internal.empty:
        return frame
    public_keys = _standard_row_keys(frame)
    internal_keys = _standard_row_keys(internal)
    internal_columns = [
        column for column in STANDARD_COLUMNS if column not in NARROW_TABLE_COLUMNS
    ]
    hidden = internal.reindex(columns=internal_columns).copy()
    hidden.columns = [f"_系统_{column}" for column in hidden.columns]
    hidden = pd.concat([internal_keys, hidden], axis=1)
    restored = pd.concat([frame.reset_index(drop=True), public_keys], axis=1).merge(
        hidden,
        on=["_标准行键", "_重复序号"],
        how="left",
        sort=False,
    )
    for column in internal_columns:
        hidden_column = f"_系统_{column}"
        if column not in restored.columns:
            restored[column] = restored[hidden_column]
        else:
            blank = restored[column].fillna("").astype(str).str.strip().eq("")
            restored.loc[blank, column] = restored.loc[blank, hidden_column]
    return restored.drop(
        columns=[
            "_标准行键",
            "_重复序号",
            *[f"_系统_{column}" for column in internal_columns],
        ]
    )


def write_standard_workbook_sheets(
    writer: pd.ExcelWriter,
    frame: pd.DataFrame,
    *,
    formula_source: pd.DataFrame | None = None,
) -> None:
    """Write the compact sheet, auditable formulas and a hidden lossless sheet."""
    canonical = standardize_uploaded_frame(frame)
    source = standardize_uploaded_frame(formula_source if formula_source is not None else canonical)
    formula_plan = _formula_plan(canonical, source)
    narrow_table_view(canonical).to_excel(writer, sheet_name="标准数据", index=False)
    if formula_plan:
        narrow_table_view(source).to_excel(
            writer,
            sheet_name=FORMULA_SOURCE_SHEET_NAME,
            index=False,
        )
        worksheet = writer.book["标准数据"]
        value_column = NARROW_TABLE_COLUMNS.index("数值") + 1
        for row_number, (formula, explanation) in formula_plan.items():
            cell = worksheet.cell(row=row_number, column=value_column)
            cell.value = formula
            cell.font = Font(color="008000")
            cell.comment = Comment(
                f"系统计算逻辑：{explanation}\n公式引用“{FORMULA_SOURCE_SHEET_NAME}”页中的标准化依据。",
                "偿付能力平台",
            )
        writer.book.calculation.calcMode = "auto"
        writer.book.calculation.fullCalcOnLoad = True
        writer.book.calculation.forceFullCalc = True
    canonical.to_excel(
        writer,
        sheet_name=INTERNAL_STANDARD_SHEET_NAME,
        index=False,
    )
    writer.book[INTERNAL_STANDARD_SHEET_NAME].sheet_state = "hidden"


def standard_workbook_bytes(frame: pd.DataFrame) -> bytes:
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        write_standard_workbook_sheets(writer, frame)
    return output.getvalue()


def read_standard_workbook(workbook_bytes: bytes, filename: str) -> pd.DataFrame:
    formula_value_rows: set[int] = set()
    header_row = None
    formula_book = load_workbook(io.BytesIO(workbook_bytes), data_only=False, read_only=True)
    try:
        sheet_name = "标准数据" if "标准数据" in formula_book.sheetnames else formula_book.sheetnames[0]
        formula_sheet = formula_book[sheet_name]
        # ReadOnlyWorksheet.cell() reopens and scans the XML from row 1 on
        # every call. One streaming pass keeps formula detection linear and
        # reuses the detected header for the subsequent cached-value read.
        value_column = None
        for excel_index, row in enumerate(formula_sheet.iter_rows(values_only=True)):
            if header_row is None:
                if excel_index >= 12:
                    break
                values = [str(value).strip() if value is not None else '' for value in row]
                if REQUIRED_STANDARD_HEADERS.issubset(values):
                    header_row = excel_index
                    value_column = values.index('数值')
            else:
                value = row[value_column] if value_column < len(row) else None
                if isinstance(value, str) and value.startswith("="):
                    formula_value_rows.add(excel_index - header_row - 1)
    finally:
        formula_book.close()
    if header_row is None:
        raise ValueError(f"{filename} 未找到标准窄表表头。")
    with pd.ExcelFile(io.BytesIO(workbook_bytes)) as excel:
        frame = pd.read_excel(excel, sheet_name=sheet_name, header=header_row)
        frame.columns = [str(column).strip() for column in frame.columns]
        if INTERNAL_STANDARD_SHEET_NAME in excel.sheet_names:
            internal = pd.read_excel(excel, sheet_name=INTERNAL_STANDARD_SHEET_NAME)
            internal.columns = [str(column).strip() for column in internal.columns]
            if len(internal) == len(frame) and "数值" in frame.columns and "数值" in internal.columns:
                for row_index in formula_value_rows:
                    if row_index >= len(frame):
                        continue
                    visible_value = frame.at[row_index, "数值"]
                    if pd.isna(visible_value) or str(visible_value).strip().startswith("="):
                        frame.at[row_index, "数值"] = internal.at[row_index, "数值"]
            frame = _restore_internal_fields(frame, internal)
    standardized = standardize_uploaded_frame(frame)
    supplied_company_codes = standardized["公司统一编码"].fillna("").astype(str).str.strip()
    result = apply_company_identities(standardized)
    # Codes generated from older full legal names must not split the same
    # insurer from a newer short-name workbook. Preserve custom codes only
    # for companies outside the audited alias/peer master.
    supplied_code_mask = supplied_company_codes.ne("") & ~standardized["公司"].map(
        lambda name: bool(known_company_standard_name(name))
    )
    result.loc[supplied_code_mask, "公司统一编码"] = supplied_company_codes[supplied_code_mask]
    blank_source = result["来源类型"].astype(str).str.strip() == ""
    result.loc[blank_source, "来源类型"] = "标准窄表上传"
    blank_file = result["来源文件"].astype(str).str.strip() == ""
    result.loc[blank_file, "来源文件"] = filename
    blank_sheet = result["来源工作表"].astype(str).str.strip() == ""
    result.loc[blank_sheet, "来源工作表"] = sheet_name
    blank_batch = result["导入批次"].astype(str).str.strip() == ""
    result.loc[blank_batch, "导入批次"] = Path(filename).stem
    return result
