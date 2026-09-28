from __future__ import annotations

import io
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

from .solvency_validator import validation_status_summary


VALIDATION_SHEET_NAMES = (
    "状态汇总",
    "分类汇总",
    "检查明细",
    "待处理事项",
    "勾稽规则",
)

RESULT_LABELS = {
    "check_type": "检查类别",
    "severity": "未通过时级别",
    "rule_id": "规则ID",
    "rule_name": "规则名称",
    "company": "公司",
    "report_period": "报告期",
    "period": "期间组",
    "actual": "披露值/实际数",
    "expected": "系统计算值/应有数",
    "difference": "差异",
    "tolerance": "容差",
    "status": "状态",
    "notes": "检查说明",
    "involved_metrics": "涉及指标编码",
    "source_pages": "来源页码",
    "suggestion": "处理建议",
}

FORMULA_DESCRIPTIONS = {
    "SCHEMA_REQUIRED_FIELDS": "标准窄表必填字段空白数 = 0",
    "SCHEMA_KNOWN_METRIC": "未在正式指标字典登记的指标编码数 = 0",
    "SCHEMA_METRIC_NAME": "指标名称与正式指标字典不一致的记录数 = 0",
    "SCHEMA_UNIT_MATCH": "单位与正式指标字典不一致的记录数 = 0",
    "SCHEMA_DATA_TYPE": "数据类型与正式指标字典不一致的记录数 = 0",
    "SCHEMA_ALLOWED_PERIOD": "期间口径不在指标字典允许范围内的记录数 = 0",
    "SCHEMA_NUMERIC_VALUE": "金额、百分比、数量和倍数指标不可解析的记录数 = 0",
    "SCHEMA_EXACT_DUPLICATE": "同一来源、同一指标、同一期间、同一数值的重复记录数 = 0",
    "SCHEMA_CONFLICTING_DUPLICATE": "同一来源、同一指标、同一期间存在多个数值的冲突组数 = 0",
    "ACTUAL_CAPITAL_COMPONENTS": "实际资本 = 核心一级资本 + 核心二级资本 + 附属一级资本 + 附属二级资本",
    "ACTUAL_CAPITAL_BALANCE": "实际资本 = 认可资产 - 认可负债",
    "CORE_SURPLUS": "核心偿付能力溢额 = 核心一级资本 + 核心二级资本 - 最低资本",
    "COMBINED_SURPLUS": "综合偿付能力溢额 = 实际资本 - 最低资本",
    "CORE_RATIO": "核心偿付能力充足率 = (核心一级资本 + 核心二级资本) / 最低资本 × 100",
    "COMBINED_RATIO": "综合偿付能力充足率 = 实际资本 / 最低资本 × 100",
    "MINIMUM_CAPITAL": "最低资本 = 量化风险最低资本 + 控制风险最低资本 + 附加资本",
    "DUPLICATE_ACTUAL_CAPITAL": "同一期间不同来源披露的实际资本最大值 = 最小值",
    "DUPLICATE_MINIMUM_CAPITAL": "同一期间不同来源披露的最低资本最大值 = 最小值",
    "DUPLICATE_RECOGNIZED_ASSETS": "同一期间不同来源披露的认可资产最大值 = 最小值",
}


def validation_formula_description(rule_id: Any, involved_metrics: Any = "") -> str:
    rule = str(rule_id or "").strip()
    if rule.startswith("REQUIRED_CORE_METRIC:"):
        code = rule.split(":", 1)[1]
        return f"COUNT({code} 在当前期末组的有效数值) ≥ 1"
    if rule in FORMULA_DESCRIPTIONS:
        return FORMULA_DESCRIPTIONS[rule]
    metrics = str(involved_metrics or "").strip()
    return f"按规则 {rule} 复核{('：' + metrics) if metrics else ''}"


def validation_display_frame(results: pd.DataFrame) -> pd.DataFrame:
    """Hide rule severity unless the current check actually failed."""
    display = results.copy()
    if "severity" in display.columns and "status" in display.columns:
        failed = display["status"].fillna("").astype(str).eq("未通过")
        display["severity"] = display["severity"].where(failed, "")
    return display


def _write_frame(workbook, sheet_name: str, frame: pd.DataFrame):
    if sheet_name in workbook.sheetnames:
        del workbook[sheet_name]
    worksheet = workbook.create_sheet(sheet_name)
    for row in dataframe_to_rows(frame, index=False, header=True):
        worksheet.append(list(row))
    return worksheet


def _style_sheet(worksheet) -> None:
    if worksheet.max_row < 1 or worksheet.max_column < 1:
        return
    header_fill = PatternFill("solid", fgColor="00338D")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_view.showGridLines = False
    worksheet.row_dimensions[1].height = 32

    for column_cells in worksheet.iter_cols():
        values = [str(cell.value or "") for cell in column_cells[:200]]
        width = min(max(max((len(value) for value in values), default=0) + 2, 10), 42)
        worksheet.column_dimensions[column_cells[0].column_letter].width = width
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _add_detail_formulas(worksheet) -> None:
    headers = {str(cell.value): cell.column for cell in worksheet[1]}
    actual_column = headers.get("披露值/实际数")
    expected_column = headers.get("系统计算值/应有数")
    tolerance_column = headers.get("容差")
    excel_difference_column = headers.get("Excel差异")
    threshold_column = headers.get("Excel阈值复核")
    if not all(
        (
            actual_column,
            expected_column,
            tolerance_column,
            excel_difference_column,
            threshold_column,
        )
    ):
        return

    actual_letter = get_column_letter(actual_column)
    expected_letter = get_column_letter(expected_column)
    tolerance_letter = get_column_letter(tolerance_column)
    difference_letter = get_column_letter(excel_difference_column)
    threshold_letter = get_column_letter(threshold_column)
    for row_number in range(2, worksheet.max_row + 1):
        worksheet.cell(row_number, excel_difference_column).value = (
            f'=IF(OR(ISBLANK({actual_letter}{row_number}),'
            f'ISBLANK({expected_letter}{row_number})),"",'
            f'{actual_letter}{row_number}-{expected_letter}{row_number})'
        )
        worksheet.cell(row_number, threshold_column).value = (
            f'=IF(ISBLANK({difference_letter}{row_number}),"",'
            f'IF(ABS({difference_letter}{row_number})<={tolerance_letter}{row_number},'
            '"容差内","超出容差"))'
        )
        worksheet.cell(row_number, excel_difference_column).font = Font(color="000000")
        worksheet.cell(row_number, threshold_column).font = Font(color="000000")

    if worksheet.max_row >= 2:
        green_fill = PatternFill("solid", fgColor="C6EFCE")
        red_fill = PatternFill("solid", fgColor="FFC7CE")
        target_range = f"{threshold_letter}2:{threshold_letter}{worksheet.max_row}"
        worksheet.conditional_formatting.add(
            target_range,
            FormulaRule(formula=[f'{threshold_letter}2="容差内"'], fill=green_fill),
        )
        worksheet.conditional_formatting.add(
            target_range,
            FormulaRule(formula=[f'{threshold_letter}2="超出容差"'], fill=red_fill),
        )


def _validation_detail_frame(results: pd.DataFrame) -> pd.DataFrame:
    details = validation_display_frame(results)
    details.insert(
        min(details.columns.get_loc("rule_name") + 1, len(details.columns)),
        "formula_description",
        [
            validation_formula_description(rule_id, metrics)
            for rule_id, metrics in zip(
                details.get("rule_id", pd.Series(dtype=str)),
                details.get("involved_metrics", pd.Series(dtype=str)),
            )
        ],
    )
    details["Excel差异"] = ""
    details["Excel阈值复核"] = ""
    return details.rename(
        columns={
            **RESULT_LABELS,
            "formula_description": "检查公式",
        }
    )


def validation_workbook_bytes(
    results: pd.DataFrame,
    base_workbook_bytes: bytes | None = None,
) -> bytes:
    """Continue from STEP3 and append an auditable STEP4 review workbook."""
    if base_workbook_bytes:
        workbook = load_workbook(io.BytesIO(base_workbook_bytes), data_only=False)
    else:
        workbook = Workbook()
        workbook.active.title = "说明"
        workbook["说明"]["A1"] = "未提供 STEP3 工作簿，仅导出 STEP4 勾稽结果。"

    for sheet_name in VALIDATION_SHEET_NAMES:
        if sheet_name in workbook.sheetnames:
            del workbook[sheet_name]

    summary = validation_status_summary(results)
    category_summary = (
        results.groupby(["check_type", "status"], dropna=False)
        .size()
        .rename("数量")
        .reset_index()
        .rename(columns={"check_type": "检查类别", "status": "状态"})
    )
    details = _validation_detail_frame(results)
    issues = details[
        ~details["状态"].astype(str).isin(["通过", "不适用"])
    ].copy()
    rule_catalog = _validation_detail_frame(results.copy())
    if "severity" in results.columns:
        rule_catalog["未通过时级别"] = results["severity"].fillna("").astype(str).values
    rules = (
        rule_catalog[["规则ID", "规则名称", "检查公式", "检查类别", "未通过时级别", "容差"]]
        .drop_duplicates(subset=["规则ID", "规则名称", "检查公式"], keep="first")
        .reset_index(drop=True)
    )

    worksheets = [
        _write_frame(workbook, "状态汇总", summary),
        _write_frame(workbook, "分类汇总", category_summary),
        _write_frame(workbook, "检查明细", details),
        _write_frame(workbook, "待处理事项", issues),
        _write_frame(workbook, "勾稽规则", rules),
    ]
    for worksheet in worksheets:
        _style_sheet(worksheet)
    _add_detail_formulas(workbook["检查明细"])
    _add_detail_formulas(workbook["待处理事项"])

    status_fills = {
        "通过": PatternFill("solid", fgColor="C6EFCE"),
        "未通过": PatternFill("solid", fgColor="FFC7CE"),
        "缺失": PatternFill("solid", fgColor="FFC7CE"),
        "需复核": PatternFill("solid", fgColor="FFEB9C"),
        "执行错误": PatternFill("solid", fgColor="FFC7CE"),
        "不适用": PatternFill("solid", fgColor="E7E6E6"),
    }
    for sheet_name in ("检查明细", "待处理事项"):
        worksheet = workbook[sheet_name]
        headers = {str(cell.value): cell.column for cell in worksheet[1]}
        status_column = headers.get("状态")
        if status_column:
            for row_number in range(2, worksheet.max_row + 1):
                cell = worksheet.cell(row_number, status_column)
                if str(cell.value) in status_fills:
                    cell.fill = status_fills[str(cell.value)]

    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
