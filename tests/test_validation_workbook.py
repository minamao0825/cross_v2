from __future__ import annotations

import io
import unittest

import pandas as pd
from openpyxl import load_workbook

from services.solvency_dataset_adapter import standard_workbook_bytes
from services.solvency_validation_workbook import (
    validation_display_frame,
    validation_workbook_bytes,
)


class ValidationWorkbookTests(unittest.TestCase):
    def test_validation_display_frame_only_shows_severity_for_failed_checks(self):
        results = pd.DataFrame(
            {
                "status": ["通过", "未通过", "缺失", "需复核", "执行错误", "不适用"],
                "severity": ["错误", "警告", "错误", "警告", "错误", "警告"],
            }
        )

        display = validation_display_frame(results)

        self.assertEqual(display["severity"].tolist(), ["", "警告", "", "", "", ""])

    def test_step4_export_continues_from_step3_and_adds_formulas(self):
        standard_data = pd.DataFrame([
            {
                "公司": "测试人寿",
                "报告期": "2026Q1",
                "指标编码": "ACTUAL_CAPITAL",
                "指标名称": "实际资本",
                "期间口径": "本季度末数",
                "数值": 100.0,
                "单位": "万元",
                "数据类型": "金额",
                "来源类型": "报告提取",
                "指标属性": "披露",
            }
        ])
        results = pd.DataFrame([
            {
                "check_type": "公式勾稽",
                "severity": "错误",
                "rule_id": "ACTUAL_CAPITAL_BALANCE",
                "rule_name": "实际资本平衡",
                "company": "测试人寿",
                "report_period": "2026Q1",
                "period": "当前期末",
                "actual": 100.0,
                "expected": 100.0,
                "difference": 0.0,
                "tolerance": 1.0,
                "status": "通过",
                "notes": "检查通过",
                "involved_metrics": "ACTUAL_CAPITAL、RECOGNIZED_ASSETS、RECOGNIZED_LIABILITIES",
                "source_pages": "1",
                "suggestion": "",
            }
        ])

        payload = validation_workbook_bytes(
            results,
            standard_workbook_bytes(standard_data),
        )
        workbook = load_workbook(io.BytesIO(payload), data_only=False)
        detail = workbook["检查明细"]
        headers = {cell.value: cell.column for cell in detail[1]}

        self.assertIn("标准数据", workbook.sheetnames)
        self.assertIn("勾稽规则", workbook.sheetnames)
        self.assertIn("未通过时级别", headers)
        self.assertIsNone(detail.cell(2, headers["未通过时级别"]).value)
        self.assertEqual(
            detail.cell(2, headers["检查公式"]).value,
            "实际资本 = 认可资产 - 认可负债",
        )
        self.assertTrue(str(detail.cell(2, headers["Excel差异"]).value).startswith("=IF("))
        self.assertTrue(str(detail.cell(2, headers["Excel阈值复核"]).value).startswith("=IF("))

        rule_sheet = workbook["勾稽规则"]
        rule_headers = {cell.value: cell.column for cell in rule_sheet[1]}
        self.assertIn("未通过时级别", rule_headers)
        self.assertEqual(rule_sheet.cell(2, rule_headers["未通过时级别"]).value, "错误")


if __name__ == "__main__":
    unittest.main()
