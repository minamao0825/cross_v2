from __future__ import annotations

import io
import unittest

import openpyxl

from services.solvency_ai_table_extractor import (
    AIExtractionBundle,
    PageGrid,
    extract_unit_records,
    reconstructed_workbook_bytes,
    unit_records_from_payload,
    unit_recovery_needed,
)
from services.solvency_table_extractor import ExtractedTable


class UnitExtractionTests(unittest.TestCase):
    def setUp(self):
        self.rows = [
            [
                "\u884c\u6b21",
                "\u9879\u76ee",
                "\u671f\u672b\u6570\uff08\u5143\uff09",
                "\u671f\u521d\u6570",
            ],
            ["1", "\u6838\u5fc3\u4e00\u7ea7\u8d44\u672c\uff08\u5143\uff09", "100", "90"],
            ["2", "\u6838\u5fc3\u4e8c\u7ea7\u8d44\u672c", "80", "70"],
            [
                "3",
                "\u6838\u5fc3\u507f\u4ed8\u80fd\u529b\u5145\u8db3\u7387",
                "130%",
                "120%",
            ],
        ]
        self.grid = PageGrid(
            12,
            "\n".join([
                "000|\u5b9e\u9645\u8d44\u672c\u8868",
                "001|\u5355\u4f4d\uff1a\u4eba\u6c11\u5e01\u5143",
                "002|\u884c\u6b21 \u9879\u76ee \u671f\u672b\u6570 \u671f\u521d\u6570",
                "003|\u6838\u5fc3\u4e00\u7ea7\u8d44\u672c 100 90",
                "004|\u5b9e\u9645\u8d44\u672c\u5408\u8ba1 180 160",
            ]),
            30,
        )

    def _table(self) -> ExtractedTable:
        units = extract_unit_records("ACTUAL_CAPITAL", self.rows, [self.grid])
        return ExtractedTable(
            "ACTUAL_CAPITAL",
            "\u5b9e\u9645\u8d44\u672c\u8868",
            12,
            1,
            self.rows,
            source_pages=[12],
            unit_records=units,
        )

    def test_extracts_table_column_row_and_percent_units(self):
        units = self._table().unit_records
        found = {(item.scope, item.normalized_unit) for item in units}
        self.assertIn(("\u8868\u7ea7", "\u5143"), found)
        self.assertIn(("\u5217\u7ea7", "\u5143"), found)
        self.assertIn(("\u884c\u7ea7", "\u5143"), found)
        self.assertTrue(
            any(
                item.normalized_unit == "%"
                and item.target == "\u6838\u5fc3\u507f\u4ed8\u80fd\u529b\u5145\u8db3\u7387"
                for item in units
            )
        )

    def test_unit_footer_does_not_modify_data_rows(self):
        table = self._table()
        self.assertEqual(len(table.to_frame()), len(self.rows))
        preview = table.to_frame(include_unit_footer=True)
        self.assertEqual(len(preview), len(self.rows) + 1)
        self.assertEqual(preview.iloc[-1, 0], "\u3010\u5355\u4f4d\u5907\u6ce8\u3011")

    def test_workbook_contains_unit_sheet_and_footer(self):
        table = self._table()
        payload = reconstructed_workbook_bytes(AIExtractionBundle([table], []))
        workbook = openpyxl.load_workbook(io.BytesIO(payload), data_only=True)
        self.assertIn("\u5355\u4f4d\u4fe1\u606f", workbook.sheetnames)
        table_sheet = workbook["\u5b9e\u9645\u8d44\u672c\u8868_P12"]
        self.assertEqual(
            table_sheet.cell(table_sheet.max_row, 1).value,
            "\u3010\u5355\u4f4d\u5907\u6ce8\u3011",
        )

    def test_scanned_payload_maps_original_label_unit_to_normalized_row(self):
        rows = [
            ["\u884c\u6b21", "\u9879\u76ee", "\u672c\u5b63\u5ea6\u672b\u6570", "\u4e0a\u5b63\u5ea6\u672b\u6570"],
            [
                "5",
                "\u5b9e\u9645\u8d44\u672c\u5408\u8ba1",
                "3,593,086,798.99",
                "3,435,399,358.42",
            ],
        ]
        payload = {
            "unit_records": [{
                "scope": "\u884c\u7ea7",
                "target": "\u5b9e\u9645\u8d44\u672c\uff08\u5143\uff09",
                "raw_unit": "\u5143",
                "source_page": 8,
                "source_text": "\u5b9e\u9645\u8d44\u672c\uff08\u5143\uff09",
                "confidence": "\u9ad8",
            }],
        }

        records = unit_records_from_payload(payload, rows, [8])

        self.assertEqual(len(records), 1)
        self.assertEqual(records[0].target, "\u5b9e\u9645\u8d44\u672c\u5408\u8ba1")
        self.assertEqual(records[0].normalized_unit, "\u5143")
        self.assertEqual(records[0].source_page, 8)
        self.assertFalse(unit_recovery_needed(rows, records))

    def test_scanned_numeric_row_without_unit_requests_recovery(self):
        rows = [
            ["\u884c\u6b21", "\u9879\u76ee", "\u672c\u5b63\u5ea6\u672b\u6570", "\u4e0a\u5b63\u5ea6\u672b\u6570"],
            [
                "5",
                "\u5b9e\u9645\u8d44\u672c\u5408\u8ba1",
                "3,593,086,798.99",
                "3,435,399,358.42",
            ],
        ]

        self.assertTrue(unit_recovery_needed(rows, []))

    def test_actual_capital_unit_is_read_from_table_title(self):
        rows = [
            ["行次", "项目", "期末数", "期初数"],
            ["1", "核心一级资本", "6,974,645,244.61", "7,945,195,913.43"],
            ["5", "实际资本合计", "8,100,000,000.00", "8,000,000,000.00"],
        ]
        grid = PageGrid(
            20,
            "\n".join([
                "000|S02-实际资本（元）",
                "001|行次 项目 期末数 期初数",
                "002|1 核心一级资本 6,974,645,244.61 7,945,195,913.43",
                "003|5 实际资本合计 8,100,000,000.00 8,000,000,000.00",
            ]),
            20,
        )

        records = extract_unit_records("ACTUAL_CAPITAL", rows, [grid])

        self.assertTrue(any(
            record.scope == "表级"
            and record.normalized_unit == "元"
            and record.source_text == "S02-实际资本（元）"
            for record in records
        ))
        self.assertFalse(unit_recovery_needed(rows, records))

    def test_mixed_units_are_read_separately_from_operating_title(self):
        rows = [
            ["指标名称", "本季度数", "本年度累计数"],
            ["（一）保险业务收入", "10,103,443,945.92", "10,103,443,945.92"],
            ["（九）投资收益率", "0.53", "0.53"],
            ["（十）综合投资收益率", "0.36", "0.36"],
        ]
        grid = PageGrid(
            11,
            "\n".join([
                "000|（四）其他经营指标（元，%）",
                "001|指标名称 本季度数 本年度累计数",
                "002|（一）保险业务收入 10,103,443,945.92 10,103,443,945.92",
                "003|（九）投资收益率 0.53 0.53",
                "004|（十）综合投资收益率 0.36 0.36",
            ]),
            26,
        )

        records = extract_unit_records("OPERATING_METRICS", rows, [grid])
        title_units = {
            record.normalized_unit
            for record in records
            if record.scope == "表级"
            and record.source_text == "（四）其他经营指标（元，%）"
        }

        self.assertEqual(title_units, {"元", "%"})
        self.assertFalse(unit_recovery_needed(rows, records))


if __name__ == "__main__":
    unittest.main()

