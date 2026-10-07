from __future__ import annotations

import io
import unittest
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

from services.solvency_dataset_adapter import (
    read_standard_workbook,
    supported_metric_catalog,
)
from services.solvency_normalizer import NARROW_TABLE_COLUMNS, STANDARD_COLUMNS
from services.solvency_metric_registry import extend_taxonomy
from services.solvency_step3_standardizer import (
    infer_step3_session_metadata,
    infer_step3_upload_metadata,
    normalize_report_period,
    read_target_template,
    result_workbook_bytes,
    standardize_to_target,
    step3_metric_catalog,
    target_template_workbook_bytes,
)
from services.solvency_table_extractor import ExtractedTable


class Step3StandardizerTests(unittest.TestCase):
    def setUp(self):
        self.taxonomy = pd.DataFrame([
            {
                "指标编码": "RECOGNIZED_ASSETS",
                "指标名称": "认可资产",
                "别名": "认可资产合计",
                "一级模块": "主要指标",
                "二级模块": "偿付能力",
                "标准单位": "万元",
                "数据类型": "金额",
            },
            {
                "指标编码": "ACTUAL_CAPITAL",
                "指标名称": "实际资本",
                "别名": "实际资本合计",
                "一级模块": "主要指标",
                "二级模块": "偿付能力",
                "标准单位": "万元",
                "数据类型": "金额",
            },
            {
                "指标编码": "STEP3_ONLY_DISCLOSURE",
                "指标名称": "STEP3新增披露指标",
                "别名": "",
                "一级模块": "实际资本",
                "二级模块": "补充披露",
                "标准单位": "万元",
                "数据类型": "金额",
                "STEP5宽表映射": "否",
            },
        ]).fillna("")

    @staticmethod
    def _workbook(frame: pd.DataFrame, sheet_name: str = "指标清单") -> bytes:
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            frame.to_excel(writer, sheet_name=sheet_name, index=False)
        return output.getvalue()

    def test_step2_session_metadata_falls_back_to_pdf_filename(self):
        detected = infer_step3_session_metadata(
            {"公司": "", "报告期": "", "披露日期": "2026-05-01"},
            "中银三星2026Q1偿付能力季度报告摘要.pdf",
            ({"company_name": "中银三星人寿保险有限公司", "company_type": "寿险"},),
            peer_group_map={"中银三星": "银行系"},
            default_peer_group="其他",
        )
        self.assertEqual(detected.company, "中银三星人寿保险有限公司")
        self.assertEqual(detected.company_type, "寿险")
        self.assertEqual(detected.report_period, "2026Q1")
        self.assertEqual(detected.company_source, "PDF文件名")
        self.assertEqual(detected.period_source, "PDF文件名")
        self.assertEqual(detected.peer_group, "银行系")
        self.assertEqual(detected.disclosure_date, "2026-05-01")
        self.assertEqual(detected.warnings, ())

    def test_step2_session_metadata_can_use_peer_short_name(self):
        detected = infer_step3_session_metadata(
            {},
            "中银三星2025Q3偿付能力季度报告摘要.pdf",
            (),
            peer_group_map={"中银三星": "银行系"},
        )
        self.assertEqual(detected.company, "中银三星")
        self.assertEqual(detected.report_period, "2025Q3")
        self.assertEqual(detected.peer_group, "银行系")

    def test_step2_session_metadata_keeps_pdf_identity_and_warns_on_conflict(self):
        detected = infer_step3_session_metadata(
            {"公司": "中银三星人寿保险有限公司", "报告期": "2025Q3"},
            "招商信诺2026Q1偿付能力季度报告摘要.pdf",
            (
                {"company_name": "中银三星", "company_type": "寿险"},
                {"company_name": "招商信诺", "company_type": "寿险"},
            ),
        )
        self.assertEqual(detected.company, "中银三星")
        self.assertEqual(detected.report_period, "2025Q3")
        self.assertEqual(len(detected.warnings), 2)

    def test_step2_session_metadata_prompts_for_unidentified_company(self):
        detected = infer_step3_session_metadata(
            {}, "偿付能力报告.pdf", ({"company_name": "中银三星", "company_type": "寿险"},)
        )
        self.assertEqual(detected.company, "")
        self.assertEqual(detected.report_period, "")
        self.assertEqual(len(detected.warnings), 2)

    def test_step5_wide_catalog_is_a_subset_of_step3_catalog(self):
        step5_catalog = supported_metric_catalog(
            self.taxonomy,
            include_derived=True,
        )
        step3_catalog = step3_metric_catalog(
            self.taxonomy,
            include_derived=True,
        )

        step5_codes = set(step5_catalog["指标编码"])
        step3_codes = set(step3_catalog["指标编码"])
        self.assertTrue(step5_codes < step3_codes)
        self.assertNotIn("STEP3_ONLY_DISCLOSURE", step5_codes)
        self.assertIn("STEP3_ONLY_DISCLOSURE", step3_codes)
        self.assertEqual(step3_catalog["指标编码"].nunique(), len(step3_catalog))
        self.assertIn("ACTUAL_CAPITAL_TO_RECOGNIZED_ASSETS", set(step3_catalog["指标编码"]))

    def test_upload_metadata_is_inferred_from_unit_source_text(self):
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            pd.DataFrame({
                "来源原文": ["公司名称：爱心人寿保险股份有限公司 2026-03-31 单位：万元"]
            }).to_excel(writer, sheet_name="单位信息", index=False)
            pd.DataFrame([["项目", "期末数"], ["核心偿付能力充足率", "156%"]]).to_excel(
                writer, sheet_name="偿付能力充足率指标_P1", index=False, header=False
            )

        detected = infer_step3_upload_metadata(
            output.getvalue(),
            "偿付能力报告_STEP2标准化提取.xlsx",
            ({"company_name": "爱心人寿", "company_type": "寿险"},),
            peer_group_map={"爱心人寿": "小型公司"},
            default_peer_group="其他",
        )

        self.assertEqual(detected.company, "爱心人寿")
        self.assertEqual(detected.company_type, "寿险")
        self.assertEqual(detected.report_period, "2026Q1")
        self.assertEqual(detected.disclosure_date, "2026-03-31")
        self.assertEqual(detected.peer_group, "小型公司")
        self.assertEqual(detected.company_source, "工作簿内容")
        self.assertEqual(detected.period_source, "工作簿日期")

    def test_embedded_metadata_takes_priority_over_filename(self):
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            pd.DataFrame([
                {"字段": "公司", "值": "爱心人寿"},
                {"字段": "报告期", "值": "2026Q1"},
            ]).to_excel(writer, sheet_name="报告元信息", index=False)

        detected = infer_step3_upload_metadata(
            output.getvalue(),
            "国民养老_2025Q4_STEP2.xlsx",
            (
                {"company_name": "爱心人寿", "company_type": "寿险"},
                {"company_name": "国民养老", "company_type": "养老险"},
            ),
        )

        self.assertEqual(detected.company, "爱心人寿")
        self.assertEqual(detected.report_period, "2026Q1")
        self.assertEqual(len(detected.warnings), 2)

    def test_matching_pdf_and_upload_filenames_correct_stale_embedded_period(self):
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            pd.DataFrame([
                {"字段": "公司", "值": "中银三星人寿保险有限公司"},
                {"字段": "报告期", "值": "2005Q1"},
                {"字段": "披露日期", "值": "2005-5-26"},
                {"字段": "来源文件", "值": "中银三星2025Q1偿付能力季度报告摘要.pdf"},
            ]).to_excel(writer, sheet_name="报告元信息", index=False)

        detected = infer_step3_upload_metadata(
            output.getvalue(),
            "偿付能力报告_STEP2结构化提取_中银三星2025Q1偿付能力季度报告摘要.xlsx",
            ({"company_name": "中银三星", "company_type": "寿险"},),
        )
        self.assertEqual(detected.report_period, "2025Q1")
        self.assertEqual(detected.period_source, "原PDF与上传文件名一致")
        self.assertEqual(detected.disclosure_date, "")

    def test_filename_is_used_as_metadata_fallback(self):
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            pd.DataFrame([["项目", "本期"]]).to_excel(
                writer, sheet_name="数据", index=False, header=False
            )
        detected = infer_step3_upload_metadata(
            output.getvalue(),
            "国民养老_2025Q4_STEP2.xlsx",
            ({"company_name": "国民养老", "company_type": "养老险"},),
        )
        self.assertEqual(detected.company, "国民养老")
        self.assertEqual(detected.company_source, "文件名")
        self.assertEqual(detected.report_period, "2025Q4")
        self.assertEqual(detected.period_source, "文件名")

    def test_normalize_report_period_does_not_read_year_digit_as_quarter(self):
        self.assertEqual(normalize_report_period("2026Q1"), (2026, "Q1", "2026Q1"))
        self.assertEqual(normalize_report_period("2026年第二季度"), (2026, "Q2", "2026Q2"))
        self.assertEqual(normalize_report_period("2026年一季报"), (2026, "Q1", "2026Q1"))
        self.assertEqual(normalize_report_period("2026-1Q"), (2026, "Q1", "2026Q1"))
        self.assertEqual(normalize_report_period("Q1_2026"), (2026, "Q1", "2026Q1"))

    def test_embedded_year_and_quarter_are_combined(self):
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            pd.DataFrame([
                {"字段": "公司全称", "值": "长生人寿保险有限公司"},
                {"字段": "年度", "值": "2026"},
                {"字段": "季度", "值": "第一季度"},
            ]).to_excel(writer, sheet_name="报告信息", index=False)

        detected = infer_step3_upload_metadata(
            output.getvalue(),
            "偿付能力报告_STEP2结构化提取.xlsx",
            ({"company_name": "长生人寿", "company_type": "寿险"},),
        )

        self.assertEqual(detected.company, "长生人寿")
        self.assertEqual(detected.report_period, "2026Q1")
        self.assertEqual(detected.period_source, "工作簿元信息")

    def test_original_pdf_filename_embedded_in_workbook_is_used(self):
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            pd.DataFrame([
                {"字段": "来源文件", "值": "长生人寿2026年一季报偿付能力报告.pdf"},
            ]).to_excel(writer, sheet_name="元信息", index=False)

        detected = infer_step3_upload_metadata(
            output.getvalue(),
            "偿付能力报告_STEP2结构化提取.xlsx",
            ({"company_name": "长生人寿", "company_type": "寿险"},),
        )

        self.assertEqual(detected.company, "长生人寿")
        self.assertEqual(detected.company_source, "原PDF文件名")
        self.assertEqual(detected.report_period, "2026Q1")
        self.assertEqual(detected.period_source, "原PDF文件名")

    def test_people_insurance_pension_legal_name_maps_to_short_name_and_peer_group(self):
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            pd.DataFrame([
                {"字段": "公司", "值": "中国人民养老保险有限责任公司"},
                {"字段": "报告期", "值": "2026Q1"},
                {"字段": "来源文件", "值": "人保养老2026Q1偿付能力季度报告摘要.pdf"},
            ]).to_excel(writer, sheet_name="报告元信息", index=False)

        detected = infer_step3_upload_metadata(
            output.getvalue(),
            "偿付能力报告_STEP2结构化提取_人保养老.xlsx",
            ({"company_name": "人保养老", "company_type": "养老险"},),
            peer_group_map={"人保养老": "养老健康"},
            default_peer_group="其他",
        )

        self.assertEqual(detected.company, "人保养老")
        self.assertEqual(detected.company_type, "养老险")
        self.assertEqual(detected.peer_group, "养老健康")
        self.assertEqual(detected.company_source, "工作簿元信息")

    def test_quarter_end_date_beats_later_disclosure_date(self):
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            pd.DataFrame({
                "来源原文": [
                    "截至2025年12月31日 上季度末数；截至2026年3月31日 本季度末数",
                    "披露日期：2026年4月29日",
                ]
            }).to_excel(writer, sheet_name="单位信息", index=False)

        detected = infer_step3_upload_metadata(
            output.getvalue(),
            "长生人寿_STEP2.xlsx",
            ({"company_name": "长生人寿", "company_type": "寿险"},),
        )

        self.assertEqual(detected.report_period, "2026Q1")
        self.assertEqual(detected.period_source, "工作簿日期")

    def test_matching_current_pdf_context_fills_legacy_workbook_period(self):
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            pd.DataFrame({"指标名称": ["实际资本"], "标准数值": [40302]}).to_excel(
                writer, sheet_name="VLM_v2指标结果", index=False,
            )

        detected = infer_step3_upload_metadata(
            output.getvalue(),
            "偿付能力报告_STEP2结构化提取_长生人寿.xlsx",
            (
                {"company_name": "长生人寿", "company_type": "寿险"},
                {"company_name": "平安养老", "company_type": "养老险"},
            ),
            context_metadata={"公司": "长生人寿保险有限公司", "报告期": "2026Q1"},
            context_filename="长生人寿2026Q1偿付能力季度报告摘要.pdf",
        )

        self.assertEqual(detected.company, "长生人寿")
        self.assertEqual(detected.report_period, "2026Q1")
        self.assertEqual(detected.period_source, "当前PDF会话（公司一致）")

    def test_mismatched_current_pdf_context_is_not_reused(self):
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            pd.DataFrame({"指标名称": ["实际资本"], "标准数值": [40302]}).to_excel(
                writer, sheet_name="VLM_v2指标结果", index=False,
            )

        detected = infer_step3_upload_metadata(
            output.getvalue(),
            "偿付能力报告_STEP2结构化提取_长生人寿.xlsx",
            (
                {"company_name": "长生人寿", "company_type": "寿险"},
                {"company_name": "平安养老", "company_type": "养老险"},
            ),
            context_metadata={"公司": "平安养老", "报告期": "2026Q1"},
            context_filename="平安养老2026Q1偿付能力季度报告摘要.pdf",
        )

        self.assertEqual(detected.company, "长生人寿")
        self.assertEqual(detected.report_period, "")

    def test_cross_capital_details_are_available_in_step5_and_other_details_remain_step3_only(self):
        taxonomy = extend_taxonomy(self.taxonomy)
        step5_catalog = supported_metric_catalog(taxonomy, include_derived=False)
        step3_catalog = step3_metric_catalog(taxonomy, include_derived=False)
        cross_detail_codes = {
            "CORE_T1_INVESTMENT_PROPERTY_FAIR_VALUE_ADJUSTMENT",
            "ANC_T1_DEFERRED_TAX_ASSET",
        }
        step3_only_codes = {
            "CORE_T2_EXCESS_DEDUCTION",
            "ANC_T2_EXCESS_DEDUCTION",
        }

        self.assertTrue(cross_detail_codes | step3_only_codes <= set(step3_catalog["指标编码"]))
        self.assertTrue(cross_detail_codes <= set(step5_catalog["指标编码"]))
        self.assertTrue(step3_only_codes.isdisjoint(set(step5_catalog["指标编码"])))

    def test_custom_target_template_is_filtered_by_enable_flag(self):
        supported = step3_metric_catalog(self.taxonomy, include_derived=True)
        custom = pd.DataFrame([
            {"启用": "是", "指标编码": "ACTUAL_CAPITAL"},
            {"启用": "否", "指标编码": "RECOGNIZED_ASSETS"},
        ])

        selected, warnings = read_target_template(
            self._workbook(custom),
            "custom.xlsx",
            supported,
        )

        self.assertEqual(selected["指标编码"].tolist(), ["ACTUAL_CAPITAL"])
        self.assertEqual(warnings, ())

    def test_custom_target_rejects_code_outside_step3_dictionary(self):
        supported = step3_metric_catalog(self.taxonomy, include_derived=True)
        custom = pd.DataFrame([{"启用": "是", "指标编码": "OUT_OF_SCOPE"}])

        with self.assertRaisesRegex(ValueError, "STEP3 正式指标字典之外"):
            read_target_template(
                self._workbook(custom),
                "custom.xlsx",
                supported,
            )

    def test_target_driven_standardization_calculates_dependencies_but_outputs_target_only(self):
        table = ExtractedTable(
            table_id="TEST",
            table_name="测试表",
            page=1,
            table_index=1,
            rows=[
                ["项目", "期末数（万元）"],
                ["认可资产", "200"],
                ["实际资本", "100"],
            ],
            source_pages=[1],
        )
        catalog = step3_metric_catalog(self.taxonomy, include_derived=True)
        target = catalog[
            catalog["指标编码"] == "ACTUAL_CAPITAL_TO_RECOGNIZED_ASSETS"
        ].copy()

        result = standardize_to_target(
            [table],
            self.taxonomy,
            {
                "公司": "测试人寿保险有限公司",
                "报告年度": 2026,
                "报告季度": "Q1",
                "报告期": "2026Q1",
                "来源文件": "测试报告.pdf",
                "导入批次": "测试报告:2026Q1",
            },
            "寿险",
            target,
            report_profile_id="LIFE_SOLVENCY",
            allowed_company_types=("寿险", "健康险", "养老险"),
            peer_group_map={"测试人寿": "小型公司"},
            default_peer_group="其他",
            include_derived=True,
        )

        self.assertEqual(list(result.data.columns), STANDARD_COLUMNS)
        self.assertEqual(
            result.data["指标编码"].tolist(),
            ["ACTUAL_CAPITAL_TO_RECOGNIZED_ASSETS"],
        )
        self.assertAlmostEqual(float(result.data.iloc[0]["数值"]), 0.5)
        self.assertEqual(result.data.iloc[0]["同业分类"], "小型公司")
        self.assertEqual(result.target_summary.iloc[0]["填报状态"], "已填报")

    def test_policy_surplus_ratios_keep_missing_and_zero_semantics(self):
        taxonomy = extend_taxonomy(self.taxonomy)
        catalog = step3_metric_catalog(taxonomy, include_derived=True)
        target = catalog[catalog["指标编码"].isin({
            "POLICY_SURPLUS_CORE_TO_CORE_CAPITAL",
            "POLICY_SURPLUS_TO_INSURANCE_LIABILITIES",
        })].copy()
        records = [
            {"指标编码": "POLICY_SURPLUS_CORE_T1", "状态": "found", "标准数值": 20, "原始值": "20", "期间口径": "本季度末数"},
            {"指标编码": "POLICY_SURPLUS_CORE_T2", "状态": "not_disclosed", "标准数值": None, "原始值": "", "期间口径": "本季度末数"},
            {"指标编码": "POLICY_SURPLUS_ANC_T1", "状态": "found", "标准数值": 5, "原始值": "5", "期间口径": "本季度末数"},
            {"指标编码": "POLICY_SURPLUS_ANC_T2", "状态": "not_disclosed", "标准数值": None, "原始值": "", "期间口径": "本季度末数"},
            {"指标编码": "CORE_T1_CAPITAL", "状态": "found", "标准数值": 100, "原始值": "100", "期间口径": "本季度末数"},
            {"指标编码": "CORE_T2_CAPITAL", "状态": "found", "标准数值": 20, "原始值": "20", "期间口径": "本季度末数"},
            {"指标编码": "INSURANCE_CONTRACT_LIABILITY", "状态": "found", "标准数值": 250, "原始值": "250", "期间口径": "本季度末数"},
        ]
        table = ExtractedTable("TEST", "测试表", 1, 1, [], metric_records=records)
        metadata = {
            "公司": "测试人寿保险有限公司",
            "报告年度": 2026,
            "报告季度": "Q1",
            "报告期": "2026Q1",
        }
        missing = standardize_to_target(
            [table], taxonomy, metadata, "寿险", target,
            report_profile_id="LIFE_SOLVENCY",
            allowed_company_types=("寿险", "健康险", "养老险"),
            include_derived=True,
        )
        partial = missing.data.set_index("指标编码")
        self.assertAlmostEqual(partial.loc["POLICY_SURPLUS_CORE_TO_CORE_CAPITAL", "数值"], 20 / 120)
        self.assertAlmostEqual(partial.loc["POLICY_SURPLUS_TO_INSURANCE_LIABILITIES", "数值"], 25 / 250)
        self.assertTrue(partial["披露状态"].eq("已计算").all())
        exported = load_workbook(io.BytesIO(result_workbook_bytes(missing)), data_only=False)
        data_rows = list(exported["标准数据"].iter_rows(values_only=True))
        code_column = data_rows[0].index("指标编码")
        value_column = data_rows[0].index("数值")
        ratio_row = next(row for row in data_rows[1:] if row[code_column] == "POLICY_SURPLUS_TO_INSURANCE_LIABILITIES")
        self.assertTrue(str(ratio_row[value_column]).startswith("=IFERROR("))
        self.assertEqual(read_standard_workbook(result_workbook_bytes(missing), "partial.xlsx").loc[
            lambda frame: frame["指标编码"].eq("POLICY_SURPLUS_TO_INSURANCE_LIABILITIES"), "数值"
        ].iloc[0], 25 / 250)

        for record in records:
            if record["指标编码"] in {"POLICY_SURPLUS_CORE_T2", "POLICY_SURPLUS_ANC_T2"}:
                record.update({"状态": "disclosed_zero", "标准数值": 0, "原始值": "0"})
        calculated = standardize_to_target(
            [ExtractedTable("TEST", "测试表", 1, 1, [], metric_records=records)],
            taxonomy, metadata, "寿险", target,
            report_profile_id="LIFE_SOLVENCY",
            allowed_company_types=("寿险", "健康险", "养老险"),
            include_derived=True,
        )
        values = calculated.data.set_index("指标编码")["数值"].to_dict()
        self.assertAlmostEqual(values["POLICY_SURPLUS_CORE_TO_CORE_CAPITAL"], 20 / 120)
        self.assertAlmostEqual(values["POLICY_SURPLUS_TO_INSURANCE_LIABILITIES"], 25 / 250)

    def test_balance_sheet_ratios_bridge_current_quarter_and_closing_periods(self):
        taxonomy = extend_taxonomy(self.taxonomy)
        catalog = step3_metric_catalog(taxonomy, include_derived=True)
        target_codes = {
            "TOTAL_ASSETS_TO_REGISTERED_CAPITAL",
            "POLICY_SURPLUS_TO_INSURANCE_LIABILITIES",
        }
        target = catalog[catalog["指标编码"].isin(target_codes)].copy()
        records = [
            {"指标编码": "TOTAL_ASSETS", "状态": "found", "标准数值": 500, "原始值": "500", "期间口径": "本季度数"},
            {"指标编码": "REGISTERED_CAPITAL", "状态": "found", "标准数值": 50, "原始值": "50", "期间口径": "本季度末数"},
            {"指标编码": "POLICY_SURPLUS_CORE_T1", "状态": "found", "标准数值": 15, "原始值": "15", "期间口径": "本季度末数"},
            {"指标编码": "POLICY_SURPLUS_ANC_T1", "状态": "found", "标准数值": 5, "原始值": "5", "期间口径": "本季度末数"},
            {"指标编码": "INSURANCE_CONTRACT_LIABILITY", "状态": "found", "标准数值": 250, "原始值": "250", "期间口径": "本季度数"},
        ]
        result = standardize_to_target(
            [ExtractedTable("TEST", "测试表", 1, 1, [], metric_records=records)],
            taxonomy,
            {"公司": "测试人寿保险有限公司", "报告年度": 2026, "报告季度": "Q2", "报告期": "2026Q2"},
            "寿险",
            target,
            report_profile_id="LIFE_SOLVENCY",
            allowed_company_types=("寿险", "健康险", "养老险"),
            include_derived=True,
        )
        rows = result.data.set_index("指标编码")
        self.assertAlmostEqual(rows.loc["TOTAL_ASSETS_TO_REGISTERED_CAPITAL", "数值"], 10.0)
        self.assertAlmostEqual(rows.loc["POLICY_SURPLUS_TO_INSURANCE_LIABILITIES", "数值"], 0.08)
        self.assertTrue(rows.loc[list(target_codes), "披露状态"].eq("已计算").all())

        workbook = load_workbook(io.BytesIO(result_workbook_bytes(result)), data_only=False)
        data_rows = list(workbook["标准数据"].iter_rows(values_only=True))
        source_rows = list(workbook["计算依据"].iter_rows(values_only=True))
        data_code_column = data_rows[0].index("指标编码")
        data_value_column = data_rows[0].index("数值")
        source_code_column = source_rows[0].index("指标编码")
        source_period_column = source_rows[0].index("期间口径")
        source_row_by_code_period = {
            (row[source_code_column], row[source_period_column]): index
            for index, row in enumerate(source_rows[1:], start=2)
        }
        formula_by_code = {
            row[data_code_column]: row[data_value_column]
            for row in data_rows[1:]
        }
        total_assets_row = source_row_by_code_period[("TOTAL_ASSETS", "本季度数")]
        liability_row = source_row_by_code_period[("INSURANCE_CONTRACT_LIABILITY", "本季度数")]
        self.assertIn(f"$I${total_assets_row}", formula_by_code["TOTAL_ASSETS_TO_REGISTERED_CAPITAL"])
        self.assertIn(f"$I${liability_row}", formula_by_code["POLICY_SURPLUS_TO_INSURANCE_LIABILITIES"])

    def test_supplied_pdf_corrects_legacy_invented_policy_surplus_zeros(self):
        taxonomy = extend_taxonomy(self.taxonomy)
        catalog = step3_metric_catalog(taxonomy, include_derived=True)
        codes = {
            "POLICY_SURPLUS_CORE_T1", "POLICY_SURPLUS_CORE_T2",
            "POLICY_SURPLUS_ANC_T1", "POLICY_SURPLUS_ANC_T2",
            "INSURANCE_CONTRACT_LIABILITY", "POLICY_SURPLUS_TO_INSURANCE_LIABILITIES",
        }
        target = catalog[catalog["指标编码"].isin(codes)].copy()
        records = [
            {"指标编码": "POLICY_SURPLUS_CORE_T1", "状态": "found", "标准数值": 20, "原始值": "20", "期间口径": "本季度末数"},
            {"指标编码": "POLICY_SURPLUS_CORE_T2", "状态": "disclosed_zero", "标准数值": 0, "原始值": "-", "期间口径": "本季度末数"},
            {"指标编码": "POLICY_SURPLUS_ANC_T1", "状态": "disclosed_zero", "标准数值": 0, "原始值": "-", "期间口径": "本季度末数"},
            {"指标编码": "POLICY_SURPLUS_ANC_T2", "状态": "disclosed_zero", "标准数值": 0, "原始值": "-", "期间口径": "本季度末数"},
            {"指标编码": "INSURANCE_CONTRACT_LIABILITY", "状态": "found", "标准数值": 250, "原始值": "250", "期间口径": "2025年第一季度"},
        ]
        metadata = {"公司": "中银三星人寿保险有限公司", "报告年度": 2025, "报告季度": "Q1", "报告期": "2025Q1"}
        with patch("services.solvency_step3_standardizer.extract_report_metadata", return_value=metadata), patch(
            "services.solvency_step3_standardizer.visible_policy_surplus_codes",
            return_value=frozenset({"POLICY_SURPLUS_CORE_T1"}),
        ):
            result = standardize_to_target(
                [ExtractedTable("TEST", "测试表", 18, 1, [], metric_records=records)],
                taxonomy, metadata, "寿险", target,
                report_profile_id="LIFE_SOLVENCY",
                allowed_company_types=("寿险", "健康险", "养老险"),
                include_derived=True,
                source_pdf_bytes=b"pdf",
            )
        rows = result.data.set_index("指标编码")
        for code in ("POLICY_SURPLUS_CORE_T2", "POLICY_SURPLUS_ANC_T1", "POLICY_SURPLUS_ANC_T2"):
            self.assertEqual(rows.loc[code, "披露状态"], "未披露")
        self.assertEqual(rows.loc["POLICY_SURPLUS_TO_INSURANCE_LIABILITIES", "披露状态"], "已计算")
        self.assertAlmostEqual(rows.loc["POLICY_SURPLUS_TO_INSURANCE_LIABILITIES", "数值"], 20 / 250)

        with patch("services.solvency_step3_standardizer.extract_report_metadata", return_value=metadata), patch(
            "services.solvency_step3_standardizer.visible_policy_surplus_codes",
            return_value=frozenset(codes - {"INSURANCE_CONTRACT_LIABILITY", "POLICY_SURPLUS_TO_INSURANCE_LIABILITIES"}),
        ):
            explicitly_zero = standardize_to_target(
                [ExtractedTable("TEST", "测试表", 18, 1, [], metric_records=records)],
                taxonomy, metadata, "寿险", target,
                report_profile_id="LIFE_SOLVENCY",
                allowed_company_types=("寿险", "健康险", "养老险"),
                include_derived=True,
                source_pdf_bytes=b"pdf",
            )
        self.assertAlmostEqual(
            explicitly_zero.data.set_index("指标编码").loc["POLICY_SURPLUS_TO_INSURANCE_LIABILITIES", "数值"],
            20 / 250,
        )

    def test_registered_capital_dependency_is_preserved_as_standalone_metric(self):
        taxonomy = extend_taxonomy(self.taxonomy)
        catalog = step3_metric_catalog(taxonomy, include_derived=True)
        target = catalog[
            catalog["指标编码"].eq("ACTUAL_CAPITAL_TO_REGISTERED_CAPITAL")
        ].copy()
        records = [
            {"指标编码": "ACTUAL_CAPITAL", "状态": "found", "标准数值": 100, "原始值": "100", "期间口径": "本季度末数"},
            {"指标编码": "REGISTERED_CAPITAL", "状态": "found", "标准数值": 50, "原始值": "50", "期间口径": "本季度末数"},
        ]
        result = standardize_to_target(
            [ExtractedTable("TEST", "测试表", 1, 1, [], metric_records=records)],
            taxonomy,
            {"公司": "测试人寿保险有限公司", "报告年度": 2026, "报告季度": "Q1", "报告期": "2026Q1"},
            "寿险",
            target,
            report_profile_id="LIFE_SOLVENCY",
            allowed_company_types=("寿险", "健康险", "养老险"),
            include_derived=True,
        )
        values = result.data.set_index("指标编码")["数值"].to_dict()
        self.assertEqual(values["REGISTERED_CAPITAL"], 50.0)
        self.assertEqual(values["ACTUAL_CAPITAL_TO_REGISTERED_CAPITAL"], 2.0)

    def test_current_quarter_period_labels_are_unified_and_duplicate_metrics_removed(self):
        taxonomy = extend_taxonomy(self.taxonomy)
        catalog = step3_metric_catalog(taxonomy, include_derived=True)
        selected_codes = {
            "POLICY_SURPLUS_CORE_T1",
            "POLICY_SURPLUS_CORE_T2",
            "POLICY_SURPLUS_ANC_T1",
            "POLICY_SURPLUS_ANC_T2",
            "CORE_T1_CAPITAL",
            "CORE_T2_CAPITAL",
            "INSURANCE_CONTRACT_LIABILITY",
            "POLICY_SURPLUS_CORE_TO_CORE_CAPITAL",
            "POLICY_SURPLUS_TO_INSURANCE_LIABILITIES",
        }
        target = catalog[catalog["指标编码"].isin(selected_codes)].copy()
        records = [
            {"指标编码": "POLICY_SURPLUS_CORE_T1", "状态": "not_disclosed", "标准数值": None, "原始值": "", "期间口径": "本季度末"},
            {"指标编码": "POLICY_SURPLUS_CORE_T1", "状态": "found", "标准数值": 20, "原始值": "20", "期间口径": "2026年第1季度"},
            {"指标编码": "POLICY_SURPLUS_CORE_T2", "状态": "found", "标准数值": 10, "原始值": "10", "期间口径": "2026年第一季度"},
            {"指标编码": "POLICY_SURPLUS_ANC_T1", "状态": "found", "标准数值": 5, "原始值": "5", "期间口径": "本季度"},
            {"指标编码": "POLICY_SURPLUS_ANC_T2", "状态": "disclosed_zero", "标准数值": 0, "原始值": "0", "期间口径": "2026Q1"},
            {"指标编码": "CORE_T1_CAPITAL", "状态": "found", "标准数值": 100, "原始值": "100", "期间口径": "2026-3-31"},
            {"指标编码": "CORE_T2_CAPITAL", "状态": "found", "标准数值": 20, "原始值": "20", "期间口径": "本季度末数"},
            {"指标编码": "INSURANCE_CONTRACT_LIABILITY", "状态": "found", "标准数值": 250, "原始值": "250", "期间口径": "2026.3.31"},
        ]
        result = standardize_to_target(
            [ExtractedTable("TEST", "测试表", 1, 1, [], metric_records=records)],
            taxonomy,
            {"公司": "测试人寿保险有限公司", "报告年度": 2026, "报告季度": "Q1", "报告期": "2026Q1"},
            "寿险",
            target,
            report_profile_id="LIFE_SOLVENCY",
            allowed_company_types=("寿险", "健康险", "养老险"),
            include_derived=True,
        )

        self.assertTrue(result.data["期间口径"].eq("本季度末数").all())
        self.assertEqual(
            int(result.data["指标编码"].eq("POLICY_SURPLUS_CORE_T1").sum()),
            1,
        )
        self.assertEqual(
            int(result.formula_source_data["指标编码"].eq("POLICY_SURPLUS_CORE_T1").sum()),
            1,
        )
        values = result.data.set_index("指标编码")["数值"].to_dict()
        self.assertAlmostEqual(values["POLICY_SURPLUS_CORE_TO_CORE_CAPITAL"], 30 / 120)
        self.assertAlmostEqual(values["POLICY_SURPLUS_TO_INSURANCE_LIABILITIES"], 35 / 250)

    def test_derived_metric_is_exported_as_excel_formula_and_round_trips(self):
        table = ExtractedTable(
            table_id="TEST",
            table_name="测试表",
            page=1,
            table_index=1,
            rows=[
                ["项目", "期末数（万元）"],
                ["认可资产", "200"],
                ["实际资本", "100"],
            ],
            source_pages=[1],
        )
        catalog = step3_metric_catalog(self.taxonomy, include_derived=True)
        target = catalog[
            catalog["指标编码"] == "ACTUAL_CAPITAL_TO_RECOGNIZED_ASSETS"
        ].copy()
        result = standardize_to_target(
            [table],
            self.taxonomy,
            {
                "公司": "测试人寿保险有限公司",
                "报告年度": 2026,
                "报告季度": "Q1",
                "报告期": "2026Q1",
            },
            "寿险",
            target,
            report_profile_id="LIFE_SOLVENCY",
            allowed_company_types=("寿险", "健康险", "养老险"),
            include_derived=True,
        )

        workbook_bytes = result_workbook_bytes(result)
        workbook = load_workbook(io.BytesIO(workbook_bytes), data_only=False)
        formula = workbook["标准数据"]["I2"].value
        restored = read_standard_workbook(workbook_bytes, "formula_step3.xlsx")

        self.assertIn("计算依据", workbook.sheetnames)
        self.assertTrue(str(formula).startswith("=IFERROR("))
        self.assertIn("'计算依据'!$I$", str(formula))
        self.assertAlmostEqual(float(restored.iloc[0]["数值"]), 0.5)

    def test_generated_workbooks_are_readable_by_step5(self):
        catalog = step3_metric_catalog(self.taxonomy, include_derived=False)
        template_bytes = target_template_workbook_bytes(catalog, "测试偿付能力报告")
        template_standard = pd.read_excel(
            io.BytesIO(template_bytes),
            sheet_name="标准数据",
        )
        self.assertEqual(list(template_standard.columns), NARROW_TABLE_COLUMNS)
        selected, _ = read_target_template(template_bytes, "target.xlsx", catalog)
        self.assertEqual(set(selected["指标编码"]), set(catalog["指标编码"]))

        table = ExtractedTable(
            table_id="TEST",
            table_name="测试表",
            page=1,
            table_index=1,
            rows=[["项目", "期末数（万元）"], ["实际资本", "100"]],
            source_pages=[1],
        )
        result = standardize_to_target(
            [table],
            self.taxonomy,
            {
                "公司": "测试人寿保险有限公司",
                "报告年度": 2026,
                "报告季度": "Q1",
                "报告期": "2026Q1",
            },
            "寿险",
            catalog[catalog["指标编码"] == "ACTUAL_CAPITAL"],
            report_profile_id="LIFE_SOLVENCY",
            allowed_company_types=("寿险", "健康险", "养老险"),
            include_derived=False,
        )
        workbook = result_workbook_bytes(result)
        exported = pd.read_excel(io.BytesIO(workbook), sheet_name="标准数据")
        step5_data = read_standard_workbook(workbook, "step3.xlsx")

        self.assertEqual(list(exported.columns), NARROW_TABLE_COLUMNS)
        self.assertEqual(list(step5_data.columns), STANDARD_COLUMNS)
        self.assertEqual(step5_data["指标编码"].tolist(), ["ACTUAL_CAPITAL"])
        self.assertTrue(step5_data.iloc[0]["公司统一编码"])
        self.assertEqual(step5_data.iloc[0]["公司类型"], "寿险")


if __name__ == "__main__":
    unittest.main()
