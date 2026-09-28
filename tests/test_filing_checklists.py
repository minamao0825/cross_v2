from __future__ import annotations

import io
import unittest
from collections import Counter
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from services.solvency_filing_catalog import filing_entries, filing_codes, filing_row_code
from services.solvency_metric_registry import extend_taxonomy
from services.solvency_step3_standardizer import (
    step3_metric_catalog, standardize_to_target, result_workbook_bytes,
    target_template_workbook_bytes,
)
from services.solvency_table_extractor import ExtractedTable
from services.solvency_dataset_adapter import read_standard_workbook
from services.solvency_normalizer import load_taxonomy, narrow_table_view, standardize_uploaded_frame
from services.solvency_filing_catalog import extend_filing_taxonomy, THREE_YEAR_TARGETS
from services.solvency_vlm_v2_pipeline import build_vlm_v2_target_cards


class FilingChecklistTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.taxonomy = extend_taxonomy(load_taxonomy(
            Path(__file__).parents[1] / 'config' / 'solvency_taxonomy.xlsx'))
        cls.catalog = step3_metric_catalog(cls.taxonomy, include_derived=True, use_checklists=True)

    def standardize(self, records, codes=None, derived=False):
        target = self.catalog if codes is None else self.catalog[self.catalog['指标编码'].isin(codes)]
        table = ExtractedTable(table_id='OPERATING_METRICS', table_name='主要经营指标',
                               page=11, table_index=1, rows=[], metric_records=records)
        return standardize_to_target([table], self.taxonomy, {
            '公司': '人保养老', '报告期': '2026Q1', '报告年度': 2026, '报告季度': 'Q1',
        }, '养老险', target, report_profile_id='LIFE_SOLVENCY',
            allowed_company_types=('寿险', '养老险', '健康险'),
            peer_group_map={'人保养老': '养老健康'}, include_derived=derived)

    @staticmethod
    def record(code, status, value=None, raw=''):
        return {'指标编码': code, '状态': status, '标准数值': value, '原始值': raw,
                '期间口径': '本季度末数', '物理页码': 11, '证据原文': '来源证据'}

    def test_all_five_lists_are_coded_and_section_headings_excluded(self):
        self.assertEqual(Counter(e['table_id'] for e in filing_entries()), {
            'SOLVENCY_MAIN': 15, 'ACTUAL_CAPITAL': 30, 'MINIMUM_CAPITAL': 34,
            'OPERATING_METRICS': 29, 'RECOGNIZED_ASSETS': 49,
        })
        self.assertEqual(len(filing_codes()), 147)
        self.assertEqual(len(self.catalog), 184)
        self.assertTrue(filing_codes() <= set(self.catalog['指标编码']))
        self.assertFalse(self.catalog['指标编码'].duplicated().any())
        source = self.catalog[self.catalog['指标编码'].isin(filing_codes())]
        self.assertTrue(source['填报规则'].str.len().gt(0).all())

    def test_homonyms_and_capital_components_do_not_match_totals(self):
        self.assertEqual(filing_row_code('OPERATING_METRICS', '净资产'), 'NET_ASSETS')
        self.assertEqual(filing_row_code('ACTUAL_CAPITAL', '净资产'), 'FINANCIAL_STATEMENT_NET_ASSETS')
        self.assertEqual(filing_row_code('MINIMUM_CAPITAL', '其他附加资本'), 'OTHER_ADDITIONAL_CAPITAL')
        self.assertEqual(filing_row_code('OPERATING_METRICS', '13个月续保率'), 'RENEWAL_RATE_13M')
        # Repeated deduction labels require row context, never an arbitrary first match.
        deductions = [e for e in filing_entries() if e['name'] == '减：超限额应扣除的部分']
        self.assertGreater(len(deductions), 1)
        self.assertEqual(filing_row_code('ACTUAL_CAPITAL', '减：超限额应扣除的部分'), '')
        for entry in deductions:
            self.assertEqual(filing_row_code('ACTUAL_CAPITAL', entry['name'], entry['row_number']), entry['code'])

    def test_extraction_cards_cover_exactly_the_checklists_with_row_context(self):
        base = load_taxonomy(Path(__file__).parents[1] / 'config' / 'solvency_taxonomy.xlsx')
        taxonomy = extend_filing_taxonomy(base)
        ids = list(dict.fromkeys(e['table_id'] for e in filing_entries()))
        cards = build_vlm_v2_target_cards([{'table_id': key, 'table_name': key} for key in ids], taxonomy)
        for card in cards:
            self.assertEqual({metric['metric_id'] for metric in card.metrics}, filing_codes(card.table_id))
        capital = next(card for card in cards if card.table_id == 'ACTUAL_CAPITAL')
        self.assertTrue(all(metric['source_row_number'] for metric in capital.metrics))

    def test_missing_na_and_zero_remain_distinct_in_narrow_output(self):
        result = self.standardize([
            self.record('NET_ASSETS', 'found', 502251.77, '502,251.77'),
            self.record('ADDITIONAL_CAPITAL', 'disclosed_zero', 0, '-'),
            self.record('INSURANCE_REVENUE', 'disclosed_na', raw='<不适用>'),
            self.record('AGENT_COUNT', 'not_disclosed'),
        ], ['NET_ASSETS', 'ADDITIONAL_CAPITAL', 'INSURANCE_REVENUE', 'AGENT_COUNT', 'SIGNED_PREMIUM'])
        rows = result.data.set_index('指标编码')
        self.assertEqual(len(rows), 5)
        self.assertEqual(rows.loc['ADDITIONAL_CAPITAL', '数值'], 0)
        self.assertEqual(rows.loc['ADDITIONAL_CAPITAL', '原始披露值'], '-')
        self.assertEqual(rows.loc['INSURANCE_REVENUE', '披露状态'], '不适用')
        for code in ['AGENT_COUNT', 'SIGNED_PREMIUM']:
            self.assertEqual(rows.loc[code, '披露状态'], '未披露')
            self.assertTrue(pd.isna(rows.loc[code, '数值']))
        self.assertTrue(result.data['公司'].eq('人保养老').all())
        self.assertTrue(result.data['同业分类'].eq('养老健康').all())
        restored = read_standard_workbook(result_workbook_bytes(result), 'test.xlsx').set_index('指标编码')
        self.assertEqual(restored.loc['SIGNED_PREMIUM', '披露状态'], '未披露')
        self.assertTrue(pd.isna(restored.loc['SIGNED_PREMIUM', '数值']))
        preview = narrow_table_view(result.data).set_index('指标编码')
        for code in ['AGENT_COUNT', 'SIGNED_PREMIUM']:
            self.assertEqual(preview.loc[code, '数值'], '未披露')
        workbook = load_workbook(io.BytesIO(result_workbook_bytes(result)))
        sheet = workbook['标准数据']
        headers = [cell.value for cell in sheet[1]]
        cells = {row[headers.index('指标编码')].value: row[headers.index('数值')]
                 for row in list(sheet.rows)[1:]}
        self.assertEqual(cells['SIGNED_PREMIUM'].value, '未披露')
        self.assertEqual(cells['AGENT_COUNT'].value, '未披露')
        self.assertEqual(cells['ADDITIONAL_CAPITAL'].value, 0)
        self.assertEqual(cells['ADDITIONAL_CAPITAL'].data_type, 'n')
        self.assertIsNone(cells['INSURANCE_REVENUE'].value)

    def test_old_canonical_dash_becomes_zero_but_absent_row_does_not(self):
        result = self.standardize([
            self.record('POLICY_SURPLUS_CORE_T2', 'disclosed_na', raw='—'),
            self.record('NET_ASSETS', 'found', 0, '0.00'),
            self.record('AGENT_COUNT', 'not_disclosed', raw='-'),
        ], ['POLICY_SURPLUS_CORE_T2', 'NET_ASSETS', 'AGENT_COUNT'])
        rows = result.data.set_index('指标编码')
        for code in ['POLICY_SURPLUS_CORE_T2', 'NET_ASSETS']:
            self.assertEqual(rows.loc[code, '数值'], 0)
            self.assertEqual(rows.loc[code, '披露状态'], '已披露为0')
        self.assertEqual(rows.loc['POLICY_SURPLUS_CORE_T2', '原始披露值'], '—')
        self.assertEqual(narrow_table_view(result.data).set_index('指标编码').loc['AGENT_COUNT', '数值'], '未披露')

    def test_literal_missing_marker_is_not_a_numeric_input(self):
        frame = standardize_uploaded_frame(pd.DataFrame([
            {'指标编码': 'NET_ASSETS', '数值': '未披露'},
            {'指标编码': 'ACTUAL_CAPITAL', '数值': 0},
        ]))
        self.assertTrue(pd.isna(frame.iloc[0]['数值']))
        self.assertEqual(frame.iloc[0]['披露状态'], '未披露')
        self.assertEqual(frame.iloc[1]['数值'], 0)

    def test_no_extracted_values_still_outputs_every_enabled_target(self):
        result = self.standardize([])
        self.assertEqual(set(result.data['指标编码']), set(self.catalog['指标编码']))
        self.assertTrue(result.data['数值'].isna().all())
        self.assertTrue(result.data[result.data['指标编码'].isin(filing_codes())]['披露状态'].eq('未披露').all())

    def test_text_product_list_is_retained(self):
        result = self.standardize([self.record('TOP_FIVE_PRODUCTS', 'found', raw='产品甲；产品乙')], ['TOP_FIVE_PRODUCTS'])
        self.assertEqual(result.data.iloc[0]['数值'], '产品甲；产品乙')
        self.assertEqual(result.data.iloc[0]['数据类型'], '文本')

    def test_missing_dependency_does_not_become_zero_via_excel_formula(self):
        code = 'ACTUAL_CAPITAL_TO_RECOGNIZED_ASSETS'
        result = self.standardize([
            self.record('ACTUAL_CAPITAL', 'not_disclosed'),
            self.record('RECOGNIZED_ASSETS', 'found', 200, '200'),
        ], [code], derived=True)
        self.assertTrue(pd.isna(result.data.iloc[0]['数值']))
        self.assertEqual(result.data.iloc[0]['披露状态'], '无法计算')
        workbook = load_workbook(io.BytesIO(result_workbook_bytes(result)), data_only=False)
        self.assertIsNone(workbook['标准数据']['I2'].value)

    def test_downloadable_target_contains_codes_rules_and_disclosure_column(self):
        book = pd.ExcelFile(io.BytesIO(target_template_workbook_bytes(self.catalog, '偿付能力报告')))
        targets = pd.read_excel(book, sheet_name='指标清单')
        self.assertEqual(len(targets), 184)
        self.assertIn('填报规则', targets.columns)
        self.assertIn('披露状态', pd.read_excel(book, sheet_name='标准数据').columns)

    def test_three_year_targets_are_individually_named_and_period_scoped(self):
        data = target_template_workbook_bytes(self.catalog, '偿付能力报告')
        target = pd.read_excel(io.BytesIO(data), sheet_name='指标清单').set_index('指标编码')
        for code, name in THREE_YEAR_TARGETS.values():
            self.assertEqual(target.loc[code, '指标名称'], name)
            self.assertEqual(target.loc[code, '期间口径'], '本季度')
            self.assertEqual(target.loc[code, '目标表ID'], 'THREE_YEAR_INVESTMENT_RETURN')
            self.assertEqual(target.loc[code, '指标属性'], '披露')

    def test_old_step2_three_year_records_do_not_overwrite_quarterly_returns(self):
        records = []
        expected = {}
        for index, (old, (new, _)) in enumerate(THREE_YEAR_TARGETS.items()):
            quarterly = self.record(old, 'found', 1.2 + index, str(1.2 + index))
            average = self.record(old, 'found', 4.5 + index, str(4.5 + index))
            average['指标语义键'] = old + ':THREE_YEAR_AVERAGE'
            average['期间口径'] = '近三年平均'
            records.extend([quarterly, average])
            expected[old] = 1.2 + index
            expected[new] = 4.5 + index
        result = self.standardize(records, list(expected))
        rows = result.data.set_index('指标编码')
        self.assertEqual(len(rows), 4)
        for code, value in expected.items():
            self.assertAlmostEqual(rows.loc[code, '数值'], value)
        self.assertTrue(result.target_summary['填报状态'].eq('已填报').all())
        restored = read_standard_workbook(result_workbook_bytes(result), 'three_year.xlsx')
        self.assertEqual(set(restored['指标编码']), set(expected))

    def test_only_three_year_targets_can_read_legacy_step2_codes(self):
        new_codes = [new for new, _ in THREE_YEAR_TARGETS.values()]
        records = []
        for old in THREE_YEAR_TARGETS:
            row = self.record(old, 'found', 3.25, '3.25%')
            row['期间口径'] = '近三年平均'
            records.append(row)
        result = self.standardize(records, new_codes)
        self.assertEqual(len(result.data), 2)
        self.assertTrue(result.data['数值'].eq(3.25).all())
        self.assertTrue(result.data['期间口径'].eq('本季度').all())

    def test_missing_three_year_return_is_not_filled_from_current_quarter(self):
        new_codes = [new for new, _ in THREE_YEAR_TARGETS.values()]
        result = self.standardize([
            self.record(old, 'found', 8, '8%') for old in THREE_YEAR_TARGETS
        ], new_codes)
        self.assertEqual(len(result.data), 2)
        self.assertTrue(result.data['披露状态'].eq('未披露').all())
        self.assertTrue(result.data['数值'].isna().all())
        self.assertTrue(result.data['期间口径'].eq('本季度').all())

    def test_legacy_per_table_three_year_workbook_rows_are_supported(self):
        table = ExtractedTable('THREE_YEAR_INVESTMENT_RETURN', '近三年（综合）投资收益率', 12, 1,
                               [['项目', '近三年平均'], ['投资收益率', '3.5%'], ['综合投资收益率', '4.5%']])
        target = self.catalog[self.catalog['指标编码'].isin([new for new, _ in THREE_YEAR_TARGETS.values()])]
        result = standardize_to_target([table], self.taxonomy, {
            '公司': '人保养老', '报告期': '2026Q1', '报告年度': 2026, '报告季度': 'Q1',
        }, '养老险', target, report_profile_id='LIFE_SOLVENCY',
            allowed_company_types=('寿险', '养老险', '健康险'), include_derived=False)
        rows = result.data.set_index('指标编码')
        self.assertEqual(rows.loc['THREE_YEAR_AVG_INVESTMENT_RETURN', '数值'], 3.5)
        self.assertEqual(rows.loc['THREE_YEAR_AVG_COMPREHENSIVE_INVESTMENT_RETURN', '数值'], 4.5)


if __name__ == '__main__':
    unittest.main()
