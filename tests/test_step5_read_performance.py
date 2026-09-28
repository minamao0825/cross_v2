import ast
import io
import unittest
from pathlib import Path
from unittest.mock import patch

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

from services.solvency_dataset_adapter import read_standard_workbook, standard_workbook_bytes


def standard_fixture(rows=180, company='测试公司'):
    records = [dict(公司=company, 报告期='2026Q1', 指标编码=code, 指标名称=name,
                    数值=value, 单位=unit, 数据类型='金额', 期间口径='本季度末数')
               for code, name, value, unit in (
                   ('ACTUAL_CAPITAL', '实际资本', 100, '万元'),
                   ('RECOGNIZED_ASSETS', '认可资产', 500, '万元'),
                   ('ACTUAL_CAPITAL_TO_RECOGNIZED_ASSETS', '实际资本/认可资产', .2, '倍'),
               )]
    records.extend(dict(公司=company, 报告期='2026Q1', 指标编码=f'TEST_{i}', 指标名称=f'测试指标{i}',
                        数值=0 if i%3==0 else '未披露' if i%3==1 else i/7, 单位='万元',
                        期间口径='本季度末数') for i in range(rows-3))
    return standard_workbook_bytes(pd.DataFrame(records))


class Step5ReadPerformanceTests(unittest.TestCase):
    def test_formula_scan_never_random_accesses_read_only_cells(self):
        data = standard_fixture()
        with patch.object(ReadOnlyWorksheet, 'cell', side_effect=AssertionError('Quadratic read-only cell access')):
            frame = read_standard_workbook(data, 'test.xlsx')
        self.assertEqual(len(frame),180)
        rows = frame.set_index('指标编码')
        self.assertEqual(rows.loc['ACTUAL_CAPITAL_TO_RECOGNIZED_ASSETS','数值'],.2)
        self.assertEqual(rows.loc['TEST_0','数值'],0)
        self.assertTrue(pd.isna(rows.loc['TEST_1','数值']))

    def test_formula_scan_is_one_stream_even_for_larger_files(self):
        data = standard_fixture(400)
        original = ReadOnlyWorksheet.iter_rows
        streams = []
        def counted(sheet, *args, **kwargs):
            if not sheet.parent.data_only:
                streams.append(sheet.title)
            return original(sheet,*args,**kwargs)
        with patch.object(ReadOnlyWorksheet,'iter_rows',counted):
            result = read_standard_workbook(data,'large.xlsx')
        self.assertEqual(streams,['标准数据'])
        self.assertEqual(len(result),400)

    def test_header_offsets_and_formula_fallback_are_preserved(self):
        data = standard_fixture(12)
        book = load_workbook(io.BytesIO(data))
        book['标准数据'].insert_rows(1,amount=11)
        book['标准数据']['A1']='说明：第十二行为表头'
        output = io.BytesIO()
        book.save(output)
        book.close()
        frame = read_standard_workbook(output.getvalue(),'offset.xlsx')
        self.assertEqual(len(frame),12)
        self.assertEqual(frame.set_index('指标编码').loc['ACTUAL_CAPITAL_TO_RECOGNIZED_ASSETS','数值'],.2)

    def test_invalid_header_closes_workbook(self):
        book = Workbook()
        book.active.append(['不是标准表'])
        output = io.BytesIO()
        book.save(output)
        book.close()
        from services import solvency_dataset_adapter as adapter
        original = adapter.load_workbook
        closed = []
        def tracked(*args,**kwargs):
            model = original(*args,**kwargs)
            close = model.close
            def close_book():
                closed.append(True)
                close()
            model.close = close_book
            return model
        with patch.object(adapter,'load_workbook',tracked):
            with self.assertRaisesRegex(ValueError,'未找到标准窄表表头'):
                adapter.read_standard_workbook(output.getvalue(),'bad.xlsx')
        self.assertEqual(closed,[True])

    def test_first_sheet_fallback_and_user_overrides_remain(self):
        book = load_workbook(io.BytesIO(standard_fixture(8)))
        sheet = book['标准数据']
        sheet.title='用户数据'
        headers = [c.value for c in sheet[1]]
        sheet.cell(4,headers.index('数值')+1,0.35)
        output = io.BytesIO()
        book.save(output)
        book.close()
        frame = read_standard_workbook(output.getvalue(),'edited.xlsx')
        self.assertEqual(frame.set_index('指标编码').loc['ACTUAL_CAPITAL_TO_RECOGNIZED_ASSETS','数值'],.35)

    def test_step5_download_is_deferred_and_cache_covers_batch(self):
        tree = ast.parse((Path(__file__).parents[1]/'app.py').read_text(encoding='utf-8-sig'))
        download = next(node for node in ast.walk(tree) if isinstance(node,ast.Call)
                        and isinstance(node.func,ast.Attribute) and node.func.attr=='download_button'
                        and node.args and isinstance(node.args[0],ast.Constant) and node.args[0].value=='下载行业集成数据')
        self.assertIsInstance(download.args[1],ast.Lambda)
        cached = next(node for node in ast.walk(tree) if isinstance(node,ast.FunctionDef) and node.name=='cached_read_standard_workbook')
        capacity = next(kw.value.value for dec in cached.decorator_list if isinstance(dec,ast.Call)
                        for kw in dec.keywords if kw.arg=='max_entries')
        self.assertGreaterEqual(capacity,100)

    def test_step5_preview_progress_confirmation_and_deferred_export(self):
        from streamlit.testing.v1 import AppTest
        from services import solvency_dataset_adapter as adapter
        source = (Path(__file__).parents[1]/'app.py').read_text(encoding='utf-8-sig')
        tree = ast.parse(source)
        lines = source.splitlines()
        helpers = []
        for node in tree.body:
            if isinstance(node,ast.FunctionDef) and node.name in {'cached_read_standard_workbook','cached_complete_step5_metrics'}:
                start = min([node.lineno]+[d.lineno for d in node.decorator_list])-1
                helpers.append('\n'.join(lines[start:node.end_lineno]))
        section = source[source.index('    with tabs[5]:'):source.index('\nraw_analysis_source =')]
        setup = '''
import pandas as pd
import streamlit as st
from types import SimpleNamespace
from time import perf_counter
from services.solvency_dataset_adapter import read_standard_workbook, add_missing_derived_metrics, standard_workbook_bytes
from services.solvency_normalizer import STANDARD_COLUMNS, narrow_table_view, standardize_uploaded_frame
from services.solvency_display import dataframe_for_display
upgrade_standard_frame = standardize_uploaded_frame
active_profile = SimpleNamespace(profile_id='LIFE_SOLVENCY')
class Tab:
    open = True
    def __enter__(self):
        self.container=st.container()
        return self.container.__enter__()
    def __exit__(self,*args):
        return self.container.__exit__(*args)
tabs=[Tab() for _ in range(9)]
for key in ['integration_preview','integration_sheet_summary','integration_mapping_summary','integration_logic_checks','integrated_data']:
    if key not in st.session_state:
        st.session_state[key]=pd.DataFrame(columns=STANDARD_COLUMNS)
if 'integration_warnings' not in st.session_state:
    st.session_state.integration_warnings=[]
    st.session_state.integration_preview_mode=''
'''
        script = setup+'\n'+'\n\n'.join(helpers)+'\nif True:\n'+section
        uploads = []
        for i in range(2):
            item = io.BytesIO(standard_fixture(12,f'Step5界面测试{i}'))
            item.name = f'ui_test_{i}.xlsx'
            uploads.append(item)
        with patch('streamlit.file_uploader',return_value=uploads), \
             patch.object(adapter,'read_standard_workbook',wraps=adapter.read_standard_workbook) as reader, \
             patch.object(adapter,'standard_workbook_bytes',wraps=adapter.standard_workbook_bytes) as export:
            app = AppTest.from_string(script).run(timeout=15)
            app.button(key='preview_standard_data').click().run(timeout=15)
            self.assertEqual(len(app.exception),0)
            self.assertEqual(len(app.error),0)
            self.assertEqual(reader.call_count,2)
            self.assertEqual(app.get('progress')[0].proto.value,100)
            self.assertTrue(app.session_state['integrated_data'].empty)
            app.button(key='preview_standard_data').click().run(timeout=15)
            self.assertEqual(reader.call_count,2)
            app.button(key='confirm_integrated_data').click().run(timeout=15)
            self.assertEqual(len(app.exception),0)
            self.assertFalse(app.session_state['integrated_data'].empty)
            self.assertEqual(export.call_count,0)


if __name__ == '__main__':
    unittest.main()
